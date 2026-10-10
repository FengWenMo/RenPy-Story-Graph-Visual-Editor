"""为什么 `append_rows_to_table('character', ...)` 返回 0？（临时探针）"""

import sys
import tempfile
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook                       # noqa: E402
from storygraph.assets import asset_registry as ar                 # noqa: E402

print('_STANDARD_TABLES 里 character ->', ar._STANDARD_TABLES.get('character'))

with tempfile.TemporaryDirectory(prefix='aw_') as td:
    td = Path(td)
    wb = Workbook()
    ws = wb.active
    ws.append(['ID', 'Name', 'Color', 'ImageTag', 'Aliases', 'Notes'])
    ws.append(['mm', '妈妈', '', '', 'MOM', ''])
    wb.save(td / 'characters.xlsx')

    print('\n--- 追加新行')
    n = ar.append_rows_to_table(td, 'character', [{'ID': 'TJF', 'Name': 'TJF',
                                                  'Notes': 'x'}])
    print('  ->', n)
    print('  表:', [list(r) for r in
                   load_workbook(td / 'characters.xlsx').active.iter_rows(
                       values_only=True)])

    print('\n--- 更新已有行（别名列）')
    n = ar.append_rows_to_table(td, 'character', [{'ID': 'mm', 'Aliases': 'MOM|MM2'}],
                               update_keys=('Aliases',))
    print('  ->', n)
    print('  表:', [list(r) for r in
                   load_workbook(td / 'characters.xlsx').active.iter_rows(
                       values_only=True)])

    print('\n--- 直接复现内部逻辑（看异常）')
    wb2 = ar.openpyxl.load_workbook(td / 'characters.xlsx')
    ws2 = wb2.active
    headers = ar._sheet_headers(ws2)
    print('  headers =', headers)
    try:
        ar._update_row_cells(ws2, 2, headers, {'Aliases': 'MOM|MM3'}, ('Aliases',))
        print('  _update_row_cells ok；cell =', ws2.cell(row=2, column=5).value)
        wb2.save(td / 'characters.xlsx')
        print('  save ok；读回 =',
              load_workbook(td / 'characters.xlsx').active.cell(row=2, column=5).value)
    except Exception:
        traceback.print_exc()
