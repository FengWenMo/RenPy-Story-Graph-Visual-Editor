"""
探针：**更刁的四种结构** —— 洗牌 / 嵌套 / translate / extend·voice。

前两轮的结构场景都是"一条直路 + 一个分支"。这一轮挑真实项目里存在、
但工具可能没见过形状的情况：

  A. **整章洗牌**：把 5 拍随机重排（策划重写一场戏的常见动作）
  B. **嵌套**：`menu` 的选项体里再嵌 `if`，删掉 if 分支里的唯一语句
  C. **`translate` 块**：多语言项目的 `translate chinese strings:` / 一段
     `translate chinese ch01_abc:` —— 工具认不出时**必须原样保真**，
     绝不能当"不认识的拍"删掉
  D. **`extend` / `voice` / `with`**：同一拍里的附加行

判据：产物必须能过 Ren'Py 引擎 `lint`（parse error = P0），
且 `translate` 块的内容一个字都不许少（那是译文，丢了就是丢工作量）。

用法：
    D:/Python/python.exe tools/probe_hard_structures.py
"""
from __future__ import annotations

import random
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.app import tables                      # noqa: E402
from storygraph.pipeline.engine_lint import _sdk, lint_text   # noqa: E402

SRC_SHUFFLE = '''\
label ch01:
    e "第1句"
    e "第2句"
    e "第3句"
    e "第4句"
    e "第5句"
    return
'''

SRC_NEST_MENU = '''\
label ch02:
    e "A1"
    menu:
        "去左边":
            e "A2 左边"
            if has_key:
                e "A3 有钥匙"
            e "A4 左边继续"
        "去右边":
            e "A5 右边"
    e "A6"
    return
'''

SRC_TRANSLATE = '''\
label ch03:
    e "B1"
    e "B2"
    return

translate chinese strings:
    old "Start"
    new "开始"

translate chinese ch03_abc:
    e "B1 的译文"
'''

SRC_EXTRAS = '''\
label ch04:
    voice "voice/c1.ogg"
    e "C1 有语音"
    e "C2"
    extend "（接上一句）"
    with dissolve
    e "C3"
    return
'''


def sheet(table: Path):
    import openpyxl
    wb = openpyxl.load_workbook(str(table))
    ws = wb['Beats']
    heads = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    rows = []
    for r in range(2, ws.max_row + 1):
        kind = str(ws.cell(r, heads['类型']).value or '').strip()
        text = str(ws.cell(r, heads['台词']).value or '').strip()
        if kind or text:
            rows.append({'row': r, 'kind': kind, 'text': text})
    return wb, ws, heads, rows


def run(name: str, body: str, mutate, work: Path, *, must_contain=()) -> dict:
    root = work / name
    (root / 'game').mkdir(parents=True, exist_ok=True)
    src = root / 'game' / 'ch.rpy'
    src.write_text(body, encoding='utf-8')
    gen = tables.generate_table(src, project_root=root)
    if not gen.get('ok'):
        return {'name': name, 'verdict': 'SKIP', 'note': str(gen.get('error'))[:60]}
    table = gen['table_path']
    wb, ws, heads, rows = sheet(table)
    note = mutate(ws, heads, rows)
    wb.save(str(table))
    res = tables.export_table(table, project_root=root, source_rpy=src)
    done = tables.apply_export(res, project_root=root)
    d = {'name': name, 'note': note or '', 'written': bool(done.get('written')),
         'patch': bool(res.get('used_patch')),
         'patch_note': str(res.get('note'))[:80]}
    if not d['written']:
        d['verdict'] = '拒写'
        d['text'] = src.read_text(encoding='utf-8')
        d['note'] += '（' + str(res.get('blocking') or done.get('error'))[:50] + '）'
        return d
    text = src.read_text(encoding='utf-8')
    d['text'] = text
    missing = [m for m in must_contain if m not in text]
    if missing:
        d['verdict'] = 'P0'
        d['note'] += ' 丢了：' + '、'.join(missing)
        return d
    rep = lint_text(text)
    d['ran'] = rep.ran
    d['errors'] = rep.parse_errors
    d['verdict'] = ('P0' if (rep.ran and rep.parse_errors)
                    else ('PASS' if rep.ran else 'SKIP'))
    if not rep.ran:
        d['note'] += f'（lint 没跑：{rep.context_note[:40]}）'
    return d


def m_shuffle(ws, h, rows):
    rnd = random.Random(20261021)
    texts = [r['text'] for r in rows]
    rnd.shuffle(texts)
    for r, t in zip(rows, texts):
        ws.cell(r['row'], h['台词']).value = t
    return '洗牌：' + ' '.join(texts)


def m_nested_delete(ws, h, rows):
    for r in rows:
        if r['text'] == 'A3 有钥匙':
            ws.delete_rows(r['row'])
            return '删掉 menu→if 里唯一那句'
    return None


def m_translate_edit(ws, h, rows):
    for r in rows:
        if r['text'] == 'B2':
            ws.cell(r['row'], h['台词']).value = 'B2 改过'
            return '改 B2（patch 里带着 translate 块）'
    return None


def m_extras_edit(ws, h, rows):
    for r in rows:
        if r['text'] == 'C2':
            ws.cell(r['row'], h['台词']).value = 'C2 改过'
            return '改 C2（同一拍里有 voice/extend/with）'
    return None


CASES = [
    ('洗牌', SRC_SHUFFLE, m_shuffle, ()),
    ('嵌套删一拍', SRC_NEST_MENU, m_nested_delete, ()),
    ('translate块', SRC_TRANSLATE, m_translate_edit,
     ('translate chinese strings:', 'new "开始"', 'translate chinese ch03_abc:',
      'B1 的译文')),
    ('voice/extend/with', SRC_EXTRAS, m_extras_edit,
     ('voice "voice/c1.ogg"', 'extend "（接上一句）"', 'with dissolve')),
]


def main() -> int:
    print(f'SDK：{_sdk()}')
    rows = []
    with tempfile.TemporaryDirectory(prefix='hard_') as td:
        work = Path(td)
        for name, body, fn, must in CASES:
            try:
                r = run(name, body, fn, work, must_contain=must)
            except Exception as e:              # noqa: BLE001
                import traceback
                traceback.print_exc()
                r = {'name': name, 'verdict': f'抛异常 {type(e).__name__}',
                     'note': str(e)[:70]}
            rows.append(r)
            print('=' * 76)
            print(f'[{name}] {r.get("note", "")}')
            print(f'  判定 = {r["verdict"]}   written={r.get("written")} '
                  f'patch={r.get("patch")}')
            print(f'  {r.get("patch_note", "")}')
            if r.get('text'):
                for ln in r['text'].splitlines():
                    print(f'     | {ln}')
            for e in (r.get('errors') or [])[:3]:
                print(f'  ✗ 引擎：{e}')
    print('\n' + '=' * 76)
    for r in rows:
        print(f'{r["name"]:18} {r["verdict"]:8} {r.get("note", "")[:44]}')
    bad = [r for r in rows if r['verdict'] == 'P0']
    print(f'\n★ P0：{len(bad)} / {len(rows)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
