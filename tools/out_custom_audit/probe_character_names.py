"""角色表 / 变量名 / show tag 审计：编剧看到的候选能填进剧本吗？

    python tools/out_custom_audit/probe_character_names.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
os.environ.setdefault('QT_QPA_FONTDIR', r'C:\Windows\Fonts')
os.environ.setdefault('PYTHONIOENCODING', 'utf-8')

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(HERE))

from PySide6.QtWidgets import QApplication                        # noqa: E402

from audit_project import build_audit_project                     # noqa: E402

PROJ = HERE / 'proj_names_audit'
SHOTS = HERE / 'shots'
SHOTS.mkdir(parents=True, exist_ok=True)
root = build_audit_project(PROJ)

from storygraph.authoring.learn_statement import (choice_sources,  # noqa: E402
                                                  suggest_shape)
from storygraph.story import statement_registry as SR              # noqa: E402
from storygraph.ui_qt.dialog_learn import CustomStatementDialog    # noqa: E402

print('=' * 80)
print('① 角色表（characters.xlsx）里到底有哪几列、哪列被藏起来了')
import openpyxl                                                    # noqa: E402
ws = openpyxl.load_workbook(root / 'asset_tables' / 'characters.xlsx').active
headers = [c.value for c in ws[1]]
print('   表头：', headers)
for i, h in enumerate(headers, 1):
    letter = openpyxl.utils.get_column_letter(i)
    hidden = ws.column_dimensions[letter].hidden
    print(f'   第{i}列 {letter} {h!r} hidden={hidden}')
print('   数据行：', [c.value for c in ws[2]])

print('\n' + '=' * 80)
print('② 自定义模块里的「角色表」来源给了什么值')
src = choice_sources(root, 'show')
role_src = next((s for s in src if s.get('key') == 'roles'), {})
print('   角色表来源：', role_src)
print('   —— 编剧能看到的值：', role_src.get('values'))
print('   —— 显示名标签：', role_src.get('labels'))

shape = suggest_shape(root, 'show')
print('\n   suggest_shape(show) 第 1 项：')
for item in (shape.get('items') or [])[:3]:
    print('   ', {k: item.get(k) for k in
                  ('index', 'kind', 'name', 'note', 'choices_from', 'choices',
                   'labels')})

print('\n' + '=' * 80)
print('③ 弹窗「试切」：真剧本的 show tag / 变量名 vs 显示名')
app = QApplication.instance() or QApplication([])
SR.clear_cache()
dlg = CustomStatementDialog(None, project_root=root, discovered=[],
                            registry_report=SR.format_registry_report(None))
dlg.name.setText('show')
dlg.params.clear()
role_vals = list(role_src.get('values') or [])
c1 = dlg.params.add_param('角色', '角色', ','.join(role_vals), kind='choice')
c1._limit.setChecked(True)
c2 = dlg.params.add_param('服装', '服装', 'swim_suit,uniform', kind='choice')
c2._limit.setChecked(True)
c3 = dlg.params.add_param('动作', '动作', 'stand,wave', kind='choice')
c3._limit.setChecked(True)
dlg.resize(1100, 760)
dlg.show()
app.processEvents()


def try_cut(line):
    dlg.sample.setText(line)
    dlg._refresh_cut()
    app.processEvents()
    print(f'   试切 {line!r} -> {dlg.cut_view.text()!r}')


print('   —— 角色表来源给的是变量名 `lan`，但剧本里 `show` 要的是 ImageTag `lan_tag`：')
dlg.tabs.setCurrentIndex(1)
app.processEvents()
try_cut('show lan_tag swim_suit stand')
try:
    bar = dlg.declare_scroll.verticalScrollBar()
    bar.setValue(bar.maximum())
    app.processEvents()
    dlg.grab().save(str(SHOTS / '11b_cut_valid_tag_rejected.png'))
    print('   截图（有效 tag 被判切不上）：', SHOTS / '11b_cut_valid_tag_rejected.png')
except Exception as e:
    print('   （截切截图失败）', e)
try_cut('show lan swim_suit stand')
try:
    bar = dlg.declare_scroll.verticalScrollBar()
    bar.setValue(bar.maximum())
    app.processEvents()
    dlg.grab().save(str(SHOTS / '11c_cut_invalid_var_accepted.png'))
    print('   截图（无效变量被判切上）：', SHOTS / '11c_cut_invalid_var_accepted.png')
except Exception as e:
    print('   （截切截图失败2）', e)
print('   —— 显示名（岚）不再是候选：')
try_cut('show 岚 swim_suit stand')
dlg.tabs.setCurrentIndex(1)
app.processEvents()
dlg.grab().save(str(SHOTS / '11_dialog_cut_display_name.png'))
print('   截图：', SHOTS / '11_dialog_cut_display_name.png')

# 只看参数卡上的候选块，也顺手截一张
dlg2 = CustomStatementDialog(None, project_root=root, discovered=[],
                             registry_report=SR.format_registry_report(None))
dlg2.params.clear()
cc = dlg2.params.add_param('角色', '角色', '', kind='choice')
cc._values.add_values(role_vals,
                      labels=dict(role_src.get('labels') or {}))
cc._limit.setChecked(True)
print('   角色芯片显示：', list(cc._values._chips))
dlg2.resize(1100, 600)
dlg2.show()
app.processEvents()
dlg2.grab().save(str(SHOTS / '12_dialog_role_chip_display_name.png'))
print('   角色候选卡截图：', SHOTS / '12_dialog_role_chip_display_name.png')

print('\n' + '=' * 80)
print('④ 「资产 / 角色统计」里角色台词量那一栏的取键')
from storygraph.app.asset_panel import AssetPanelData               # noqa: E402
from storygraph.parsing.character_script_linker import CharacterScriptLinker  # noqa: E402
# 真实返回形状（见 CharacterScriptLinker.character_activity_ranking）
row = {'variable_name': 'lan', 'display_name': '岚',
       'dialogue_count': 12, 'label_count': 2}
print('   数据源给的键：', row)
print('   show_asset_overview 读的键：name / character / count / lines')
print('   → 会显示成：', f'{row.get("name") or row.get("character")}：'
                       f'{row.get("count") or row.get("lines")}')

# 真的跑一遍那个方法（把模态 TextViewDialog 换掉）
from storygraph.ui_qt import window_graph_ops as WGO               # noqa: E402


class _DummyDialog:
    def __init__(self, *a, **k):
        pass

    def exec(self):
        return 0


class _S:
    def asset_overview(self):
        return {'ok': True, 'statistics': {}, 'categories': {},
                'top_characters': [row], 'undefined': [],
                'unused_characters': [], 'unused_assets': [], 'broken': []}


class _Host:
    _session = _S()

    def _log(self, *a):
        pass

    def _status(self, *a):
        pass


orig = WGO.__dict__.get('TextViewDialog')
import storygraph.ui_qt.dialogs as _dialogs                       # noqa: E402
_dialogs.TextViewDialog = _DummyDialog
try:
    text = WGO.GraphOpsMixin.show_asset_overview(_Host())
finally:
    if orig is not None:
        _dialogs.TextViewDialog = orig
print('   实际生成的报告片段：',
      [ln for ln in text.splitlines() if 'lan' in ln or 'None' in ln])

dlg.close()
dlg2.close()
print('\n完成。截图目录：', SHOTS)
