"""
探针（走**真实路径**）：单元格里的控制字符会不会变成产物里的物理换行。

上一版直接给 `Beat` 字段塞新行，**绕过了 hydrate**，可能不是真实可达的。
这一版走完整链路：

    .rpy → beats → xlsx → **用 openpyxl 往单元格里塞真换行/真 TAB**
         → import_beats_from_excel（会 hydrate）→ render

单元格里出现真换行是**真实可达**的：Excel 的 Alt+Enter、从别处粘一段
多行文本、或脚本写表，都会留下它。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.assets.asset_registry import AssetRegistry
from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
from storygraph.story.beat_build import beats_from_parse_result
from storygraph.story.beat_model import BEAT_FIELDS
from storygraph.story.beat_render import make_symbols, render_beats_to_rpy
from storygraph.story.excel_export import export_beats_to_excel
from storygraph.story.excel_import import import_beats_from_excel

SRC = ('label a:\n'
       '    show ling happy\n'
       '    ling "第一句"\n'
       '    ling "第二句"\n')

#: 往哪一列塞什么，看产物会不会断行
CASES = {
    '说话人 + 真换行': ('Speaker', 'ling\nX'),
    '说话人 + 真制表符': ('Speaker', 'ling\tX'),
    '立绘 + 真换行': ('Show', 'ling\nhappy'),
    '背景 + 真换行': ('Scene', 'bg\nroom'),
    '位置 + 真换行': ('At', 'left\nright'),
    '语音 + 真换行': ('Voice', 'v/a\nb.ogg'),
    '音乐 + 真换行': ('BGM', 'bgm/a\nb.mp3'),
    '台词 + 真换行': ('Text', '第一行\n第二行'),
}


def _loose(text: str) -> list:
    ok = ('label ', 'init', 'default', 'define', 'transform', 'screen',
          'image ', 'style ', 'pass', 'return', '#', 'menu', 'if ', 'elif ',
          'else', 'while ', 'for ')
    out = []
    for i, ln in enumerate(text.splitlines(), 1):
        s = ln.strip()
        if not s or ln[:1].isspace() or s.startswith(ok):
            continue
        out.append((i, ln))
    return out


def main() -> int:
    reg = AssetRegistry()
    syms = make_symbols(reg)
    bad = 0
    for name, (col, value) in CASES.items():
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            f = tmp / 's.rpy'
            f.write_text(SRC, encoding='utf-8')
            pr, _ = IncrementalRenPyAnalyzer().analyze(files=[f], graph=None)
            xl = tmp / 's.xlsx'
            export_beats_to_excel(beats_from_parse_result(pr), xl)

            from openpyxl import load_workbook
            wb = load_workbook(xl)
            ws = wb['Beats']
            headers = [c.value for c in ws[1]]
            col_i = headers.index(
                {'Speaker': '说话人', 'Show': '立绘', 'Scene': '背景',
                 'At': '位置', 'Voice': '语音', 'BGM': '音乐',
                 'Text': '台词'}[col])
            # 第 2 行（第一张拍）与第 3 行都塞，覆盖"有/无台词"两种
            for r in (2, 3):
                ws.cell(row=r, column=col_i + 1).value = value
            wb.save(xl)

            back = import_beats_from_excel(xl)
            out = render_beats_to_rpy(back, registry=reg, symbols=syms)

            f2 = tmp / 'o.rpy'
            f2.write_text(out, encoding='utf-8')
            try:
                pr2, _ = IncrementalRenPyAnalyzer().analyze(files=[f2], graph=None)
                info = f'{len(pr2.statements)} 条语句'
                err = ''
            except Exception as e:                   # noqa: BLE001
                info, err = '-', f'{type(e).__name__}: {e}'

        loose = _loose(out)
        flag = '✗' if (loose or err) else '✓'
        print(f'{flag} {name}（{col} 列）')
        for ln in out.splitlines()[2:]:
            print('      ', repr(ln))
        if loose:
            print(f'      被断开的尾巴：{loose[:2]}')
            bad += 1
        if err:
            print(f'      回读失败：{err}')
            bad += 1
        elif not loose:
            print(f'      回读：{info}')
    print('\n结论：', '单元格控制字符会写坏产物' if bad else '单元格控制字符被挡住了')
    return 1 if bad else 0


if __name__ == '__main__':
    raise SystemExit(main())
