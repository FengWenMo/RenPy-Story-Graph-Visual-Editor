#!/usr/bin/env python
"""
极为变态的测试：随机差分 + 变异 + 脏数据 + 嵌套地狱 + 规模压测。

与 torture_test（手挑用例）的区别：这里**不假设任何具体形态**，
用随机生成去撞隐含假设，用变异去打鲁棒性，用脏数据去打边界。

  F1  随机剧本差分：随机生成器产出合法 .rpy，往返后比对 statement 召回
  F2  真实项目变异：删行 / 交换 / 插入 / 改缩进，要求不崩且产物可解析
  F3  Excel 脏数据注入：None / 超长 / 换行 / 公式前缀 / emoji / 注入串
  F4  嵌套地狱：if 套 if 套(if 内 menu 内 if)、分支内 show 层层嵌套
  F5  规模压测：100 选项 menu、3000 行台词、50 个 label
  F6  五次往返幂等（不是三次）
  F7  对抗性边界：空文件、只有 label、超大单行、重复 label

判据统一：不崩 + 产物可被解析器接受 + 内容 statement 全召回。
"""
from __future__ import annotations

import random
import re
import sys
import tempfile
import traceback
from collections import Counter
from pathlib import Path

PROJECT = Path(r'D:\千连万理\RenPyTest\sr_general\game')
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer  # noqa
from storygraph.story.beat_build import beats_from_parse_result        # noqa
from storygraph.story.excel_export import export_beats_to_excel        # noqa
from storygraph.story.excel_import import import_beats_from_excel      # noqa
from storygraph.pipeline.compile import compile_beats                  # noqa

import openpyxl  # noqa

TMP = Path(tempfile.mkdtemp(prefix='brutal_'))
A = IncrementalRenPyAnalyzer()

CONTENT_KINDS = {'say', 'show', 'scene', 'hide', 'play', 'voice',
                 'python', 'comment', 'menu_option', 'menu', 'extend'}
STATS = Counter()
FAILS = []


def norm(x):
    return re.sub(r'\s+', ' ', str(x or '')).strip()


def key(s):
    return (s.kind, tuple(sorted((k, norm(v)) for k, v in s.payload.items())))


def rt_file(path: Path, tag='x'):
    """原文文件 → Excel → 产物文本；返回 (产物文本, 解析到的 statements)"""
    pr, _ = A.analyze(files=[path], graph=None)
    beats = beats_from_parse_result(pr)
    if not beats:
        return None, pr.statements
    xl = TMP / f'{tag}.xlsx'
    export_beats_to_excel(beats, xl)
    res = compile_beats(import_beats_from_excel(xl))
    out = TMP / f'{tag}_o.rpy'
    out.write_text(res.rpy_text, encoding='utf-8')
    pr2, _ = A.analyze(files=[out], graph=None)
    return res.rpy_text, pr2.statements


def recall_loss(src_stmts, out_stmts):
    """原文的内容 statement，产物里有没有（多重集差）。"""
    ca = Counter(key(s) for s in src_stmts if s.kind in CONTENT_KINDS)
    cb = Counter(key(s) for s in out_stmts if s.kind in CONTENT_KINDS)
    return ca - cb


def report(name, ok, detail=''):
    STATS['total'] += 1
    if ok:
        STATS['pass'] += 1
    else:
        STATS['fail'] += 1
        FAILS.append((name, detail))
    flag = 'PASS' if ok else 'FAIL'
    print(f'  [{flag}] {name}' + (f'  — {detail}' if detail else ''))


# ==================================================================
#  F1  随机剧本差分
# ==================================================================
SPEAKERS = ['e', 'q', 'mm', 'ling']
TEXTS = ['你好', '这是一个测试。', 'Ciallo~', '嗯……', '带"引号"的话',
         '含{w}标签', '尾部有空格 ', '[[转义']
IMAGES = ['bg room', 'bg cafe', 'ling happy', 'ling sad']
POSES = ['left', 'right', 'h11', 'center']
CONDS = ['flag', 'hp > 3', 'x == 1', 'pomoplace == "cafe"', '_return == "ok"']


def gen_label_body(rng, indent, depth, budget):
    """随机产出一段 label 体（可含嵌套块）。"""
    out = []
    n = rng.randint(1, 4)
    for _ in range(n):
        if budget[0] <= 0:
            break
        budget[0] -= 1
        r = rng.random()
        pad = ' ' * indent
        if r < 0.30:
            out.append(f'{pad}{rng.choice(SPEAKERS)} "{rng.choice(TEXTS)}"')
        elif r < 0.40:
            out.append(f'{pad}"{rng.choice(TEXTS)}"')
        elif r < 0.55:
            at = f' at {rng.choice(POSES)}' if rng.random() < 0.4 else ''
            wth = ' with dissolve' if rng.random() < 0.3 else ''
            out.append(f'{pad}show {rng.choice(IMAGES)}{at}{wth}')
        elif r < 0.62:
            out.append(f'{pad}hide {rng.choice(IMAGES)}')
        elif r < 0.70:
            out.append(f'{pad}play music "bgm/{rng.randint(1,3)}.mp3"')
        elif r < 0.78:
            out.append(f'{pad}$ score += {rng.randint(1,9)}')
        elif r < 0.84:
            out.append(f'{pad}# 注释 {rng.randint(1,99)}')
        elif r < 0.92 and depth < 2:
            # if / elif 链，分支内可能带 show
            out.append(f'{pad}if {rng.choice(CONDS)}:')
            if rng.random() < 0.6:
                out.append(f'{pad}    show {rng.choice(IMAGES)}')
            out += gen_label_body(rng, indent + 4, depth + 1, budget)
            for _ in range(rng.randint(0, 2)):
                out.append(f'{pad}elif {rng.choice(CONDS)}:')
                if rng.random() < 0.6:
                    out.append(f'{pad}    show {rng.choice(IMAGES)}')
                out += gen_label_body(rng, indent + 4, depth + 1, budget)
            if rng.random() < 0.5:
                out.append(f'{pad}else:')
                out += gen_label_body(rng, indent + 4, depth + 1, budget)
        else:
            # menu
            if depth < 2:
                out.append(f'{pad}menu:')
                for k in range(rng.randint(1, 3)):
                    out.append(f'{pad}    "分支{k+1}":')
                    if rng.random() < 0.7:
                        out.append(f'{pad}        jump lbl{k+1}')
                    else:
                        out += gen_label_body(rng, indent + 8, depth + 1, budget)
            else:
                out.append(f'{pad}jump lbl{rng.randint(1,3)}')
    return out


def gen_script(seed):
    rng = random.Random(seed)
    lines = []
    for li in range(rng.randint(1, 3)):
        lines.append(f'label generated_{li}:')
        lines += gen_label_body(rng, 4, 0, [6])
        if rng.random() < 0.7:
            lines.append('    return')
    return '\n'.join(lines) + '\n'


def f1_fuzz(n=120):
    print('\n=== F1 随机剧本差分（fuzz）===')
    crash = unparsable = loss_hits = 0
    worst = []
    for seed in range(n):
        src = gen_script(seed)
        f = TMP / f'f1_{seed}.rpy'
        f.write_text(src, encoding='utf-8')
        try:
            pr, _ = A.analyze(files=[f], graph=None)
            txt, out_stmts = rt_file(f, tag=f'f1_{seed}')
        except Exception as e:
            crash += 1
            if len(worst) < 5:
                worst.append((seed, f'异常 {type(e).__name__}: {e}'))
            continue
        if txt is None:
            continue
        try:
            A.analyze(files=[TMP / f'f1_{seed}_o.rpy'], graph=None)
        except Exception:
            unparsable += 1
            worst.append((seed, '产物无法解析'))
            continue
        lost = recall_loss(pr.statements, out_stmts)
        if lost:
            loss_hits += 1
            if len(worst) < 5:
                sample = list(lost)[0]
                worst.append((seed, f'丢 {sum(lost.values())} 条，例：{str(sample)[:70]}'))
    report('F1 无异常崩溃', crash == 0, f'{crash}/{n} 崩')
    report('F1 产物全部可解析', unparsable == 0, f'{unparsable}/{n} 不可解析')
    report('F1 全部样本零丢失', loss_hits == 0, f'{loss_hits}/{n} 有丢失')
    print('  失败样本（前5）：')
    for s, d in worst[:5]:
        print(f'    seed {s}: {d}')
    return worst[:1]


# ==================================================================
#  F2  真实项目变异
# ==================================================================
def mutate(lines, rng):
    ls = list(lines)
    if len(ls) < 3:
        return ls
    op = rng.choice(['del', 'swap', 'insert', 'indent'])
    if op == 'del':
        ls.pop(rng.randrange(len(ls)))
    elif op == 'swap':
        i = rng.randrange(len(ls) - 1)
        ls[i], ls[i + 1] = ls[i + 1], ls[i]
    elif op == 'insert':
        ls.insert(rng.randrange(len(ls)), '    # 变异插入')
    else:
        i = rng.randrange(len(ls))
        ls[i] = '    ' + ls[i]
    return ls


def f2_mutate(n=60):
    print('\n=== F2 真实项目变异（鲁棒性）===')
    files = sorted(set(PROJECT.rglob('*.rpy')))
    rng = random.Random(20261003)
    crash = unparsable = 0
    for i in range(n):
        src_file = files[rng.randrange(len(files))]
        lines = src_file.read_text(encoding='utf-8', errors='replace').splitlines()
        if len(lines) < 5:
            continue
        mutated = mutate(lines, rng)
        f = TMP / f'f2_{i}.rpy'
        f.write_text('\n'.join(mutated) + '\n', encoding='utf-8')
        try:
            txt, _ = rt_file(f, tag=f'f2_{i}')
            if txt is None:
                continue
            A.analyze(files=[TMP / f'f2_{i}_o.rpy'], graph=None)
        except Exception:
            unparsable += 1
    report('F2 变异后不崩且产物可解析', unparsable == 0,
           f'{unparsable}/{n} 出问题')


# ==================================================================
#  F3  Excel 脏数据注入
# ==================================================================
DIRTY = {
    '超长串': 'あ' * 5000,
    '含换行': '第一行\n第二行',
    '公式前缀': '=1+1',
    '前导尾随空格': '   两边有空格   ',
    'emoji': '你好🌸🎉',
    '注入串': '"; drop table beats; --',
    '反斜杠': 'C:\\path\\to\\file',
    '百分号': '100%完成',
    '控制字符样子': 'a\tb\tc',
    '纯空格': '     ',
}


def f3_dirty():
    print('\n=== F3 Excel 脏数据注入 ===')
    base = ('label a:\n'
            '    menu:\n'
            '        "甲":\n'
            '            jump x\n'
            '        "乙":\n'
            '            jump y\n'
            '    e "正常台词"\n')
    f = TMP / 'f3.rpy'
    f.write_text(base, encoding='utf-8')
    pr, _ = A.analyze(files=[f], graph=None)
    beats = beats_from_parse_result(pr)
    xl = TMP / 'f3.xlsx'
    export_beats_to_excel(beats, xl)

    wb = openpyxl.load_workbook(xl)
    ws = wb['Beats']
    hdr = [c.value for c in ws[1]]
    row = 2
    crash_ok, lost_ok = True, True
    detail = []
    for label, dirty in DIRTY.items():
        try:
            wb2 = openpyxl.load_workbook(xl)
            ws2 = wb2['Beats']
            # 往「台词」「说话人」「注释」列灌脏值
            for col_name in ('台词', '注释', '说话人'):
                if col_name in hdr:
                    ci = hdr.index(col_name) + 1
                    ws2.cell(row=2, column=ci, value=dirty)
            p = TMP / f'f3_{label}.xlsx'
            wb2.save(p)
            res = compile_beats(import_beats_from_excel(p))
            out = TMP / f'f3_{label}_o.rpy'
            out.write_text(res.rpy_text, encoding='utf-8')
            A.analyze(files=[out], graph=None)
            if norm(dirty) and not _value_survives(dirty, res.rpy_text):
                detail.append(f'{label} 值未出现在产物')
                lost_ok = False
        except Exception as e:
            detail.append(f'{label} 崩了 {type(e).__name__}: {e}')
            crash_ok = False
    report('F3 脏数据不崩', crash_ok)
    report('F3 脏数据未被静默吞掉', lost_ok, '；'.join(detail[:4]))


def _value_survives(dirty: str, rpy_text: str) -> bool:
    """
    脏值是否**没被吞掉**（考虑到 `.rpy` 里必须做的转义）。

    `C:\\path\\to\\file` 这种值写进 `.rpy` 时，反斜杠会被 `_quote` 再转义
    一层（`C:\\\\path\\\\to\\\\file`）—— 那是**正确**的（不转义才是错的）。
    所以判据不能是"原串逐字包含"，而是"**未转义的原文**逐字包含"。

    纯空格被 trim 属于合理行为，调用方已先排除。`=` 开头的值另有一层：
    Excel 会把它当公式，`data_only=True` 读回来是 `None` —— 那是**真 bug**，
    由 `story.excel_import` 的 `data_only=False` 修掉，这里照旧逐字判。
    """
    if dirty in rpy_text:
        return True
    unescaped = dirty.replace('\\\\', '\\')
    return unescaped in rpy_text or dirty.replace('\\', '\\\\') in rpy_text


# ==================================================================
#  F4  嵌套地狱
# ==================================================================
HELL_SRC = (
    'label hell:\n'
    '    if outer:\n'
    '        show bg_outer\n'
    '        if middle:\n'
    '            show bg_middle\n'
    '            menu:\n'
    '                "深层选项":\n'
    '                    show bg_deep\n'
    '                    if deepest:\n'
    '                        e "最深层"\n'
    '                    else:\n'
    '                        e "最深层的 else"\n'
    '                    jump deep_end\n'
    '                "另一个":\n'
    '                    jump other\n'
    '        else:\n'
    '            show bg_middle_else\n'
    '            e "middle 的 else"\n'
    '    else:\n'
    '        show bg_outer_else\n'
    '        menu:\n'
    '            "外层选项":\n'
    '                jump out\n'
)


def f4_hell():
    print('\n=== F4 嵌套地狱（if×3 + if内menu + menu内if/else）===')
    f = TMP / 'f4.rpy'
    f.write_text(HELL_SRC, encoding='utf-8')
    pr, _ = A.analyze(files=[f], graph=None)
    src_n = sum(1 for s in pr.statements if s.kind in CONTENT_KINDS)
    print(f'  原文内容 statement: {src_n}')
    for s in pr.statements:
        if s.kind in ('say', 'menu_option', 'menu'):
            print(f'    [{s.kind}] {str(s.payload)[:70]}')
    try:
        txt, out_stmts = rt_file(f, tag='f4')
    except Exception as e:
        report('F4 不崩', False, f'{type(e).__name__}: {e}')
        return
    report('F4 不崩', True)
    lost = recall_loss(pr.statements, out_stmts)
    report('F4 零丢失', not lost, f'丢 {sum(lost.values())} 条')
    if lost:
        for k in list(lost)[:6]:
            print(f'      丢: {str(k)[:90]}')
    # 结构检查：`if outer` / `if deepest` 这类要按**条件**找，不能按字面前缀找。
    # 解析器会把嵌套的条件**累积**起来（`if ((outer) and (middle)) and (deepest):`），
    # 所以 `'if deepest' in text` 这种判据必然误报 —— 条件对但字符串对不上。
    for token in ('menu:', 'if outer', 'if deepest', 'bg_deep', '最深层的 else'):
        if token.startswith('if '):
            name = token[3:]
            ok = bool(re.search(
                r'^\s*if\b[^\n]*\b' + re.escape(name) + r'\b', txt or '',
                re.M))
        else:
            ok = token in (txt or '')
        report(f'F4 关键结构 {token!r}', ok)
    print('  --- 产物 ---')
    for l in (txt or '').splitlines():
        print(f'    | {l}')


# ==================================================================
#  F5  规模压测
# ==================================================================
def f5_scale():
    print('\n=== F5 规模压测 ===')
    # 5a: 100 选项
    lines = ['label big:']
    lines.append('    menu:')
    for i in range(100):
        lines.append(f'        "选项{i}":')
        lines.append(f'            jump target{i}')
    f = TMP / 'f5a.rpy'
    f.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    try:
        pr, _ = A.analyze(files=[f], graph=None)
        txt, out_stmts = rt_file(f, tag='f5a')
        lost = recall_loss(pr.statements, out_stmts)
        report('F5a 100 选项零丢失', not lost,
               f'丢 {sum(lost.values())} 条')
    except Exception as e:
        report('F5a', False, f'{type(e).__name__}: {e}')

    # 5b: 3000 行台词
    lines = ['label many:']
    for i in range(3000):
        lines.append(f'    e "第 {i} 句话。"')
    f = TMP / 'f5b.rpy'
    f.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    try:
        pr, _ = A.analyze(files=[f], graph=None)
        txt, out_stmts = rt_file(f, tag='f5b')
        lost = recall_loss(pr.statements, out_stmts)
        report('F5b 3000 句零丢失', not lost,
               f'丢 {sum(lost.values())} 条')
    except Exception as e:
        report('F5b', False, f'{type(e).__name__}: {e}')

    # 5c: 50 个 label
    lines = []
    for li in range(50):
        lines.append(f'label multi_{li}:')
        lines.append(f'    e "章节 {li}"')
        lines.append('    return')
    f = TMP / 'f5c.rpy'
    f.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    try:
        txt, _ = rt_file(f, tag='f5c')
        n = sum(1 for l in (txt or '').splitlines() if l.startswith('label '))
        report('F5c 50 个 label 全在', n == 50, f'产物 {n} 个')
    except Exception as e:
        report('F5c', False, f'{type(e).__name__}: {e}')


# ==================================================================
#  F6  五次往返幂等
# ==================================================================
def f6_idempotent():
    print('\n=== F6 五次往返幂等 ===')
    src = ('label a:\n'
           '    menu:\n'
           '        "甲":\n'
           '            jump x\n'
           '        "乙":\n'
           '            show bg1 at left\n'
           '            jump y\n'
           '    if hp > 3:\n'
           '        show hero\n'
           '        e "战斗"\n'
           '    elif hp > 1:\n'
           '        show hero weak\n'
           '        e "苦战"\n')
    f = TMP / 'f6.rpy'
    f.write_text(src, encoding='utf-8')
    pr, _ = A.analyze(files=[f], graph=None)
    beats = beats_from_parse_result(pr)
    sigs = []
    for i in range(5):
        xl = TMP / f'f6_{i}.xlsx'
        export_beats_to_excel(beats, xl)
        beats = import_beats_from_excel(xl)
        res = compile_beats(beats)
        sigs.append(re.sub(r'\s+', ' ', res.rpy_text).strip())
    stable = len(set(sigs)) == 1
    report('F6 五次产物完全一致', stable)
    if not stable:
        for i, s in enumerate(sigs):
            print(f'    轮{i+1}: {s[:150]}')


# ==================================================================
#  F7  对抗性边界
# ==================================================================
def f7_edges():
    print('\n=== F7 对抗性边界 ===')
    cases = {
        '空 label': 'label e1:\n',
        '只有 return': 'label e2:\n    return\n',
        '超大单行台词': 'label e3:\n    e "' + '字' * 20000 + '"\n',
        '重复 label': 'label e4:\n    e "一"\nlabel e4:\n    e "二"\n',
        '无 label': '    e "没有label"\n',
        '只有注释': 'label e5:\n    # 只有注释\n',
        'menu 无选项': 'label e6:\n    menu:\n',
        '空 if': 'label e7:\n    if flag:\n',
    }
    for name, src in cases.items():
        f = TMP / f'f7_{abs(hash(name)) % 9999}.rpy'
        f.write_text(src, encoding='utf-8')
        try:
            pr, _ = A.analyze(files=[f], graph=None)
            beats = beats_from_parse_result(pr)
            if not beats:
                report(f'F7 {name}', True, '无拍（跳过）')
                continue
            txt, _ = rt_file(f, tag=f'f7_{abs(hash(name)) % 9999}')
            ok = txt is not None
            report(f'F7 {name}', ok, '可往返' if ok else '失败')
        except Exception as e:
            report(f'F7 {name}', False, f'{type(e).__name__}: {e}')


def main() -> int:
    which = set(sys.argv[1:]) or {'all'}
    n_fuzz = int(next((x[2:] for x in which if x.startswith('n=')), 50))
    print('=' * 62)
    print('  极为变态的测试：差分 / 变异 / 脏数据 / 嵌套 / 规模 / 幂等 / 边界')
    print('=' * 62)
    steps = [
        ('f1', lambda: f1_fuzz(n_fuzz)),
        ('f2', lambda: f2_mutate(30)),
        ('f3', f3_dirty),
        ('f4', f4_hell),
        ('f5', f5_scale),
        ('f6', f6_idempotent),
        ('f7', f7_edges),
    ]
    try:
        for tag, fn in steps:
            if 'all' in which or tag in which:
                fn()
    except Exception:
        traceback.print_exc()
        return 2

    print('\n' + '=' * 62)
    print(f'  总计：{STATS["pass"]}/{STATS["total"]} 通过，'
          f'{STATS["fail"]} 失败')
    print('=' * 62)
    for name, detail in FAILS:
        print(f'  FAIL  {name}: {detail}')
    return 1 if STATS['fail'] else 0


if __name__ == '__main__':
    sys.exit(main())
