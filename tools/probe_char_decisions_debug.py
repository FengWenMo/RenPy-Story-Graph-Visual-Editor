"""诊断 char_decisions 的 7 条失败（临时探针，只读 + 临时目录）。"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.authoring.char_decisions import (                    # noqa: E402
    _evidence_for, _rank_candidates, apply_alias_plans,
    apply_new_character_plans, plan_alias, plan_new_character, suggestions)
from storygraph.authoring.txt_import import CharacterMap, convert_text  # noqa: E402

ROWS = {
    'jc': {'ID': 'jc', 'Name': '「蒋忱」'},
    'mm': {'ID': 'mm', 'Name': '妈妈'},
    'txf': {'ID': 'txf', 'Name': '汤X凤'},
    'aaa': {'ID': 'aaa', 'Name': '舍友A'},
}

print('=== A. 证据函数')
rows = dict(ROWS, mm={'ID': 'mm', 'Name': '妈妈', 'Aliases': 'Mom'})
print('  CASE Mom/MOM ->', _evidence_for('MOM', 'mm', rows['mm']))
print('  SEP  L_J_K/ljk ->', _evidence_for('L_J_K', 'ljk', {}))
print('  FUZZY MOM/mm ->', _evidence_for('MOM', 'mm', ROWS['mm']))

print('\n=== B. 候选排序（L_J_K 用带 ljk 的表）')
rows2 = dict(ROWS, ljk={'ID': 'ljk', 'Name': '「李佳凯」'})
conv = convert_text('L_J_K"我"', characters=CharacterMap(rows2))
print('  unmapped=', dict(conv.unmapped), 'ambiguous=', dict(conv.ambiguous))
sug = suggestions(conv, CharacterMap(rows2), rows2)
for s in sug:
    print(f'  [{s.token}] note={s.note!r}')
    for c in s.candidates:
        print(f'      {c.var} {c.evidence} {c.score:.2f} {c.reason}')
print('  rank 直接调用 ->', [(c.var, c.evidence)
                             for c in _rank_candidates('L_J_K', rows2,
                                                       CharacterMap(rows2), 5)])

print('\n=== C. plan_alias 的两个拒绝分支')
r3 = {'ljk': {'ID': 'ljk', 'Name': '「李佳凯」'}}
p = plan_alias(r3, 'LJK', 'ljk')
print('  LJK→ljk :', p.ok, '|', p.reason)
r4 = {'a': {'ID': 'a', 'Name': '甲'}, 'b': {'ID': 'b', 'Name': '乙'}}
p = plan_alias(r4, '乙', 'a')
print('  alias 乙→a :', p.ok, '|', p.reason)

print('\n=== D. plan_new_character')
p = plan_new_character(ROWS, 'MOM', var='mom_new', name='妈妈')
print('  MOM var=mom_new name=妈妈 :', p.ok, '|', p.reason)
p = plan_new_character(ROWS, 'MOM', var='mom2', name='妈妈')
print('  MOM var=mom2 name=妈妈 :', p.ok, '|', p.reason)

print('\n=== E. 落盘（写进临时表，再读回来）')
from openpyxl import Workbook, load_workbook                     # noqa: E402

with tempfile.TemporaryDirectory(prefix='cd_') as td:
    td = Path(td)
    wb = Workbook()
    ws = wb.active
    ws.append(['ID', 'Name', 'Color', 'ImageTag', 'Aliases', 'Notes'])
    ws.append(['mm', '妈妈', '', '', '', '手写的备注'])
    wb.save(td / 'characters.xlsx')
    rows5 = {'mm': {'ID': 'mm', 'Name': '妈妈', 'Notes': '手写的备注'}}
    plan = plan_alias(rows5, 'MM2', 'mm')
    print('  plan:', plan.ok, plan.after)
    res = apply_alias_plans(td, [plan])
    print('  apply_alias ->', res)
    ws2 = load_workbook(td / 'characters.xlsx').active
    print('  表内容:', [list(r) for r in ws2.iter_rows(values_only=True)])

with tempfile.TemporaryDirectory(prefix='cd2_') as td:
    td = Path(td)
    wb = Workbook()
    ws = wb.active
    ws.append(['ID', 'Name', 'Color', 'ImageTag', 'Aliases', 'Notes'])
    ws.append(['mm', '妈妈', '', '', 'MOM', ''])
    wb.save(td / 'characters.xlsx')
    plan = plan_new_character({'mm': {'ID': 'mm', 'Name': '妈妈'}}, 'TJF',
                              notes='由纯文字导入创建')
    print('\n  new plan:', plan.ok, plan.row)
    res = apply_new_character_plans(td, [plan])
    print('  apply_new ->', res)
    ws2 = load_workbook(td / 'characters.xlsx').active
    print('  表内容:', [list(r) for r in ws2.iter_rows(values_only=True)])
