"""
探针：**非 UTF-8 源文件**（GBK）在工具眼里是什么样、被写回来会变成什么。

假设（待验证）：`source_style.read_text` 用 `errors='replace'`，
所以 GBK 的中文被读成 U+FFFD（替换符），而不是"读不出来"。
如果没人拦，写盘时这些替换符就变成真的 UTF-8 `ef bf bd` ——
**用户文件里所有中文永久损坏**，而且退出码是 0。

用法：
    D:/Python/python.exe tools/probe_non_utf8.py
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable

BODY = '''\
label ch1:
    jc "第一句台词。"
    jc "第二句台词。"
    return
'''


def run(*args):
    p = subprocess.run([PY, str(ROOT / 'cli.py'), *args], cwd=str(ROOT),
                       capture_output=True, text=True,
                       encoding='utf-8', errors='replace')
    return p.returncode, (p.stdout or '') + (p.stderr or '')


def dump_sheet(xlsx: Path):
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx))
    print(f'  表文件：{xlsx.name}，{len(wb.worksheets)} 个页')
    for ws in wb.worksheets:
        print(f'  -- [{ws.title}] {ws.max_row} 行 x {ws.max_column} 列')
        for row in ws.iter_rows(max_row=8):
            vals = [c.value for c in row]
            if any(v is not None for v in vals):
                print('     ', [str(v)[:40] if v is not None else '' for v in vals])


def main() -> int:
    with tempfile.TemporaryDirectory(prefix='nonutf8_') as td:
        work = Path(td)
        src = work / 'ch1.rpy'
        src.write_bytes(BODY.encode('gbk'))
        before = src.read_bytes()
        print(f'原文件：{len(before)} 字节，GBK 编码')
        print(f'  GBK 原文可正常解码：{before.decode("gbk")!r}')

        xl = work / 'ch1.xlsx'
        rc, log = run('table', 'rpy2xlsx', str(src), '--out', str(xl))
        print(f'\n[1] rpy2xlsx 退出码 = {rc}')
        for ln in log.strip().splitlines()[-8:]:
            print('   ', ln)
        dump_sheet(xl)

        # 模拟用户在表里"改一句台词"：改掉表里那个被糟蹋过的格子
        import openpyxl
        wb = openpyxl.load_workbook(str(xl))
        hit = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and '\ufffd' in c.value:
                        c.value = c.value.replace('\ufffd', 'X')
                        hit += 1
        wb.save(str(xl))
        print(f'\n[2] 表里含替换符的格子：{hit} 个（已把它们改成 X，模拟用户编辑）')

        out = work / 'out.rpy'
        rc, log = run('table', 'xlsx2rpy', str(xl), '--out', str(out),
                      '--source', str(src), '--no-engine-lint')
        print(f'\n[3] xlsx2rpy 退出码 = {rc}')
        for ln in log.strip().splitlines()[-8:]:
            print('   ', ln)

        print('\n[4] 产物检查')
        for p in (src, out):
            if not p.exists():
                print(f'   {p.name}: 不存在')
                continue
            b = p.read_bytes()
            print(f'   {p.name}: {len(b)} 字节')
            print(f'      是 GBK 原文吗    : {b == before}')
            print(f'      含 U+FFFD 字节吗 : {b.count(b"\xef\xbf\xbd")} 处')
            try:
                print(f'      按 GBK 解        : {b.decode("gbk")!r}')
            except UnicodeDecodeError as e:
                print(f'      按 GBK 解        : 解不了（{e}）')
            print(f'      按 UTF-8 解      : {b.decode("utf-8", "replace")!r}')
        print('\n[5] 原文件还在吗（--out 另存时不该动原文件）')
        print(f'   源文件字节未变: {src.read_bytes() == before}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
