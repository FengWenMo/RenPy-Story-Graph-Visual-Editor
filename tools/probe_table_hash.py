#!/usr/bin/env python
"""
建表**内容指纹**：优化前后逐格比对（"性能改了但产物一个字节都不许变"）
========================================================================

做法：把真机拍导出成 xlsx，再**逐格**读回来算一个稳定 hash
（表名 / 行 / 列 / 值 / 单元格类型），并打出一行 `HASH=…`。
优化前后各跑一次，`HASH` 必须**一字不差**。

    python tools/probe_table_hash.py [项目根]
"""

from __future__ import annotations

import hashlib
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl                                            # noqa: E402

from storygraph.parsing.discover import discover_scripts   # noqa: E402
from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer  # noqa: E402
from storygraph.story.beat_build import beats_from_parse_result  # noqa: E402
from storygraph.story.excel_export import export_beats_to_excel  # noqa: E402


def main() -> int:
    root = Path(sys.argv[1] if len(sys.argv) > 1
                else r'D:\千连万理\RenpyTest\sr_general')
    beats = []
    for f in discover_scripts(root):
        pr, _ = IncrementalRenPyAnalyzer().analyze(files=[f], graph=None)
        beats.extend(beats_from_parse_result(pr))

    out = Path(tempfile.mkdtemp(prefix='probe_table_hash_')) / 't.xlsx'
    export_beats_to_excel(beats, out)

    h = hashlib.sha256()
    wb = openpyxl.load_workbook(out)
    try:
        for name in sorted(wb.sheetnames):
            ws = wb[name]
            h.update(f'#SHEET {name}\n'.encode())
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    h.update(f'{cell.row},{cell.column},'
                             f'{cell.data_type},{cell.value!r}\n'.encode())
    finally:
        wb.close()
    print(f'表 {len(wb.sheetnames) if False else "-"} · 拍 {len(beats)} · '
          f'体积 {out.stat().st_size} 字节')
    print(f'HASH={h.hexdigest()}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
