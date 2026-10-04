#!/usr/bin/env python
"""
复核：把真换行塞进 `SINGLE_LINE_BEAT_FIELDS` 的**每一列**
（不只是已经测过的那 10 列），看产物有没有断行 / 不可解析。

为什么单独跑这一遍：`tests/test_cell_control_chars.py` 的逐列用例只覆盖
Speaker / Scene / Show / At / With / Hide / BGM / Voice / Label / Condition
十列，而 `SINGLE_LINE_BEAT_FIELDS` 有 20 列。**分类清单是对的**
（`test_field_partition_is_explicit` 守着），但"清单里"不等于"测过" ——
万一哪天有人把某一列挪出清单，测试不会红。

只读。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
from storygraph.story.beat_build import beats_from_parse_result
from storygraph.story.beat_model import (BEAT_HEADERS, BEAT_FIELDS,
                                         SINGLE_LINE_BEAT_FIELDS)
from storygraph.story.beat_render import make_symbols, render_beats_to_rpy
from storygraph.story.excel_export import export_beats_to_excel
from storygraph.story.excel_import import import_beats_from_excel
from storygraph.assets.asset_registry import AssetRegistry

SRC = ('label a:\n'
       '    show ling happy\n'
       '    play music "m.ogg" fadein 1.0\n'
       '    ling "第一句"\n'
       '    ling "第二句"\n')

TOP_OK = ('label ', 'init', 'default', 'define', 'transform', 'screen',
          'image ', 'style ', 'pass', 'return', '#', 'menu', 'if ',
          'elif ', 'else', 'while ', 'for ', 'onlayer')


def loose(text: str) -> list:
    return [ln for ln in text.splitlines()
            if ln.strip() and not ln[:1].isspace()
            and not ln.strip().startswith(TOP_OK)]


def stmts(text: str):
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / 'o.rpy'
        f.write_text(text, encoding='utf-8')
        try:
            pr, _ = IncrementalRenPyAnalyzer().analyze(files=[f], graph=None)
            return len(pr.statements)
        except Exception as e:                        # noqa: BLE001
            return -1


def probe(column: str, value: str):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        f = tmp / 's.rpy'
        f.write_text(SRC, encoding='utf-8')
        pr, _ = IncrementalRenPyAnalyzer().analyze(files=[f], graph=None)
        xl = tmp / 's.xlsx'
        export_beats_to_excel(beats_from_parse_result(pr), xl)
        wb = load_workbook(xl)
        ws = wb['Beats']
        headers = [c.value for c in ws[1]]
        col = headers.index(BEAT_HEADERS[column]) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).value = value
        wb.save(xl)
        back = import_beats_from_excel(xl)
        reg = AssetRegistry()
        out = render_beats_to_rpy(back, registry=reg,
                                  symbols=make_symbols(reg))
    return out


def main() -> int:
    bad = []
    for column in BEAT_FIELDS:
        if column not in SINGLE_LINE_BEAT_FIELDS:
            continue
        # Label / Seq 这些有形态约束，用"它那列可能出现的坏值"
        value = {
            'Label': 'a\nb', 'Seq': '1\n2', 'Kind': 'say\nscene',
            'LineId': 'a.0001\nX', 'ContentHash': 'ab\ncd',
            'StageOrder': 'scene\nshow', 'OpOrder': 'stage\naudio',
            'CodeBlocks': '1-3\n5-7',
        }.get(column, 'x\ny')
        try:
            out = probe(column, value)
        except Exception as e:                        # noqa: BLE001
            bad.append((column, f'抛异常 {type(e).__name__}: {e}'))
            print(f'✗ {column:14} 抛异常 {type(e).__name__}: {e}')
            continue
        l = loose(out)
        n = stmts(out)
        flag = '✗' if (l or n < 0) else '·'
        if l or n < 0:
            bad.append((column, f'断行={l[:2]} 语句数={n}'))
        print(f'{flag} {column:14} 语句={n:>3}  断行={len(l)}')
    print('-' * 46)
    if bad:
        print(f'有 {len(bad)} 列没挡住：')
        for c, why in bad:
            print(f'  ✗ {c}: {why[:120]}')
        return 1
    print(f'清单里全部 {len(SINGLE_LINE_BEAT_FIELDS)} 列都挡住了真换行')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
