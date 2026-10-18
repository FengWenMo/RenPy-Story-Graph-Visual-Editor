"""
真实项目的协作**第 N 轮**实测（2026-10-18）。
============================================

用户报的硬伤是"第二次起就失效"。这种缺陷单次往返的测试**必然漏**，
所以判据必须在真库上把第 1、2、3 轮都跑一遍、逐轮打印数字。

    python tools/probe_collab_rounds.py [项目根]

默认项目根 = `D:\\千连万理\\renpytest\\sr_general`（54 个脚本 / 1790 拍）。
产物写在 `tools/out_collab_rounds/`（每次重建，不碰项目本身）。

每一轮都在量三件事：
  ① 每行 payload 多大（曾经 69,341 —— 那是"文件级数据混进行级"的证据）；
  ② `BasePayload` 有没有被截断（解析得开吗、多长、有没有走摘要）；
  ③ 两人改同一句的**不同字段**时，是自动合并还是报假冲突。
"""
from __future__ import annotations

import json
import statistics
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.collab import collab_store as cs          # noqa: E402
from storygraph.core.cell_limits import CELL_TEXT_LIMIT    # noqa: E402
from storygraph.story.beat_fields import BEAT_HEADERS      # noqa: E402

DEFAULT_PROJECT = Path(r'D:\千连万理\renpytest\sr_general')
OUT = ROOT / 'tools' / 'out_collab_rounds'


def build_db(project: Path) -> Path:
    """建库 + import-ir（与 `cli.py collab import-ir` 同一条路）。"""
    from storygraph.pipeline.build import (discover_scripts,
                                           load_asset_registry)
    from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
    from storygraph.story.beat_build import beats_from_parse_result
    from storygraph.story.excel_export import export_beats_to_excel

    OUT.mkdir(parents=True, exist_ok=True)
    db = OUT / 'collab.db'
    for f in OUT.glob('*.xlsx'):
        f.unlink()
    if db.exists():
        db.unlink()
    scripts = discover_scripts(project)
    pr, _ = IncrementalRenPyAnalyzer().analyze(files=scripts, graph=None)
    beats = beats_from_parse_result(pr)
    cs.init_db(db, users=['alice', 'bob'])
    seed = OUT / 'seed.xlsx'
    export_beats_to_excel(beats, seed, include_dropdowns=False)
    res = cs.push(db, seed, actor='alice', force=True)
    cs.upsert_assets(db, load_asset_registry(project), actor='alice')
    print(f'入库 {len(res["added"])} 行（来自 {len(scripts)} 个脚本）')
    return db


def payload_sizes(db):
    with cs.connect(db) as conn:
        rows = conn.execute('SELECT length(payload) AS n FROM beats '
                            'WHERE deleted=0').fetchall()
        blocks = cs.load_file_blocks(conn)
    n = [r['n'] for r in rows]
    text = next(iter(blocks.values()), '')
    print(f'  payload：{len(n)} 行 · 中位 {statistics.median(n):.0f} · '
          f'最大 {max(n)}')
    print(f'  文件级保真：{len(blocks)} 份 / {len(text):,} 字符 '
          f'（真源在 file_blocks 表，不在任何一行里）')
    return max(n)


def base_stats(path, tag):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[cs.COLLAB_SHEET]
        heads = [str(c.value or '') for c in ws[1]]
        i_base = heads.index(cs.BASE_FIELD) + 1
        i_lid = heads.index('LineId') + 1
        sizes, forms, bad = [], {'full': 0, 'hashed': 0}, []
        for r in range(2, ws.max_row + 1):
            lid = str(ws.cell(row=r, column=i_lid).value or '').strip()
            if not lid or lid == cs.PULLED_AT_KEY:
                continue
            raw = ws.cell(row=r, column=i_base).value
            text = '' if raw is None else str(raw)
            sizes.append(len(text))
            if not text:
                forms['none'] = forms.get('none', 0) + 1
                bad.append(lid)
                continue
            try:
                got = json.loads(text)
            except ValueError:
                bad.append(lid)
                continue
            forms['hashed' if cs.BASE_HASHES_KEY in got else 'full'] += 1
        print(f'  [{tag}] 基线 {len(sizes)} 行：中位 {statistics.median(sizes):.0f}'
              f' / 最大 {max(sizes)}（上限 {CELL_TEXT_LIMIT}）; '
              f'full {forms["full"]} · hashed {forms["hashed"]} · 坏 {len(bad)}')
        return bad
    finally:
        wb.close()


def _set_cell(path, lid, field, value):
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb['Beats']
        heads = [str(c.value or '') for c in ws[1]]
        i_id = heads.index(BEAT_HEADERS['LineId']) + 1
        col = heads.index(BEAT_HEADERS[field]) + 1
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=i_id).value or '').strip() == lid:
                ws.cell(row=r, column=col, value=value)
                break
        else:
            raise SystemExit(f'表里没有 {lid}')
        wb.save(path)
    finally:
        wb.close()


def pick_row(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb['Beats']
        heads = [str(c.value or '') for c in ws[1]]
        i_id = heads.index(BEAT_HEADERS['LineId']) + 1
        i_tx = heads.index(BEAT_HEADERS['Text']) + 1
        for r in range(2, ws.max_row + 1):
            lid = str(ws.cell(row=r, column=i_id).value or '').strip()
            tx = str(ws.cell(row=r, column=i_tx).value or '')
            if lid and 4 < len(tx) < 200:
                return lid, tx
        return '', ''
    finally:
        wb.close()


def main(argv=None):
    project = Path((argv or sys.argv[1:] or [DEFAULT_PROJECT])[0])
    print(f'项目：{project}')
    db = build_db(project)

    print('① 库刚 import-ir')
    payload_sizes(db)

    print('② 第一次 pull（alice）')
    a1 = OUT / 'alice1.xlsx'
    print('  ', cs.pull(db, a1, actor='alice'))
    base_stats(a1, 'alice1')

    lid, old = pick_row(a1)
    print(f'③ alice 改一句：{lid}（原台词 {old[:20]!r}）')
    _set_cell(a1, lid, 'Text', 'ALICE：' + old)
    r = cs.push(db, a1, actor='alice')
    print(f'   push → updated {len(r["updated"])} · conflicts {len(r["conflicts"])}')
    payload_sizes(db)

    print('④ 第二次 pull（bob）—— 用户报的断点就在这之后')
    b1 = OUT / 'bob1.xlsx'
    cs.pull(db, b1, actor='bob')
    bad = base_stats(b1, 'bob1')
    assert not bad, f'第二次 pull 的基线又坏了：{bad[:5]}'

    print('⑤ bob 改同一句的备注（不同字段 → 必须自动合并）')
    _set_cell(b1, lid, 'Comment', '## BOB 第二轮备注')
    r2 = cs.push(db, b1, actor='bob')
    print(f'   push → updated {r2["updated"]} · degraded {r2["degraded"]} · '
          f'conflicts {[(c["kind"], [f["field"] for f in c["fields"]]) for c in r2["conflicts"]]}')

    print('⑥ 第三次：两人各拉一份、各改一格')
    a2, b2 = OUT / 'alice2.xlsx', OUT / 'bob2.xlsx'
    cs.pull(db, a2, actor='alice')
    cs.pull(db, b2, actor='bob')
    base_stats(b2, 'bob2')
    _set_cell(a2, lid, 'Text', 'ALICE 第三轮')
    ra = cs.push(db, a2, actor='alice')
    _set_cell(b2, lid, 'Voice', 'voice/bob.ogg')
    rb = cs.push(db, b2, actor='bob')
    print(f'   alice → updated {ra["updated"][:3]} · conflicts {len(ra["conflicts"])}')
    print(f'   bob   → updated {rb["updated"][:3]} · degraded {rb["degraded"]} · '
          f'conflicts {len(rb["conflicts"])} · 合并字段 {rb["auto_merged_fields"]}')

    with cs.connect(db) as conn:
        row = conn.execute('SELECT payload FROM beats WHERE line_id=?',
                           (lid,)).fetchone()
    p = json.loads(row['payload'])
    print('⑦ 库里这三格：')
    for k in ('Text', 'Comment', 'Voice'):
        print(f'   {k} = {p.get(k)!r}')
    ok = (p.get('Comment') == '## BOB 第二轮备注'
          and p.get('Text') == 'ALICE 第三轮'
          and p.get('Voice') == 'voice/bob.ogg')
    print('结论：', '✅ 三轮都按字段合，没有退化' if ok else '❌ 还有字段被吃掉')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
