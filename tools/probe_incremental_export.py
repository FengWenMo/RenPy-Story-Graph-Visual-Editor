"""
探针：CLI 的**增量路径**（`--source`）在真实项目上是否保住"第三类"。

交付评审（2026-10-04）的边界②：

> CLI 的 `table xlsx2rpy` 只有 xlsx、没有源 .rpy → patcher 无源 →
> **永远走全量** → 只能用于**新建**；改已有章节必须用 UI 卡牌编辑器。

2026-10-05 拓展之后：`--source <这一章原来的 .rpy>` 让命令行与编辑器等价。
本探针在真实项目上走一遍 `.rpy → 表 → 改一句 → 表 → .rpy`，量三件事：

  1. 走没走增量（`--source` 给对了就该走）；
  2. 顶层语句（`default` / `init python` / `transform` …）还在不在；
  3. 与原文的差异行数（理想情况 = 生成的标记 + 改掉的那句）。

用法：
    D:/Python/python.exe tools/probe_incremental_export.py [章节.rpy]
"""
from __future__ import annotations

import difflib
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REAL = Path(r'D:\千连万理\RenPyTest\sr_general')
PY = sys.executable

#: 顶层（不进表）的语句 —— 全量重生成一行都留不下
OPAQUE_RE = re.compile(r'^(?:init(?:\s+\d+)?\s+python|default|define|transform'
                       r'|screen|style|init\b)', re.M)

#: 默认样本：**同时**有顶层语句（`default` / `define` / `init python` /
#: `screen`）和可改的台词 —— 少任何一样都演示不了边界②。
#: 想要更快的样本：`whackamole.rpy`（199 行，但只有一句旁白可改）。
DEFAULT_SAMPLE = REAL / 'game' / 'minigame' / 'ball_moving_game.rpy'


def run(*args):
    p = subprocess.run([PY, str(ROOT / 'cli.py'), *args], cwd=ROOT,
                       capture_output=True, text=True, encoding='utf-8',
                       errors='replace')
    return p.returncode, (p.stdout or '') + (p.stderr or '')


def main(argv) -> int:
    src = Path(argv[1]) if len(argv) > 1 else DEFAULT_SAMPLE
    if not src.exists():
        print('真实项目不在，跳过')
        return 0
    with tempfile.TemporaryDirectory(prefix='inc_probe_') as td:
        return _run(src, Path(td))


def _run(src: Path, work: Path) -> int:
    xl, out = work / 't.xlsx', work / 't.out.rpy'
    for f in (xl, out):
        if f.exists():
            f.unlink()

    original = src.read_text(encoding='utf-8')
    opaque = sorted({m.group(0).strip() for m in OPAQUE_RE.finditer(original)})
    print(f'源：{src.name}（{len(original.splitlines())} 行；'
          f'顶层语句 {len(opaque)} 种：{opaque[:5]}）')

    code, log = run('table', 'rpy2xlsx', str(src), '--out', str(xl))
    print(f'[1] rpy2xlsx exit={code}')
    if code:
        print(log[:600])
        return 1

    from openpyxl import load_workbook
    wb = load_workbook(xl)
    ws = wb['Beats']
    headers = [c.value for c in ws[1]]
    edited = 0
    for col_name in ('台词', '注释'):
        if edited:
            break
        col = headers.index(col_name) + 1
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v and str(v).strip() and not str(v).startswith('#'):
                ws.cell(row=row, column=col).value = '改过的PROBE'
                edited = 1
                break
    wb.save(xl)
    print(f'[2] 在表里改了 {edited} 处（模拟文案动手）')

    code, log = run('table', 'xlsx2rpy', str(xl), '--source', str(src),
                    '--out', str(out))
    print(f'[3] xlsx2rpy --source exit={code}')
    print('    ' + log.strip().replace('\n', '\n    ')[:400])
    if code or not out.exists():
        print('✗ 没有产物（被门禁拦下 —— 那正是"不静默降级"）')
        return 1

    text = out.read_text(encoding='utf-8')
    gone = [t for t in opaque if t not in text]
    print(f'[4] 顶层语句：{"全部保住 ✓" if not gone else "丢了 " + str(gone)}')
    diff = [l for l in difflib.unified_diff(original.splitlines(),
                                            text.splitlines(), lineterm='',
                                            n=0)
            if l[:1] in '+-' and l[:3] not in ('---', '+++')]
    print(f'[5] 与原文的差异行数：{len(diff)}（生成的标记占 2 行）')
    for l in diff[:8]:
        print(f'      {l[:100]}')
    ok = not gone
    print('结论：', '增量路径保住了第三类内容' if ok else '不达标')
    return 0 if ok else 1


if __name__ == '__main__':
    raise SystemExit(main(sys.argv))
