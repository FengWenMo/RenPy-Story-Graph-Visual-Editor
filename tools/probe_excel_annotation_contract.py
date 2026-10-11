"""按"提示放哪"那张表**逐行验收**（真机副本，不动真项目）。

用户 2026-10-11：以实际资源为准，按那张表做。表里四行：

  | 载体                     | 放不放 | 口径                                     |
  | 资源表 images.xlsx        | ❌ 一字不动 | 人工资源登记，机器不写语义判断、不加会漂移的列 |
  | 候选表 Candidates_*       | ❌ 不动     | 下拉必须全量（以实际资源为准）              |
  | 剧本表 ·「立绘」格批注     | ✅ 主载体   | 编剧在这一格工作；不占列、不改数据            |
  | 卡片 note                 | ✅ 逐拍     | 与 N3 同一条                              |

    python tools/probe_excel_annotation_contract.py
"""

import hashlib
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

REAL = Path(r'D:\千连万理\RenpyTest\sr_general')
if not REAL.is_dir():
    print('真机项目不在本机，跳过')
    raise SystemExit(0)

import openpyxl                                                      # noqa: E402

from storygraph.core import engine_facts as EF                       # noqa: E402
from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer  # noqa: E402
from storygraph.story.beat_build import beats_from_parse_result      # noqa: E402
from storygraph.story.excel_export import (BEAT_FIELDS,              # noqa: E402
                                           export_beats_to_excel)

work = Path(tempfile.mkdtemp(prefix='xlscontract_')) / 'sr_general'
shutil.copytree(REAL, work,
                ignore=shutil.ignore_patterns('.git', 'saves', 'build',
                                              '*.rpyc', 'log.txt', '*.backup',
                                              'cache'))
facts = EF.loads((work / 'storygraph.engine_facts.json').read_text(encoding='utf-8'))
print('事实可用：', EF.is_collectable(facts))


def digest(path):
    return hashlib.sha1(Path(path).read_bytes()).hexdigest()[:12]


def snapshot(dirpath):
    """目录里所有 xlsx 的指纹（用来证明"一字不动"）。"""
    out = {}
    for p in sorted(Path(dirpath).glob('*.xlsx')):
        out[p.name] = (digest(p), p.stat().st_size)
    return out


# ① 资源表 / 候选表：导出前后必须**逐字节一致**
tables = work / 'game' / 'asset_tables'
before = snapshot(tables)
print(f'\n=== ① 导出前：asset_tables 里 {len(before)} 个 xlsx')

chapter = work / 'game' / 'chapter' / 'chapter1_2.rpy'
pr, _ = IncrementalRenPyAnalyzer().analyze(files=[chapter], graph=None)
beats = beats_from_parse_result(pr)
print(f'   这一章 {len(beats)} 拍')

out = work / 'story_table_probe.xlsx'
res = export_beats_to_excel(beats, out, facts=facts)
n_annot = res.get('annotated_unknown') if isinstance(res, dict) else None
print('   导出结果：', {k: v for k, v in (res or {}).items()
                        if k in ('ok', 'rows', 'annotated_unknown',
                                 'annotated_spec_heads', 'path')} if isinstance(res, dict) else res)

after = snapshot(tables)
print('\n=== ② 资源表 / 候选表有没有被动过')
changed = [k for k in before if before[k] != after.get(k)]
added = [k for k in after if k not in before]
print('   变化：', changed or '无', '| 新增：', added or '无')
print('   判定：', '✅ 一字不动' if not changed and not added else '❌ 被改了')

# ③ 剧本表：批注挂了几格、挂在哪些格、内容对不对
wb = openpyxl.load_workbook(out)
ws = wb.active
col = BEAT_FIELDS.index('Show') + 1
commented, data_cells = [], []
for row in range(2, ws.max_row + 1):
    cell = ws.cell(row=row, column=col)
    if cell.value not in (None, ''):
        data_cells.append((row, str(cell.value)))
    if cell.comment is not None:
        commented.append((row, str(cell.value or ''), cell.comment.text))
print(f'\n=== ③ 剧本表「{BEAT_FIELDS[col - 1]}」列')
print(f'   有内容的格子 {len(data_cells)} 个；带批注的 {len(commented)} 个')
for row, value, text in commented[:4]:
    first = text.splitlines()[0]
    hit = [ln.strip() for ln in text.splitlines()
           if ln.strip() and not ln.startswith(('这一格', '可能是', '⚠', '引擎会'))]
    print(f'\n   第 {row} 行  格子值 = {value!r}')
    print(f'     批注首行：{first}')
    print(f'     批注提到：{hit[:4]}')

# 批注指的是不是**这一格**（值里真的含那个词）
# ⚠ 只解析**数据行**（以空格缩进的那种），别把提示行也当数据 ——
#    我第一版按"两个词就算一条"解析，把 `⚠ 只是提示：…不拦你。` 也当成了
#    `('⚠', '只是提示：…')`，于是自造出 4 处"不符"（探针自己错了）。
bad = []
for row, value, text in commented:
    for ln in text.splitlines():
        if not ln.startswith('  '):
            continue
        parts = ln.strip().split()
        if len(parts) == 2:
            tag, attr = parts
            if tag not in value or attr not in value:
                bad.append((row, value, tag, attr))
print(f'\n   批注是否都指向**它自己那一格**：'
      f'{"✅ 是" if not bad else "❌ 有 " + str(len(bad)) + " 处不符"}')
if bad:
    print('   反例：', bad[:3])

# ④ 不带事实 → 一个批注都不挂（不宣称）
out2 = work / 'story_table_probe_nofacts.xlsx'
export_beats_to_excel(beats, out2, facts=None)
ws2 = openpyxl.load_workbook(out2).active
n2 = sum(1 for row in range(2, ws2.max_row + 1)
         if ws2.cell(row=row, column=col).comment is not None)
print(f'\n=== ④ 不带事实导出 → 批注数 {n2}（应为 0：采不到就不宣称）')

# ⑤ 数据有没有被批注改动（两边的 Show 列逐格一致）
ws1 = wb.active
same = []
for row in range(2, max(ws1.max_row, ws2.max_row) + 1):
    a = ws1.cell(row=row, column=col).value
    b = ws2.cell(row=row, column=col).value
    same.append(a == b)
print(f'=== ⑤ 带事实 vs 不带事实：Show 列 {sum(same)}/{len(same)} 格一致'
      f' → {"✅ 批注不改数据" if all(same) else "❌ 数据被改了"}')
