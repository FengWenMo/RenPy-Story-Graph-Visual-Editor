"""
探针：**表的边界位置**（首行之前插一拍 / 改章节名 / 删到空）在产物里长什么样。

这三个都不是"内容对不对"的问题，而是"插进去的位置合不合法"：

  1. 在**第一张卡之前**插一拍 —— 锚点可能落到 `label` 行上，
     台词会被写到 `label ch01:` 之前（语法坏）；
  2. 把某一行的「章节」列改成别的 label —— 产物里会不会多出一个 label？
  3. 把表里所有行删掉 —— 产物是空文件、还是 `label ch01:` + `return`？

用法：
    D:/Python/python.exe tools/probe_table_edges.py
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable

BASE = '''\
label ch01:
    e "第一句"
    e "第二句"
    return
'''


def run(*args):
    p = subprocess.run([PY, str(ROOT / 'cli.py'), *args], cwd=str(ROOT),
                       capture_output=True, text=True,
                       encoding='utf-8', errors='replace')
    return p.returncode, (p.stdout or '') + (p.stderr or '')


def heads(ws):
    return {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}


def show(path: Path, title: str):
    print(f'  -- {title}（{path.stat().st_size} 字节）')
    for ln in path.read_text(encoding='utf-8').splitlines():
        print(f'     | {ln}')


def case_insert_first(work: Path):
    print('\n' + '=' * 72)
    print('场景 1：在**第一张卡之前**插一拍（锚点会落到 label 行上）')
    src = work / 'ins.rpy'
    src.write_text(BASE, encoding='utf-8')
    xl = work / 'ins.xlsx'
    rc, _ = run('table', 'rpy2xlsx', str(src), '--out', str(xl))
    assert rc == 0

    import openpyxl
    wb = openpyxl.load_workbook(str(xl))
    ws = wb['Beats']
    h = heads(ws)
    # 在标题行（第 1 行）之后、第一张卡之前插一行
    ws.insert_rows(2)
    ws.cell(2, h['章节']).value = 'ch01'
    ws.cell(2, h['类型']).value = 'say'
    ws.cell(2, h['说话人']).value = 'e'
    ws.cell(2, h['台词']).value = '【插在最前面的一句】'
    wb.save(str(xl))

    out = work / 'ins.out.rpy'
    rc, log = run('table', 'xlsx2rpy', str(xl), '--out', str(out),
                  '--source', str(src), '--no-engine-lint')
    print(f'  退出码 = {rc}')
    for ln in [x for x in log.splitlines() if x.strip()][-4:]:
        print(f'     {ln[:150]}')
    if out.exists():
        show(out, '产物')
        t = out.read_text(encoding='utf-8')
        bad = False
        for i, ln in enumerate(t.splitlines(), 1):
            if ln.strip().startswith('e "【插在最前面'):
                bad = i < next((k for k, x in enumerate(t.splitlines(), 1)
                                if x.startswith('label ')), 10 ** 9)
        print(f'  台词写在 label 之前了吗 : {bad}')


def case_rename_label(work: Path):
    print('\n' + '=' * 72)
    print('场景 2：把第 2 行的「章节」改成 ch99（表里只改了这一个格子）')
    src = work / 'ren.rpy'
    src.write_text(BASE, encoding='utf-8')
    xl = work / 'ren.xlsx'
    rc, _ = run('table', 'rpy2xlsx', str(src), '--out', str(xl))
    assert rc == 0

    import openpyxl
    wb = openpyxl.load_workbook(str(xl))
    ws = wb['Beats']
    h = heads(ws)
    changed = 0
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, h['台词']).value or '').strip() == '第二句':
            ws.cell(r, h['章节']).value = 'ch99'
            changed += 1
    wb.save(str(xl))
    print(f'  改到 {changed} 个格子')

    out = work / 'ren.out.rpy'
    rc, log = run('table', 'xlsx2rpy', str(xl), '--out', str(out),
                  '--source', str(src), '--no-engine-lint')
    print(f'  退出码 = {rc}')
    for ln in [x for x in log.splitlines() if x.strip()][-4:]:
        print(f'     {ln[:150]}')
    if out.exists():
        show(out, '产物')
    else:
        print('  产物没生成（门禁拦下）')


def case_empty_table(work: Path):
    print('\n' + '=' * 72)
    print('场景 3：把表里**所有行**删掉（等于"这一章不要了"）')
    src = work / 'empty.rpy'
    src.write_text(BASE, encoding='utf-8')
    xl = work / 'empty.xlsx'
    rc, _ = run('table', 'rpy2xlsx', str(src), '--out', str(xl))
    assert rc == 0

    import openpyxl
    wb = openpyxl.load_workbook(str(xl))
    ws = wb['Beats']
    ws.delete_rows(2, ws.max_row)
    wb.save(str(xl))

    out = work / 'empty.out.rpy'
    rc, log = run('table', 'xlsx2rpy', str(xl), '--out', str(out),
                  '--source', str(src), '--no-engine-lint')
    print(f'  退出码 = {rc}')
    for ln in [x for x in log.strip().splitlines()][-4:]:
        print(f'     {ln[:150]}')
    print(f'  源文件还在吗 : {src.exists()}（{src.stat().st_size} 字节）')
    if out.exists():
        show(out, '产物')


def main() -> int:
    with tempfile.TemporaryDirectory(prefix='edges_') as td:
        work = Path(td)
        for fn in (case_insert_first, case_rename_label, case_empty_table):
            try:
                fn(work)
            except Exception as e:              # noqa: BLE001 - 探针要跑完
                print(f'  !! {fn.__name__} 抛了：{type(e).__name__}: {e}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
