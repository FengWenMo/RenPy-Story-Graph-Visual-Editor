"""
B 阶段三的三条证据（退掉 `stamp_menu_bodies` 之后）：

  ① **表内容变干净**：选项体的 `Condition` 格子不再被补 `True`；
  ② **产物零变化**：那层 `if True:` 本来就等于"累积条件"、渲染时会被剥掉
     （所以两条路哈希与阶段二一致）；
  ③ **老表照样能读**：有 `True`、**没有** `Nest` 列的表仍然渲染正确 ——
     老推断要的正是那些 `True`。

只写临时目录。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

SRC = '''label a:
    e "开场"
    menu:
        "选项A":
            e "A 的内容"
        "选项B":
            e "B 的内容"
    e "菜单之后"
'''


def _beats(text: str, tmp: Path) -> list:
    from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
    from storygraph.story.beat_build import beats_from_parse_result
    f = tmp / 'ch.rpy'
    f.write_text(text, encoding='utf-8')
    pr, _ = IncrementalRenPyAnalyzer().analyze(files=[f], graph=None)
    return beats_from_parse_result(pr)


def main() -> int:
    import openpyxl

    from storygraph.story.beat_model import BEAT_HEADERS
    from storygraph.story.beat_render import render_beats_to_rpy
    from storygraph.story.excel_export import export_beats_to_excel
    from storygraph.story.excel_import import import_beats_from_excel

    with tempfile.TemporaryDirectory(prefix='b3_') as td:
        tmp = Path(td)
        beats = _beats(SRC, tmp)
        xl = tmp / 'new.xlsx'
        export_beats_to_excel(beats, xl)

        # ① 表里还有没有我们补的 True
        wb = openpyxl.load_workbook(xl)
        ws = wb['Beats']
        heads = [c.value for c in ws[1]]
        ci = heads.index(BEAT_HEADERS['Condition']) + 1
        conds = [ws.cell(row=r, column=ci).value
                 for r in range(2, ws.max_row + 1)]
        stamped = [c for c in conds if str(c or '').strip() == 'True']
        print(f'① 表里 Condition 列：{conds}')
        print(f'   我们自己补的 `True`：**{len(stamped)} 个**（阶段三是 0 个）')

        # ② 产物（过表）与阶段二应当逐字一致
        back = import_beats_from_excel(xl)
        print('② 过表产物：')
        for ln in render_beats_to_rpy(back).splitlines():
            if ln.strip() and not ln.startswith('#'):
                print(f'     {ln}')

        # ③ 模拟**老表**：删掉 Nest 列、把**选项体里的拍**条件写成 True。
        #    老 `stamp_menu_bodies` 只补"选项体里 + 条件为空"的那些 ——
        #    第一版探针把块外的拍也补了，那是**模拟错了**（不是代码错），
        #    结果 `菜单之后` 被吸进选项体。
        wb2 = openpyxl.load_workbook(xl)
        ws2 = wb2['Beats']
        heads2 = [c.value for c in ws2[1]]
        ni = heads2.index(BEAT_HEADERS['Nest']) + 1
        ci2 = heads2.index(BEAT_HEADERS['Condition']) + 1
        for r in range(2, ws2.max_row + 1):
            nest = str(ws2.cell(row=r, column=ni).value or '')
            in_body = '>' in nest          # `menu:n>X…` = 在某个选项体里
            if in_body and not str(
                    ws2.cell(row=r, column=ci2).value or '').strip():
                ws2.cell(row=r, column=ci2).value = 'True'
        ws2.delete_cols(ni)
        old_xl = tmp / 'old.xlsx'
        wb2.save(old_xl)
        old_back = import_beats_from_excel(old_xl)
        print(f'③ 老表（删掉 Nest、条件补 True）读回：'
              f'nest 全空 = {all(b.nest_path is None for b in old_back)}')
        print('   老表产物：')
        for ln in render_beats_to_rpy(old_back).splitlines():
            if ln.strip() and not ln.startswith('#'):
                print(f'     {ln}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
