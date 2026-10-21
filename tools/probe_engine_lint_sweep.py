"""
探针：把导出的产物**丢给真的 Ren'Py 引擎**读一遍（`lint`）。

工具自己的三层校验验的是"内部前后一致"；"引擎能不能解析这个产物"是另一件事
—— 上一轮审计（2026-10-18）的 P0-1~P0-6 全是"引擎会报、工具不会报"的错。

这一轮挑最可能写出**非法语法**的十个场景，每份产物都过 `lint`：

  1. 删 if 分支的唯一语句（`else` 要改写成 `if not (...)`）
  2. 删 elif 分支的唯一语句（链上要改写）
  3. 删菜单选项体的唯一语句（体换成 `pass`）
  4. 菜单只剩一个选项
  5. 说话人与正文同时改（槽位配对）
  6. 未定义角色（渲染器写 TODO + 降级旁白）
  7. 台词里含引号 / 换行 / 反斜杠
  8. 空章节（`label x:` + `return`）
  9. 两层嵌套 if 里删一拍
 10. `python:` 块里的语句

判据只有一条：**引擎报 parse error = P0**（工具说"已导出"、游戏起不来）。

用法：
    D:/Python/python.exe tools/probe_engine_lint_sweep.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.app import tables                      # noqa: E402
from storygraph.pipeline.engine_lint import _sdk, lint_text   # noqa: E402

SRC_SIMPLE = '''\
label ch01:
    e "第一句"
    e "第二句"
    e "第三句"
    return
'''

SRC_IF = '''\
label ch02:
    e "S1"
    if flag:
        e "S2 在分支里"
    else:
        e "S3 在另一支里"
    e "S4"
    return
'''

SRC_ELIF = '''\
label ch03:
    e "T1"
    if n == 1:
        e "T2 一"
    elif n == 2:
        e "T3 二"
    else:
        e "T4 其它"
    return
'''

SRC_MENU = '''\
label ch04:
    e "M1"
    menu:
        "选左":
            e "M2 左"
        "选右":
            e "M3 右"
    return
'''

SRC_SINGLE_MENU = '''\
label ch05:
    e "N1"
    menu:
        "只有一个":
            e "N2 唯一"
    return
'''

SRC_NEST = '''\
label ch06:
    e "P1"
    if a:
        if b:
            e "P2 双层"
        e "P3 外层"
    else:
        e "P4 另一支"
    return
'''

SRC_PY = '''\
label ch07:
    e "Q1"
    python:
        x = 1
        y = 2
    e "Q2"
    return
'''

SRC_EMPTY = '''\
label ch08:
    return
'''

SRC_UNDEF = '''\
label ch09:
    qing "我是未定义角色"
    e "我是定义过的"
    return
'''


def build(root: Path, name: str, body: str) -> Path:
    (root / 'game').mkdir(parents=True, exist_ok=True)
    fp = root / 'game' / name
    fp.write_text(body, encoding='utf-8')
    return fp


def sheet(table: Path):
    import openpyxl
    wb = openpyxl.load_workbook(str(table))
    ws = wb['Beats']
    heads = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    rows = []
    for r in range(2, ws.max_row + 1):
        kind = str(ws.cell(r, heads['类型']).value or '').strip()
        text = str(ws.cell(r, heads['台词']).value or '').strip()
        if kind or text:
            rows.append({'row': r, 'kind': kind, 'text': text})
    return wb, ws, heads, rows


def run_case(name: str, body: str, mutate, work: Path) -> dict:
    root = work / name
    src = build(root, 'ch.rpy', body)
    gen = tables.generate_table(src, project_root=root)
    if not gen.get('ok'):
        return {'name': name, 'verdict': 'SKIP', 'note': f'建表失败 {gen.get("error")}'}
    table = gen['table_path']
    wb, ws, heads, rows = sheet(table)
    note = mutate(ws, heads, rows)
    wb.save(str(table))
    res = tables.export_table(table, project_root=root, source_rpy=src)
    done = tables.apply_export(res, project_root=root)
    d = {'name': name, 'note': note or '', 'written': bool(done.get('written'))}
    if not d['written']:
        d['verdict'] = '拒写'
        d['note'] += '（' + str(res.get('blocking') or done.get('error'))[:60] + '）'
        return d
    text = src.read_text(encoding='utf-8')
    d['text'] = text
    rep = lint_text(text)
    d['ran'] = rep.ran
    if not rep.ran:
        d['verdict'] = 'SKIP'
        d['note'] += f'（lint 没跑：{rep.context_note[:40]}）'
        return d
    d['errors'] = rep.parse_errors
    d['verdict'] = 'P0' if rep.parse_errors else 'PASS'
    return d


# ---- 十种改动 ----
def _find(rows, text):
    for r in rows:
        if r['text'] == text:
            return r
    return None


def m_del_if(ws, h, rows):
    r = _find(rows, 'S2 在分支里')
    if r:
        ws.delete_rows(r['row'])
        return '删 if 分支唯一语句'
    return None


def m_del_elif(ws, h, rows):
    r = _find(rows, 'T3 二')
    if r:
        ws.delete_rows(r['row'])
        return '删 elif 分支唯一语句'
    return None


def m_del_menu_body(ws, h, rows):
    r = _find(rows, 'M2 左')
    if r:
        ws.delete_rows(r['row'])
        return '删菜单选项体唯一语句'
    return None


def m_menu_to_one(ws, h, rows):
    r = _find(rows, 'M3 右')
    if r:
        ws.delete_rows(r['row'])
        # 再把选项行也删掉（"选右"那一行）
        wb2, ws2, h2, rows2 = sheet_from(ws)
        for x in rows2:
            if x['text'] == '选右':
                ws.delete_rows(x['row'])
                break
        return '菜单删到一个选项'
    return None


def sheet_from(ws):
    heads = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    rows = []
    for r in range(2, ws.max_row + 1):
        kind = str(ws.cell(r, heads['类型']).value or '').strip()
        text = str(ws.cell(r, heads['台词']).value or '').strip()
        if kind or text:
            rows.append({'row': r, 'kind': kind, 'text': text})
    return None, ws, heads, rows


def m_speaker_and_text(ws, h, rows):
    r = _find(rows, '第一句')
    if r:
        ws.cell(r['row'], h['说话人']).value = 'ling'
        ws.cell(r['row'], h['台词']).value = '第一句（同改）'
        return '说话人+正文同改'
    return None


def m_undefined(ws, h, rows):
    """把说话人改成表里没有的角色（渲染器会写 TODO + 降级旁白）。"""
    r = _find(rows, '第二句')
    if r:
        ws.cell(r['row'], h['说话人']).value = 'qing'
        return '未定义角色'
    return None


def m_payload(ws, h, rows):
    r = _find(rows, '第三句')
    if r:
        ws.cell(r['row'], h['台词']).value = '他说："走吧。"\n第二行\\到这里'
        return '台词含引号/换行/反斜杠'
    return None


def m_empty_label(ws, h, rows):
    for r in rows:
        ws.delete_rows(r['row'])
    return '删光整章（label 只留 return）'


def m_nest(ws, h, rows):
    r = _find(rows, 'P2 双层')
    if r:
        ws.delete_rows(r['row'])
        return '两层嵌套里删一拍'
    return None


def m_python_block(ws, h, rows):
    r = _find(rows, 'Q2')
    if r:
        ws.cell(r['row'], h['台词']).value = 'Q2 改过'
        return 'python 块旁边改一句'
    return None


CASES = [
    ('删if分支', SRC_IF, m_del_if),
    ('删elif分支', SRC_ELIF, m_del_elif),
    ('删菜单选项体', SRC_MENU, m_del_menu_body),
    ('菜单剩一个选项', SRC_MENU, m_menu_to_one),
    ('说话人同改', SRC_SIMPLE, m_speaker_and_text),
    ('未定义角色', SRC_SIMPLE, m_undefined),
    ('台词含特殊字符', SRC_SIMPLE, m_payload),
    ('删光整章', SRC_EMPTY, m_empty_label),
    ('嵌套里删一拍', SRC_NEST, m_nest),
    ('python块旁改一句', SRC_PY, m_python_block),
    ('未定义角色源', SRC_UNDEF, lambda ws, h, rows: '原样导出'),
]


def main() -> int:
    sdk = _sdk()
    print(f'Ren\'Py SDK：{sdk}')
    if sdk is None:
        print('没有 SDK，跳过（这一套判据只在真机上有意义）')
        return 0
    results = []
    with tempfile.TemporaryDirectory(prefix='lint_') as td:
        work = Path(td)
        for name, body, fn in CASES:
            try:
                r = run_case(name, body, fn, work)
            except Exception as e:                  # noqa: BLE001
                r = {'name': name, 'verdict': f'抛异常 {type(e).__name__}',
                     'note': str(e)[:80]}
            results.append(r)
            print('=' * 76)
            print(f'[{name}] {r.get("note", "")}')
            print(f'  判定 = {r["verdict"]}')
            if r.get('text'):
                for ln in r['text'].splitlines():
                    print(f'     | {ln}')
            for e in (r.get('errors') or [])[:3]:
                print(f'  ✗ 引擎：{e}')
    print('\n' + '=' * 76)
    print(f'{"场景":16} {"判定":8} 说明')
    for r in results:
        print(f'{r["name"]:16} {r["verdict"]:8} {r.get("note", "")[:50]}')
    bad = [r for r in results if r['verdict'] == 'P0']
    print(f'\n★ 引擎读不了的产物：{len(bad)} / {len(results)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
