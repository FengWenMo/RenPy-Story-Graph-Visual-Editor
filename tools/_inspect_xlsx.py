import sys
import tempfile
from pathlib import Path

sys.path.insert(0, '.')
from storygraph.story.script_table import open_rpy_as_table

src = Path(r'D:\software\renpy-8.4.1-sdk\the_question\game\script.rpy')
tmp = Path(tempfile.mkdtemp())
x = tmp / 'q.xlsx'
open_rpy_as_table(str(src), str(x))
import openpyxl

wb = openpyxl.load_workbook(x)
print(wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print('---', name, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(3, ws.max_row),
                            values_only=True):
        print([str(c)[:26] if c is not None else '' for c in row][:10])
