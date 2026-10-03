#!/usr/bin/env python
"""定性：brutal_test 报出的 F1 丢失样本与 F3 脏数据未出现案例。"""
from __future__ import annotations

import random
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer  # noqa
from storygraph.story.beat_build import beats_from_parse_result        # noqa
from storygraph.story.excel_export import export_beats_to_excel        # noqa
from storygraph.story.excel_import import import_beats_from_excel      # noqa
from storygraph.pipeline.compile import compile_beats                  # noqa
import openpyxl  # noqa
from tools.brutal_test import gen_script, rt_file, recall_loss, A, TMP  # noqa


def inspect_f1(seed):
    src = gen_script(seed)
    f = TMP / f'seed_{seed}.rpy'
    f.write_text(src, encoding='utf-8')
    pr, _ = A.analyze(files=[f], graph=None)
    txt, out = rt_file(f, tag=f'seed_{seed}')
    lost = recall_loss(pr.statements, out)
    print(f'--- seed {seed}：丢 {sum(lost.values())} 条 ---')
    for k in list(lost):
        print(f'    丢: {str(k)[:110]}')
    # 原文里这条语句在哪
    for k in list(lost):
        kind, parts = k
        payload = dict(parts)
        img = payload.get('image')
        if img:
            hits = [i + 1 for i, l in enumerate(src.splitlines())
                    if img in l and l.strip().split()[0] in ('show', 'hide')]
            print(f'    原文出现行号: {hits}')
            for h in hits[:4]:
                print(f'      {h}: {src.splitlines()[h-1]}')
            # 产物里这类指令
            outs = [l for l in (txt or '').splitlines() if img in l]
            print(f'    产物里的 {img} 行: {outs}')
    print()


def main():
    print('=== F1 丢失样本定性 ===')
    for s in (30, 38, 53, 56):
        inspect_f1(s)

    print('=== F3 脏数据定性 ===')
    base = ('label a:\n'
            '    menu:\n'
            '        "甲":\n'
            '            jump x\n'
            '        "乙":\n'
            '            jump y\n'
            '    e "正常台词"\n')
    f = TMP / 'dirty.rpy'
    f.write_text(base, encoding='utf-8')
    pr, _ = A.analyze(files=[f], graph=None)
    beats = beats_from_parse_result(pr)
    xl = TMP / 'dirty.xlsx'
    export_beats_to_excel(beats, xl)
    wb0 = openpyxl.load_workbook(xl)
    hdr = [c.value for c in wb0['Beats'][1]]

    for label, dirty in [('公式前缀', '=1+1'),
                         ('反斜杠', 'C:\\path\\to\\file')]:
        wb = openpyxl.load_workbook(xl)
        ws = wb['Beats']
        for col in ('台词', '注释', '说话人'):
            if col in hdr:
                ws.cell(row=2, column=hdr.index(col) + 1, value=dirty)
        p = TMP / f'dirty_{label}.xlsx'
        wb.save(p)
        # 看回读到的值
        wb2 = openpyxl.load_workbook(p)
        got = wb2['Beats'].cell(row=2, column=hdr.index('台词') + 1).value
        res = compile_beats(import_beats_from_excel(p))
        hits = [l for l in res.rpy_text.splitlines()
                if '1+1' in l or 'path' in l]
        print(f'  {label}: 写入 {dirty!r}')
        print(f'    Excel 回读得到: {got!r}')
        print(f'    产物相关行: {hits}')
        if 'path' in dirty:
            # 反斜杠是否被转义？Ren'Py 里 \\ 表示一个 \
            esc = [l for l in res.rpy_text.splitlines() if 'path' in l]
            print(f'    转义检查（产物应为 C:\\\\path）: {esc}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
