import sys
import tempfile
from pathlib import Path

sys.path.insert(0, '.')
import openpyxl
from storygraph.story.script_table import open_rpy_as_table

src = Path(r'D:\software\renpy-8.4.1-sdk\the_question\game\script.rpy')
tmp = Path(tempfile.mkdtemp())
x = tmp / 'q.xlsx'
open_rpy_as_table(str(src), str(x))
wb = openpyxl.load_workbook(x)
ws = wb['Beats']
header = [c.value for c in ws[1]]
print('columns:', len(header))
print(header)
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] == 'rightaway.0030':
        for h, v in zip(header, row):
            if v not in (None, ''):
                print(f'  {h} = {v!r}'[:200])
