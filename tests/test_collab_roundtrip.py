"""
协作库 round-trip 的保真契约（调研 P0-1 / P1-1）
==============================================

P0-1 的实测形态：`collab_store._payload_from_beat()` 用
`getattr(beat, field.lower(), '')` 猜属性名 ——

    StageOrder → stageorder （实际 stage_order）→ 空
    With       → with       （实际 with_）       → 空
    SpecHead   → spechead   （实际 spec_head）   → 空
    …

于是**过一次协作库**，舞台顺序、代码块结构、菜单结构、自定义语句头、
循环/相机头、行尾注释、顶层 `init python` 全部静默丢失；`Seq` 还被 pop 掉，
pull 出来全是 0。而当时唯一的 round-trip 测试只断言了 LineId 和 Text。

这个文件守：
  · 每个 `FieldSpec.collab` 字段过 DB 都原样回来（用**渲染产物**做总判据，
    再用关键字段逐个钉住"不是碰巧渲染一样"）；
  · `Seq` / `label` 从数据库那两列恢复（不是 payload 里那个旧值）；
  · 没有「协作信息」表 = 没有基线 → **默认拒绝覆盖**（P1-1）；
  · 软删除的行重新出现在表里 → 复活（`deleted=0`），不是 `unchanged`；
  · 只改 Seq / Label（内容指纹不变）也要落库。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.story.beat_model import assign_identity          # noqa: E402

CHAPTER = '''init python:
    persistent.x = 1

label ch01:
    scene bg room
    show ling happy at left # tail
    e "hello" # say tail
    if flag:
        e "inside"
    menu:
        "go":
            e "chosen"
            jump ch02
    mycue "shandian" intensity=3:
        target = ling
    while x:
        $ a = 1
    camera background:
        perspective 1.0
    return
'''


@pytest.fixture
def beats(tmp_path):
    """解析一段"每类保真结构都有"的章节。"""
    from storygraph.parsing.renpy_parser import (ParseResultBuilder,
                                                 RenPyParser)
    from storygraph.story.beat_build import beats_from_parse_result

    src = tmp_path / 'ch01.rpy'
    src.write_text(CHAPTER, encoding='utf-8')
    pr = ParseResultBuilder().build(RenPyParser().parse(str(src)))
    out = beats_from_parse_result(pr)
    # 把那条自定义块**声明过**（T1）：spec_head 是可编辑的头行原文。
    for b in out:
        if b.opaque.startswith('mycue'):
            b.spec_name = 'mycue'
            b.spec_head = b.opaque.split('\n', 1)[0]
    assign_identity(out)
    return out


@pytest.fixture
def seeded(tmp_path, beats):
    """建库 → 首次入库（force：没有基线是正常的首次导入）→ pull 出前端表。"""
    import storygraph.collab.collab_store as cs
    from storygraph.story.excel_export import export_beats_to_excel

    db = tmp_path / 'collab.db'
    cs.init_db(db, users=['alice', 'bob'])
    seed = tmp_path / '_seed.xlsx'
    export_beats_to_excel(beats, seed)
    cs.push(db, seed, actor='alice', force=True)
    sheet = tmp_path / 'frontend.xlsx'
    cs.pull(db, sheet, actor='alice')
    return db, sheet


def _render(beats) -> str:
    from storygraph.story.beat_render import render_beats_to_rpy
    return render_beats_to_rpy(beats)


class TestPayloadKeepsFidelity:

    def test_the_whole_chapter_survives_a_db_roundtrip(self, seeded, beats):
        """总判据：过一遍 DB 后的**渲染产物**必须与原文一致。"""
        _db, sheet = seeded
        from storygraph.story.excel_import import import_beats_from_excel

        back = import_beats_from_excel(sheet)
        assert _render(back) == _render(beats)

    def test_each_fidelity_field_is_still_there(self, seeded, beats):
        """
        逐个钉住曾被 `field.lower()` 弄空的字段 —— 渲染相等还不够
        （例如 Stage 文本可能掩盖 op_order 丢失，反过来也一样）。
        """
        _db, sheet = seeded
        from storygraph.story.excel_import import import_beats_from_excel

        back = import_beats_from_excel(sheet)
        got = {b.seq: b for b in back}
        want = {b.seq: b for b in beats}

        # 顶层块：只在第一章拍上
        assert got[1].file_blocks == want[1].file_blocks
        # 舞台 / 操作顺序
        assert got[1].stage_order == want[1].stage_order
        assert got[1].op_order == want[1].op_order
        # 行尾注释
        assert got[1].inline_comment == want[1].inline_comment
        # 结构位（if / menu 选项体）
        for seq, beat in want.items():
            if beat.nest_path is not None:
                assert got[seq].nest_path == beat.nest_path, seq
        # 自定义语句头 / 循环头 / 相机头
        assert got[7].opaque == want[7].opaque
        assert got[7].spec_name == 'mycue'
        assert got[7].spec_head == want[7].spec_head
        assert got[8].loop_header == want[8].loop_header
        assert got[10].camera_header == want[10].camera_header
        # 代码块结构
        assert got[9].code_blocks == want[9].code_blocks
        assert got[9].op_order == want[9].op_order

    def test_seq_is_restored_from_the_db_column(self, seeded):
        """`Seq` 从前被 pop 掉，pull 出来全是 0 —— 现在必须原样。"""
        _db, sheet = seeded
        from storygraph.story.excel_import import import_beats_from_excel

        back = import_beats_from_excel(sheet)
        assert [b.seq for b in back] == sorted(b.seq for b in back)
        assert all(b.seq > 0 for b in back)

    def test_export_ir_from_db_keeps_fidelity(self, seeded, beats):
        """CI / 构建走的是 `export_ir_from_db`（不经 Excel），它也得保真。"""
        import storygraph.collab.collab_store as cs

        db, _sheet = seeded
        from_db, _assets = cs.export_ir_from_db(db)
        assert _render(from_db) == _render(beats)


class TestPushNeedsABaseline:

    def test_a_plain_export_sheet_cannot_overwrite(self, seeded, tmp_path):
        """
        P1-1 的原始复现：普通导出表没有「协作信息」表 →
        `base_rev = 0` → 从前永远不冲突，直接盲覆盖。

        现在必须报冲突（reason=no_baseline），并且**库里的内容一个字不变**。
        """
        import storygraph.collab.collab_store as cs
        from storygraph.story.excel_export import export_beats_to_excel

        db, _sheet = seeded
        db_beats, _assets = cs.export_ir_from_db(db)
        db_beats[1].text = 'stale（旧表想覆盖）'
        plain = tmp_path / 'plain.xlsx'
        export_beats_to_excel(db_beats, plain)      # 注意：没有协作信息表

        res = cs.push(db, plain, actor='bob')
        assert res['conflicts'], res
        assert res['conflicts'][0]['reason'] == 'no_baseline'
        assert not res['updated']
        fresh, _ = cs.export_ir_from_db(db)
        assert fresh[1].text != 'stale（旧表想覆盖）'

    def test_force_is_the_explicit_escape_hatch(self, seeded, tmp_path):
        """确实要整表覆盖 → 必须显式 `--force`（与从前一样）。"""
        import storygraph.collab.collab_store as cs
        from storygraph.story.excel_export import export_beats_to_excel

        db, _sheet = seeded
        db_beats, _ = cs.export_ir_from_db(db)
        db_beats[1].text = '整表覆盖（明确要求）'
        plain = tmp_path / 'plain.xlsx'
        export_beats_to_excel(db_beats, plain)
        res = cs.push(db, plain, actor='bob', force=True)
        assert res['updated']
        fresh, _ = cs.export_ir_from_db(db)
        assert fresh[1].text == '整表覆盖（明确要求）'


class TestSoftDeleteAndMetaChanges:

    def test_a_deleted_line_comes_back_when_it_reappears(self, seeded):
        """
        软删除的行又被加回表里 → **复活**（`deleted=0`）。

        从前 UPDATE 不把 `deleted` 置回 0，push 还可能返回 `unchanged`，
        下次 pull 它又消失了 —— 用户看到的是"加回来也不见"。
        """
        import openpyxl
        import storygraph.collab.collab_store as cs

        db, sheet = seeded
        # 删掉第 2 行
        wb = openpyxl.load_workbook(sheet)
        wb['Beats'].delete_rows(3)
        wb.save(sheet)
        first = cs.push(db, sheet, actor='alice')
        assert not first['conflicts'], first
        with cs.connect(db) as conn:
            row = conn.execute(
                "SELECT deleted FROM beats WHERE line_id='ch01.0002'").fetchone()
        assert row['deleted'] == 1

        # 再把原来那份表推回去（这一行又出现了）
        pulled = cs.pull(db, sheet, actor='alice')  # 先恢复一份含该行的前端表
        assert pulled['rows'] >= 1
        # pull 只导出 deleted=0 的行 —— 所以要用最初的副本重建
        # （模拟用户手上还留着删除前那份 Excel）
        # 这里直接检验：表里有这一行时 push 会复活。
        wb = openpyxl.load_workbook(sheet)
        ws = wb['Beats']
        from storygraph.story.script_table import BEAT_FIELDS
        row = [''] * len(BEAT_FIELDS)
        row[BEAT_FIELDS.index('Seq')] = 2
        row[BEAT_FIELDS.index('Label')] = 'ch01'
        row[BEAT_FIELDS.index('LineId')] = 'ch01.0002'
        row[BEAT_FIELDS.index('Kind')] = 'say'
        row[BEAT_FIELDS.index('Speaker')] = 'e'
        row[BEAT_FIELDS.index('Text')] = '想回来的一句'
        ws.append(row)
        wb.save(sheet)
        res = cs.push(db, sheet, actor='alice')
        assert 'ch01.0002' in res['revived'], res
        with cs.connect(db) as conn:
            row = conn.execute(
                "SELECT deleted, payload FROM beats WHERE line_id='ch01.0002'"
            ).fetchone()
        assert row['deleted'] == 0

    def test_seq_only_change_is_pushed(self, seeded):
        """
        `Seq` 不在内容指纹里（否则插一行会让后面所有行"内容变了"），
        所以 push 必须**单独**比 Seq/Label —— 否则只改序号永远落不了库。
        """
        import openpyxl
        import storygraph.collab.collab_store as cs
        from storygraph.story.script_table import BEAT_FIELDS

        db, sheet = seeded
        wb = openpyxl.load_workbook(sheet)
        ws = wb['Beats']
        # 第 2 行改成 Seq=9（内容一个字不动）
        ws.cell(row=3, column=BEAT_FIELDS.index('Seq') + 1, value=9)
        wb.save(sheet)
        res = cs.push(db, sheet, actor='alice')
        assert 'ch01.0002' in res['updated'], res
        with cs.connect(db) as conn:
            row = conn.execute(
                "SELECT seq FROM beats WHERE line_id='ch01.0002'").fetchone()
        assert row['seq'] == 9
