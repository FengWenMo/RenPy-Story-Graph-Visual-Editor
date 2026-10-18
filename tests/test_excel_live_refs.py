"""
Excel **活引用**（2026-10-18，审计 P1-6）
========================================

审计原话："下拉是**候选**，不是**引用**：改资源表，Beats 里已写的值不会
跟着变（0 个公式）"。

这一份钉住"引用"这件事本身，判据分三层（都不钉实现细节）：

  ① **有公式**：`资源引用（活）` 表里，数据行除了行号**每一格都是公式**；
  ② **引用的是真源**：说话人那一格是 `VLOOKUP(…Candidates_角色!…)`
     —— 改资源表的"说明"列，这一格就会跟着变；
  ③ **反向计数也活**：每个资源的「被引用次数」是 `COUNTIF(Beats!…)`，
     改 Beats 立即重算（删资源前先看它）。

⚠ 同时守一条**不许退化**的边界：Beats 主表**不许多出列** —— 引用是旁路
视图，不是把主表变成公式表（公式格一旦被手改就是"覆盖公式"，连接没了）。
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.story.beat_model import Beat                      # noqa: E402
from storygraph.story.excel_export import export_beats_to_excel   # noqa: E402


def _registry():
    from storygraph.assets.asset_registry import AssetRegistry
    reg = AssetRegistry()
    reg.characters = {'lj': {'Name': '「蒋忱」', 'ImageTag': 'lj'}}
    reg.images = {'lj happy': {'Path': 'images/lj_happy.png'}}
    reg.music = {'bgm_day': {'Path': 'audio/bgm/day.ogg'}}
    reg.sounds = {}
    reg.voices = {}
    reg.transforms = {'leftish': {'Params': 'xpos 0.1'}}
    return reg


def _beats():
    return [
        Beat(seq=1, label='ch01', line_id='L1', kind='say', speaker='lj',
             text='你好'),
        Beat(seq=2, label='ch01', line_id='L2', kind='say', speaker='lj',
             text='再见', show='lj happy', bgm='bgm_day'),
    ]


@pytest.fixture
def book(tmp_path):
    out = tmp_path / 'live.xlsx'
    export_beats_to_excel(_beats(), out, registry=_registry())
    return openpyxl.load_workbook(out), out


def test_live_sheet_exists_and_is_formula_driven(book):
    wb, _ = book
    assert '资源引用（活）' in wb.sheetnames
    ws = wb['资源引用（活）']
    # 找出数据行：第 3 行起、第 2 列（章节）引用 Beats 的那些行
    rows = [r for r in range(3, ws.max_row + 1)
            if str(ws.cell(row=r, column=2).value or '').startswith('=')
            and 'Beats!' in str(ws.cell(row=r, column=2).value)]
    assert len(rows) == 2, f'两拍应当有两行，实际 {rows}'
    for r in rows:
        for c in range(2, 16):
            v = ws.cell(row=r, column=c).value
            assert isinstance(v, str) and v.startswith('='), (
                f'第 {r} 行第 {c} 列不是公式：{v!r}（引用必须是活的）')


def test_speaker_reference_points_at_the_resource_sheet(book):
    wb, _ = book
    ws = wb['资源引用（活）']
    # 第 4 列 = 说话人（资源表）：必须是 VLOOKUP 到 Candidates_角色
    f = ws.cell(row=3, column=4).value
    assert f.startswith('=') and 'VLOOKUP' in f
    assert 'Candidates_角色' in f, f'说话人没有引用角色表：{f}'
    # 第 2 列 = 章节：直接引用 Beats 的那一格（改 Beats 跟着变）
    assert 'Beats!' in ws.cell(row=3, column=2).value


def test_usage_count_is_live_countif(book):
    wb, _ = book
    ws = wb['资源引用（活）']
    counts = [ws.cell(row=r, column=4).value
              for r in range(3, ws.max_row + 1)
              if isinstance(ws.cell(row=r, column=4).value, str)
              and 'COUNTIF' in str(ws.cell(row=r, column=4).value)]
    assert counts, '「被引用次数」一列必须是 COUNTIF 公式（不是死数字）'
    assert any('Beats!' in c for c in counts)
    # 说话人的计数是整格精确匹配（一个格一个 ID）
    assert any(':E' in c or '!$E' in c for c in counts), counts


def test_beats_sheet_keeps_its_column_contract(book):
    """主表**不许多列**：引用是旁路视图，不许把 Beats 变成公式表。"""
    from storygraph.story.beat_fields import BEAT_FIELDS
    wb, _ = book
    ws = wb['Beats']
    assert ws.max_column == len(BEAT_FIELDS), (
        f'Beats 列数变了：{ws.max_column} != {len(BEAT_FIELDS)}')
    # 主表里一格公式都不许有（否则用户一编辑就"覆盖公式"）
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            assert not (isinstance(cell.value, str)
                        and cell.value.startswith('=')), (
                f'Beats 里出现了公式格 {cell.coordinate}：{cell.value!r}')


def test_no_live_sheet_without_registry(tmp_path):
    """没有资源表就没有可引用的真源 —— 这张表不该凭空出现。"""
    out = tmp_path / 'plain.xlsx'
    export_beats_to_excel(_beats(), out, registry=None)
    wb = openpyxl.load_workbook(out)
    assert '资源引用（活）' not in wb.sheetnames
