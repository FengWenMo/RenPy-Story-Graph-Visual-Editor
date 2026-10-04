"""
B 阶段三：退掉 `stamp_menu_bodies` 的 `True` 打标 —— 见不变量 77。

那层打标当年必须存在，因为**表里没有"这拍在菜单块里"这一位**，只能拿
`Condition` 列记 1 bit（`True` = 在某个选项体里 / 空 = 块外）。
块路径（`Nest` 列）随表走、渲染器也按它建树之后，这一位就多余了；
留着还有两个坏处：表里多出一堆策划没写过的 `True`，产物里多出一层
`if True:`（语义等价但难读，还让"过表"与"直接渲染"两份产物对不上）。

三条要钉住的：
  ① 表里不再出现我们补的 `True`；
  ② **产物零变化**（那层 `if True:` 本来就等于"累积条件"、渲染时会剥掉）；
  ③ **老表照样能读**（有 `True`、没有 `Nest`）—— 老推断要的正是那些 `True`。
"""

from __future__ import annotations

import inspect
from pathlib import Path

import openpyxl

from storygraph.story import excel_export
from storygraph.story.beat_model import BEAT_HEADERS
from storygraph.story.beat_render import render_beats_to_rpy
from storygraph.story.generation_guard import structure_loss

from tests.test_nest_phase2 import (_beats, _via_table, MENU_THEN_IF)

MENU = '''label a:
    e "开场"
    menu:
        "选项A":
            e "A 的内容"
        "选项B":
            e "B 的内容"
    e "菜单之后"
'''


class TestStampingIsRetired:
    def test_the_exporter_does_not_call_it_anymore(self):
        src = inspect.getsource(excel_export)
        assert 'stamp_menu_bodies' not in src, \
            '打标被挂回来了 —— 新表靠 `Nest` 列说结构，这一位是多余的'

    def test_no_true_conditions_are_written(self, tmp_path):
        xl = tmp_path / 'ch.xlsx'
        excel_export.export_beats_to_excel(_beats(MENU, tmp_path), xl)
        wb = openpyxl.load_workbook(xl)
        ws = wb['Beats']
        heads = [c.value for c in ws[1]]
        ci = heads.index(BEAT_HEADERS['Condition']) + 1
        conds = [ws.cell(row=r, column=ci).value
                 for r in range(2, ws.max_row + 1)]
        assert all(not str(c or '').strip() for c in conds), (
            f'表里被补了条件（策划没写过这些）：{conds}')

    def test_no_true_wrapper_in_the_product(self, tmp_path):
        """过表产物里不该再有 `if True:` 那层壳。"""
        rpy = render_beats_to_rpy(_via_table(_beats(MENU_THEN_IF, tmp_path),
                                             tmp_path))
        assert 'if True:' not in rpy, rpy


class TestTheTwoPathsNowAgree:
    """
    阶段二 + 三的合起来的效果：**过表产物与直接渲染一致**。

    以前"过表"那条会多出 `if True:` 那层（打标留下的），两份产物对不上；
    现在结构由 `Nest` 说、条件不再被我们改写，两条路应当逐字一样。
    """

    def test_table_path_equals_source_path(self, tmp_path):
        direct = render_beats_to_rpy(_beats(MENU_THEN_IF, tmp_path))
        through = render_beats_to_rpy(
            _via_table(_beats(MENU_THEN_IF, tmp_path), tmp_path))
        assert through == direct, f'过表与直接渲染对不上：\n{through}'

    def test_the_structure_gate_is_quiet_now(self, tmp_path):
        """
        A 档那条门禁（源 vs 产物的菜单结构）在过表产物上**不该再响** ——
        结构已经保真了。（它现在只在"有人手改了结构"时才该报。）
        """
        src = MENU_THEN_IF
        through = render_beats_to_rpy(
            _via_table(_beats(src, tmp_path), tmp_path))
        assert structure_loss(src, through) == [], through

    def test_two_roundtrips_converge(self, tmp_path):
        """连写两次表：产物必须收敛（打标当年也兼着"反复往返收敛"这件事）。"""
        beats = _beats(MENU_THEN_IF, tmp_path)
        first = _via_table(beats, tmp_path)
        r1 = render_beats_to_rpy(first)
        second = _via_table(first, tmp_path)
        r2 = render_beats_to_rpy(second)
        assert r2 == r1, f'第二次往返产物变了：\n{r1}\n---\n{r2}'


class TestOldTablesStillWork:
    """
    **老表**（有 `True` 打标、没有 `Nest` 列）必须继续渲染正确 ——
    存量文件不能因为这一版而坏掉。
    """

    def _make_old_table(self, tmp_path) -> Path:
        xl = tmp_path / 'new.xlsx'
        excel_export.export_beats_to_excel(_beats(MENU, tmp_path), xl)
        wb = openpyxl.load_workbook(xl)
        ws = wb['Beats']
        heads = [c.value for c in ws[1]]
        ni = heads.index(BEAT_HEADERS['Nest']) + 1
        ci = heads.index(BEAT_HEADERS['Condition']) + 1
        for r in range(2, ws.max_row + 1):
            nest = str(ws.cell(row=r, column=ni).value or '')
            # 老口径：只补"选项体里 + 条件为空"的拍
            if '>' in nest and not str(
                    ws.cell(row=r, column=ci).value or '').strip():
                ws.cell(row=r, column=ci).value = 'True'
        ws.delete_cols(ni)                       # 老表没有这一列
        old = tmp_path / 'old.xlsx'
        wb.save(old)
        return old

    def test_it_renders_correctly_without_the_nest_column(self, tmp_path):
        from storygraph.story.excel_import import import_beats_from_excel
        old = self._make_old_table(tmp_path)
        back = import_beats_from_excel(old)
        assert all(b.nest_path is None for b in back), '前提：老表没有结构位'
        rpy = render_beats_to_rpy(back)
        assert rpy.count('menu:') == 1, rpy
        for token in ('"选项A":', '"选项B":', 'A 的内容', 'B 的内容'):
            assert token in rpy, rpy
        body = next(ln for ln in rpy.splitlines() if 'A 的内容' in ln)
        assert body.startswith(' ' * 12), rpy          # 在选项体里
        tail = next(ln for ln in rpy.splitlines() if '菜单之后' in ln)
        assert tail.startswith(' ' * 4), rpy           # 在菜单外
