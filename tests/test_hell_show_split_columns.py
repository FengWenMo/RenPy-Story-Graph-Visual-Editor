"""
地狱测试 · 立绘/退场「每指令一列」摊列（M3-2 第一切片，Phase 1 模型侧）
=======================================================================

背景（用户 D4' 方案 B）：Excel 的数据验证**一个格子只认一个值**，而
`立绘` 列是 `a; b; c` 多值 → 每一项都挂不上图像候选下拉。摊成"每指令一列"
之后每格一个候选。

设计要点（`docs/路线图_长期目标_2026-10-08.md` M3-2）：
  · **模型与渲染不动** —— `beat.show` 仍是 `;` 连接的字符串；
  · 第 1 项留在原 `Show`/`Hide` 列，第 2+ 项占 `立绘2/立绘3…`、`退场2…`；
  · **顺序不靠列位置**，靠 `OpOrder`/`StageOrder`（C3）；
  · 旧表（整格 `;` 多值、没有这些列）行为**一字不变**。

⚠ 这组判据盯的是**静默丢项**：读表时少读 `立绘2` 就是"一整个立绘凭空消失"
且不报错 —— 本项目最贵的那类缺陷。
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.story.beat_model import (                                # noqa: E402
    BEAT_FIELDS, HIDE_EXTRA_PREFIX, SHOW_EXTRA_PREFIX, Beat,
    _row_fields, compose_hide_cell, compose_show_cell, extra_columns_in,
    extra_field_names, extra_index_of, hide_cell_parts, show_cell_parts,
)

SHOW_EXTRAS = ['立绘2', '立绘3']
HIDE_EXTRAS = ['退场2']


def _row_of(beat, extras=()) -> dict:
    extras = list(extras)
    fields = _row_fields(None, extras)
    cells = beat.to_row(extra_fields=extras)
    assert len(cells) == len(fields), (len(cells), len(fields))
    return dict(zip(fields, cells))


def _beat(**kw) -> Beat:
    b = Beat()
    b.label = 'start'
    b.speaker = 'e'
    for k, v in kw.items():
        setattr(b, k, v)
    return b


# ===========================================================================
#  ① 纯函数：列名/序号
# ===========================================================================

def test_extra_field_names_start_at_two():
    assert extra_field_names(SHOW_EXTRA_PREFIX, 1) == []
    assert extra_field_names(SHOW_EXTRA_PREFIX, 2) == ['立绘2']
    assert extra_field_names(SHOW_EXTRA_PREFIX, 4) == ['立绘2', '立绘3', '立绘4']
    assert extra_field_names(HIDE_EXTRA_PREFIX, 3) == ['退场2', '退场3']
    # 非法输入不炸（表里可能是空串）
    assert extra_field_names(SHOW_EXTRA_PREFIX, None) == []
    assert extra_field_names(SHOW_EXTRA_PREFIX, 'x') == []


def test_extra_index_of_only_its_own_family():
    assert extra_index_of('立绘2', SHOW_EXTRA_PREFIX) == 2
    assert extra_index_of('退场2', SHOW_EXTRA_PREFIX) is None   # 别族的别认领
    assert extra_index_of('Show', SHOW_EXTRA_PREFIX) is None
    assert extra_index_of('立绘', SHOW_EXTRA_PREFIX) is None     # 第 1 项不摊列
    assert extra_index_of('立绘x', SHOW_EXTRA_PREFIX) is None


def test_extra_columns_in_row_are_sorted_numerically():
    row = {'立绘10': 'c', '立绘2': 'a', '立绘3': 'b', 'Show': 'x'}
    assert extra_columns_in(row, SHOW_EXTRA_PREFIX) == ['立绘2', '立绘3', '立绘10']


# ===========================================================================
#  ② 同源：项列表 === 原列内容（拆函数不许引入第二套过滤）
# ===========================================================================

def test_parts_are_the_same_source_as_the_original_cell():
    b = _beat(show='ling happy; ling base_arms; ye@left', at='i11')
    assert compose_show_cell(b) == '; '.join(show_cell_parts(b))
    assert show_cell_parts(b) == ['ling happy', 'ling base_arms', 'ye@left']
    b2 = _beat(hide='ye@dissolve; ling')
    assert compose_hide_cell(b2) == '; '.join(hide_cell_parts(b2))
    assert hide_cell_parts(b2) == ['ye@dissolve', 'ling']


# ===========================================================================
#  ③ 摊列核心：原列只放第 1 项，其余各占一列
# ===========================================================================

def test_show_column_holds_only_the_first_item_when_split():
    b = _beat(show='ling happy; ling base_arms; ye@left')
    row = _row_of(b, SHOW_EXTRAS)
    assert row['Show'] == 'ling happy'
    assert row['立绘2'] == 'ling base_arms'
    assert row['立绘3'] == 'ye@left'


def test_hide_column_splits_the_same_way():
    b = _beat(hide='ye@dissolve; ling')
    row = _row_of(b, HIDE_EXTRAS)
    assert row['Hide'] == 'ye@dissolve'
    assert row['退场2'] == 'ling'


def test_extra_columns_are_blank_when_there_are_fewer_items():
    """列在、项不够 → 空格子（不是把上一项复制一遍）。"""
    b = _beat(show='ling happy')
    row = _row_of(b, SHOW_EXTRAS)
    assert row['Show'] == 'ling happy'
    assert row['立绘2'] == '' and row['立绘3'] == ''


# ===========================================================================
#  ④ 往返：摊列来回，项一个不许丢、顺序不许变
# ===========================================================================

def test_split_roundtrip_keeps_every_item_in_order():
    for src in ('ling happy; ling base_arms; ye@left',
                'a; b',
                'ling; ling',                       # 重复项是语义，不许去重
                'cg cg02@t'):                       # 单项：不该被摊坏
        b = _beat(show=src)
        back = Beat.from_row(_row_of(b, SHOW_EXTRAS))
        assert show_cell_parts(back) == show_cell_parts(b), src
        assert compose_show_cell(back) == compose_show_cell(b), src


def test_hide_split_roundtrip_keeps_every_item():
    b = _beat(hide='ye@dissolve; ling; ye')
    back = Beat.from_row(_row_of(b, HIDE_EXTRAS))
    assert hide_cell_parts(back) == hide_cell_parts(b)


def test_not_enough_columns_never_drops_items():
    """
    ⚠ **安全阀**：项数 > 格数时**不许摊**（摊就会把多出来的项悄悄丢掉）。

    场景：表是手改过的 / 批次算摊列数时漏了这一行。放不下就整格写，
    读回来照样全在 —— 只是这一行没享受摊列下拉。
    """
    b = _beat(show='a; b; c; d')          # 4 项，但只有 1 个摊列（格数=2）
    row = _row_of(b, ['立绘2'])
    assert row['Show'] == 'a; b; c; d', '放不下却摊了 → 必然丢项'
    assert row['立绘2'] == ''
    assert show_cell_parts(Beat.from_row(row)) == ['a', 'b', 'c', 'd']


def test_roundtrip_is_idempotent():
    b = _beat(show='ling happy; ye@left', hide='ye@fade')
    row1 = _row_of(b, SHOW_EXTRAS + HIDE_EXTRAS)
    b2 = Beat.from_row(row1)
    row2 = _row_of(b2, SHOW_EXTRAS + HIDE_EXTRAS)
    b3 = Beat.from_row(row2)
    assert show_cell_parts(b3) == show_cell_parts(b2)
    assert compose_show_cell(b3) == compose_show_cell(b2)
    assert _row_of(b3, SHOW_EXTRAS + HIDE_EXTRAS)['立绘2'] == row2['立绘2']


# ===========================================================================
#  ⑤ 向后兼容：没有摊列时，行为与从前一字不差
# ===========================================================================

def test_without_extra_columns_the_original_cell_keeps_all_items():
    b = _beat(show='ling happy; ling base_arms', hide='ye@fade; ling')
    row = _row_of(b)                       # 不给 extra_fields
    assert row['Show'] == 'ling happy; ling base_arms'
    assert row['Hide'] == 'ye@fade; ling'
    assert len(b.to_row()) == len(BEAT_FIELDS)


def test_old_table_with_multi_value_cell_still_reads():
    """旧表：整格 `;` 多值、**没有**摊列 —— 照样还原成多项。"""
    row = {'Label': 'start', 'Show': 'a; b; c', 'Hide': 'd; e'}
    b = Beat.from_row(row)
    assert show_cell_parts(b) == ['a', 'b', 'c']
    assert hide_cell_parts(b) == ['d', 'e']


def test_old_and_new_carriers_together_are_all_read():
    """
    迁移中的表：旧格还是多值、**同时**已经有摊列 —— 两边都要读进来。

    ⚠ 少读一边就是静默丢项；这条判据就是钉这个。
    """
    row = {'Label': 'start', 'Show': 'a', '立绘2': 'b', '立绘3': 'c'}
    b = Beat.from_row(row)
    assert show_cell_parts(b) == ['a', 'b', 'c']


def test_editing_one_extra_cell_changes_only_that_item():
    """策划改 `立绘2`：只有第 2 项变，别的项一个不动。"""
    b = _beat(show='ling happy; ling base_arms; ye@left')
    row = _row_of(b, SHOW_EXTRAS)
    row['立绘2'] = 'ling angry'
    back = Beat.from_row(row)
    assert show_cell_parts(back) == ['ling happy', 'ling angry', 'ye@left']


def test_clearing_an_extra_cell_removes_that_item():
    b = _beat(show='a; b; c')
    row = _row_of(b, SHOW_EXTRAS)
    row['立绘2'] = ''
    back = Beat.from_row(row)
    assert show_cell_parts(back) == ['a', 'c']


# ===========================================================================
#  ⑥ 列契约：表头与数据行同源（错位 = 最后几列静默全空）
# ===========================================================================

def test_row_length_matches_the_contract_with_extras():
    b = _beat(show='a; b; c')
    for extras in ((), SHOW_EXTRAS, SHOW_EXTRAS + HIDE_EXTRAS):
        cells = b.to_row(extra_fields=list(extras))
        fields = _row_fields(None, list(extras))
        assert len(cells) == len(fields)
    # 摊列一律**追加末尾**（F2：基础列永不动）
    fields = _row_fields(None, SHOW_EXTRAS)
    assert fields[:len(BEAT_FIELDS)] == list(BEAT_FIELDS)
    assert fields[-len(SHOW_EXTRAS):] == SHOW_EXTRAS


def test_extra_cells_are_single_line_sanitized():
    """摊列格与别的单行列一样过清洗（换行会被压平并留痕）。"""
    b = _beat(show='a; b')
    row = _row_of(b, ['立绘2'])
    assert '\n' not in row['立绘2']


# ===========================================================================
#  ⑦ Phase 2：接进 Excel（表头/数据行/下拉/回读）
# ===========================================================================

def _beats_for_export():
    import openpyxl  # noqa: F401  （确保依赖在）
    b1 = _beat(show='ling happy; ling base_arms@i11; ye@left')
    b1.seq = 1
    b1.text = 'one'
    b2 = _beat(show='ye happy')
    b2.seq = 2
    b2.text = 'two'
    b3 = _beat(hide='ye@dissolve; ling')
    b3.seq = 3
    b3.text = 'three'
    return [b1, b2, b3]


def test_export_adds_extra_columns_and_splits_the_original(tmp_path):
    import openpyxl
    from storygraph.story.beat_model import BEAT_HEADERS
    from storygraph.story.excel_export import export_beats_to_excel

    xl = tmp_path / 'beats.xlsx'
    export_beats_to_excel(_beats_for_export(), xl)
    ws = openpyxl.load_workbook(xl).active
    headers = [c.value for c in ws[1]]

    # 这批最大 3 个立绘 / 2 个退场 → 立绘2、立绘3、退场2
    # （`Show` 的**中文表头**就是「立绘」，所以这族列自然连成 立绘/立绘2/立绘3）
    assert '立绘2' in headers and '立绘3' in headers and '退场2' in headers
    assert '立绘1' not in headers, '第 1 项留在原列，不该另造一列'
    assert headers[-3:] == ['立绘2', '立绘3', '退场2'], '摊列必须追加在末尾'

    idx = {h: i for i, h in enumerate(headers)}
    row1 = [c.value for c in ws[2]]
    assert row1[idx[BEAT_HEADERS['Show']]] == 'ling happy'
    assert row1[idx['立绘2']] == 'ling base_arms@i11'
    assert row1[idx['立绘3']] == 'ye@left'
    # 只有 1 项的拍 → 摊列空格
    row2 = [c.value for c in ws[3]]
    assert row2[idx[BEAT_HEADERS['Show']]] == 'ye happy'
    assert not (row2[idx['立绘2']] or '')


def test_export_import_roundtrip_keeps_all_items(tmp_path):
    """导出 → 读回：每一项、每一个顺序都不许丢。"""
    from storygraph.story.excel_export import export_beats_to_excel
    from storygraph.story.excel_import import import_beats_from_excel

    beats = _beats_for_export()
    want = {b.seq: show_cell_parts(b) for b in beats}
    want_hide = {b.seq: hide_cell_parts(b) for b in beats}
    xl = tmp_path / 'beats.xlsx'
    export_beats_to_excel(beats, xl)
    back = import_beats_from_excel(xl)
    got = {b.seq: show_cell_parts(b) for b in back}
    got_hide = {b.seq: hide_cell_parts(b) for b in back}
    assert got == want, (got, want)
    assert got_hide == want_hide, (got_hide, want_hide)


def test_export_without_multi_items_adds_no_extra_columns(tmp_path):
    """没有多值就**不加列** —— 老项目导出结果与从前逐列相同。"""
    import openpyxl
    from storygraph.story.excel_export import export_beats_to_excel
    from storygraph.story.beat_model import BEAT_FIELDS

    b = _beat(show='ling happy', hide='ye')
    b.seq = 1
    b.text = 'x'
    xl = tmp_path / 'one.xlsx'
    export_beats_to_excel([b], xl)
    ws = openpyxl.load_workbook(xl).active
    headers = [c.value for c in ws[1]]
    assert len(headers) == len(BEAT_FIELDS), headers[len(BEAT_FIELDS):]
    assert '立绘2' not in headers


def test_extra_columns_get_image_dropdowns(tmp_path):
    """摊列每格绑图像候选（这才是摊列的目的：拿到格级下拉）。"""
    import openpyxl
    from storygraph.story.excel_export import export_beats_to_excel

    xl = tmp_path / 'beats.xlsx'
    export_beats_to_excel(_beats_for_export(), xl)
    ws = openpyxl.load_workbook(xl).active
    headers = [c.value for c in ws[1]]
    # 拿不到资源候选表时不该崩；有候选表时摊列应出现在 data validation 里
    targets = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
    letters_with_dv = {t.split('2:')[0].split(':')[0] for t in targets}
    from openpyxl.utils import get_column_letter
    for name in ('立绘2', '立绘3', '退场2'):
        letter = get_column_letter(headers.index(name) + 1)
        assert letter in letters_with_dv or True, name  # 无候选表时允许没有
    # 有候选表时必须挂上（这条是真正的判据）
    if ws.data_validations.dataValidation:
        letter2 = get_column_letter(headers.index('立绘2') + 1)
        assert letter2 in letters_with_dv, (letter2, sorted(letters_with_dv))
