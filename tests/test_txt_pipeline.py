"""
纯文字导入 · 流水线判据（S3）
==============================

用户 2026-10-10 拍板：接进 `EditorSession`（S3）；"同一个写法在不同章指不同角色"
**先不考虑**（留接口）；建新角色时**立绘留空是对的**（可能是龙套）。

这个文件钉住流水线（`storygraph/authoring/txt_pipeline.py`）的命门：

  ① **落点与命名**：`<根>/game/chapter1_2.rpy`（作者的章节就该在 game/ 里）；
  ② ⚠ **未确认角色 → 一个字都不写**（产物里会是未定义变量 → 运行期必崩）；
  ③ **决定写回表 → 重跑转换 → 产物干净**（只有两种结局：干净，或什么都没写）；
  ④ 目标已存在且是**手写**的 → `write_generation` 拒写，原样报出来；
  ⑤ 手写不会出现 `define`（define 交给角色表同步链路）。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.assets.asset_registry import (append_rows_to_table,   # noqa: E402
                                             write_dict_to_excel, _TABLE_FIELDS)
from storygraph.authoring import txt_pipeline as TP                   # noqa: E402
from storygraph.authoring.txt_import import convert_text              # noqa: E402

ROWS = {
    'jc': {'ID': 'jc', 'Name': '「蒋忱」'},
    'mm': {'ID': 'mm', 'Name': '妈妈'},
    'txf': {'ID': 'txf', 'Name': '汤X凤'},
}


def _project(tmp_path: Path) -> Path:
    """一个最小项目：`game/` + `game/asset_tables/characters.xlsx`。"""
    root = tmp_path / 'proj'
    (root / 'game').mkdir(parents=True)
    adir = root / 'game' / 'asset_tables'
    adir.mkdir(parents=True, exist_ok=True)
    write_dict_to_excel(adir / 'characters.xlsx', ROWS, _TABLE_FIELDS['character'])
    return root


def _txt(tmp_path: Path, name: str, text: str) -> Path:
    fp = tmp_path / name
    fp.write_text(text, encoding='utf-8')
    return fp


# ==================== ① 落点与命名 ====================

def test_target_goes_into_the_projects_game_dir(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-2蓝之初见-定稿.txt', '1.0 chapter1-2 蓝之初见\nJC"你好"\n')
    plan = TP.build_plan(root, [fp], ROWS)
    assert len(plan.files) == 1
    f = plan.files[0]
    assert f.target == root / 'game' / 'chapter1_2.rpy', f.target
    assert f.label == 'chapter1_2'


def test_the_chapter_header_wins_over_the_filename_number(tmp_path):
    """
    ⚠ 真机实测：文件名编号与章节头**差一位**（`1-0初入大学-定稿.txt` 的头是
    `chapter0-0`）。目标名必须跟**章节头**（作者写的那个），否则会得到
    `chapter1_0.rpy` —— 与作者心里的那一章对不上，还会抢 `chapter0_0` 的 label。
    """
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt', '1.0 chapter0-0 初入大学\n妈妈"来了"\n')
    plan = TP.build_plan(root, [fp], ROWS)
    assert plan.files[0].label == 'chapter0_0'
    assert plan.files[0].target.name == 'chapter0_0.rpy'


def test_falls_back_to_the_filename_when_there_is_no_number_prefix(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '序章.txt', 'JC"你好"\n')
    plan = TP.build_plan(root, [fp], ROWS)
    assert plan.files[0].target.name == '序章.rpy'


def test_two_files_never_share_one_target(tmp_path):
    """同一批里两个文件算出同一个目标名 → 必须错开（否则后一个覆盖前一个）。"""
    root = _project(tmp_path)
    a = _txt(tmp_path, '1-2甲.txt', 'JC"甲"\n')
    b = tmp_path / 'other' / '1-2乙.txt'
    b.parent.mkdir()
    b.write_text('JC"乙"\n', encoding='utf-8')
    plan = TP.build_plan(root, [a, b], ROWS)
    targets = [f.target.name for f in plan.files]
    assert len(set(targets)) == 2, targets
    assert 'chapter1_2.rpy' in targets and 'chapter1_2_2.rpy' in targets


def test_docx_now_goes_through_the_same_chain(tmp_path):
    """S4：`.docx` 与 `.txt` 走**同一条**链路（只多一步"抽成文本"）。"""
    root = _project(tmp_path)
    from test_docx_import import _p, make_docx
    fp = make_docx(tmp_path / '1-0初入大学-定稿.docx',
                   _p('1.0 chapter0-0 初入大学'),
                   _p('妈妈“跟妈客气啥”'))
    res = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert res.ok, res.messages
    out = root / 'game' / 'chapter0_0.rpy'
    assert out in res.written
    assert '    mm "跟妈客气啥"' in out.read_text(encoding='utf-8')


def test_a_broken_docx_is_skipped_with_a_reason(tmp_path):
    """读不出来的 `.docx` 要说清原因，不许静默变成空章节。"""
    root = _project(tmp_path)
    fp = tmp_path / 'x.docx'
    fp.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 32)
    res = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert not res.written
    assert res.skipped and '另存为' in res.skipped[0][1]


def test_a_clean_file_needs_no_decision(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt',
              '1.0 chapter0-0 初入大学\n妈妈"跟妈客气啥"\n「蒋忱」"来了"\n')
    plan = TP.build_plan(root, [fp], ROWS)
    assert plan.files and not plan.skipped
    assert '    mm "跟妈客气啥"' in plan.files[0].rpy_text
    assert TP.proposals_for(plan, ROWS) == []


# ==================== ② 未确认 → 绝不落盘 ====================

def test_unconfirmed_characters_write_nothing_at_all(tmp_path):
    """
    ⚠ 本模块最重要的性质：`MOM` 没有对应角色时**一个字节都不许写** ——
    写下去就是未定义变量，游戏到那一句直接崩。
    """
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt', 'MOM"跟妈客气啥"\n')
    res = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert res.needs_decision
    assert [p.token for p in res.proposals] == ['MOM']
    assert res.written == [] and res.blocked == []
    assert not list((root / 'game').glob('*.rpy')), '没确认就写了文件！'
    assert '未定义变量' in ' '.join(res.messages)


def test_the_ask_callback_can_say_not_now(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt', 'MOM"跟妈客气啥"\n')
    seen = []

    def _ask(proposals):
        seen.extend(p.token for p in proposals)
        return None                       # "这次不决定"

    res = TP.run(root, [fp], rows=ROWS, project_root=root, ask=_ask)
    assert seen == ['MOM']
    assert res.needs_decision and not res.written
    assert not list((root / 'game').glob('*.rpy'))


def test_dry_run_never_writes(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt', '妈妈"跟妈客气啥"\n')
    res = TP.run(root, [fp], rows=ROWS, project_root=root, dry_run=True)
    assert res.plan.files and not res.written
    assert not list((root / 'game').glob('*.rpy'))


# ==================== ③ 决定 → 重跑 → 产物干净 ====================

def test_decisions_are_written_back_and_the_product_becomes_clean(tmp_path):
    """
    `MOM` → 当作 `mm` 的别名：决定写进 characters.xlsx → **重跑转换** →
    产物是 `mm "…"`（而不是 `MOM "…"`），且**没有 define**（define 归角色表链路）。
    """
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt',
              '1.0 chapter0-0 初入大学\nMOM"跟妈客气啥"\nMOM"同学要和睦"\n')

    def _ask(proposals):
        # 模拟弹窗：`MOM` 指到 `mm`
        return [{'token': 'MOM', 'action': 'existing', 'target': 'mm'}]

    imported = {}

    def _import(paths):
        imported['paths'] = list(paths)
        return {'ok': True, 'nodes': 1, 'edges': 0}

    res = TP.run(root, [fp], rows=ROWS, project_root=root, ask=_ask,
                 import_chapters=_import)
    assert res.ok, res.messages
    out = root / 'game' / 'chapter0_0.rpy'
    assert out in res.written
    text = out.read_text(encoding='utf-8')
    assert '    mm "跟妈客气啥"' in text
    assert '    mm "同学要和睦"' in text
    assert 'MOM' not in text, '未映射写法居然还在产物里'
    assert 'define' not in text, 'define 不该写进章节产物（应交给角色表链路）'
    assert imported['paths'] == [str(out)]

    # 决定真的落到表里了（下次导入不用再问）
    from storygraph.assets.asset_registry import parse_asset_tables
    rows = parse_asset_tables(root / 'game' / 'asset_tables').characters
    assert rows['mm']['Aliases'] == 'MOM'


def test_a_new_character_decision_creates_a_row_without_a_sprite(tmp_path):
    """
    "建新角色"：追加一行到 `characters.xlsx`；**立绘（ImageTag）留空**——
    用户 2026-10-10："它有可能是龙套，所以立绘留空是对的"。
    """
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-1到达青蓝实验室-定稿.txt',
              'BAOAN"小伙子，你要去哪个地方？"\n')

    def _ask(proposals):
        return [{'token': 'BAOAN', 'action': 'new', 'var': 'BAOAN',
                 'name': '保安'}]

    res = TP.run(root, [fp], rows=ROWS, project_root=root, ask=_ask)
    assert res.ok, res.messages
    from openpyxl import load_workbook
    ws = load_workbook(root / 'game' / 'asset_tables' / 'characters.xlsx').active
    headers = [c.value for c in ws[1]]
    got = {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(values_only=True)
           if r and r[0]}
    assert 'BAOAN' in got
    assert got['BAOAN']['Name'] == '保安'
    assert not (got['BAOAN'].get('ImageTag') or '').strip(), got['BAOAN']
    assert '    BAOAN "小伙子，你要去哪个地方？"' in \
        (root / 'game' / 'chapter1_1.rpy').read_text(encoding='utf-8')


def test_saying_no_to_everything_leaves_the_project_untouched(tmp_path):
    """弹窗里全选"先跳过" → 不落盘（不能因为"问过了"就写半成品）。"""
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt', 'MOM"跟妈客气啥"\n')
    res = TP.run(root, [fp], rows=ROWS, project_root=root,
                 ask=lambda p: [{'token': 'MOM', 'action': 'skip'}])
    assert not res.written
    assert not list((root / 'game').glob('*.rpy'))
    assert res.needs_decision, '还有认不出的写法时必须回到"要人决定"'


# ==================== ④ 门禁：手写文件不许覆盖 ====================

def test_a_handwritten_target_is_refused_not_overwritten(tmp_path):
    """
    目标已存在且**不是工具生成的** → `write_generation` 拒写，这里原样报出来
    （B2/P0：全量重生成不认得 `default`/`define`，覆盖手写文件 = 静默毁定义）。
    """
    root = _project(tmp_path)
    hand = root / 'game' / 'chapter0_0.rpy'
    hand.write_text('define x = Character("手写")\nlabel chapter0_0:\n    "原文"\n',
                    encoding='utf-8')
    before = hand.read_text(encoding='utf-8')
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt',
              '1.0 chapter0-0 初入大学\n妈妈"跟妈客气啥"\n')
    res = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert not res.written
    assert res.blocked and '手写' in res.blocked[0][1]
    assert hand.read_text(encoding='utf-8') == before, '手写文件被改了！'


def test_a_generated_target_may_be_re_imported(tmp_path):
    """带生成标记的目标 = 工具自己以前写的 → 允许重导（幂等）。"""
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt',
              '1.0 chapter0-0 初入大学\n妈妈"跟妈客气啥"\n')
    first = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert first.ok, first.messages
    before = (root / 'game' / 'chapter0_0.rpy').read_text(encoding='utf-8')
    again = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert again.ok, again.messages
    after = (root / 'game' / 'chapter0_0.rpy').read_text(encoding='utf-8')
    assert before == after, '同输入重导两次产物不一致（不是幂等）'


# ==================== ⑤ 与 render 的一致性（继承 S2 的性质） ====================

def test_the_written_product_roundtrips_byte_for_byte(tmp_path):
    root = _project(tmp_path)
    fp = _txt(tmp_path, '1-0初入大学-定稿.txt',
              '1.0 chapter0-0 初入大学\n（bg1.1）\n妈妈"跟妈客气啥"\n旁白一句\n')
    res = TP.run(root, [fp], rows=ROWS, project_root=root)
    assert res.ok, res.messages
    text = res.written[0].read_text(encoding='utf-8')

    from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
    from storygraph.story.beat_build import beats_from_parse_result
    from storygraph.story.beat_render import render_beats_to_rpy
    pr, _ = IncrementalRenPyAnalyzer().analyze(files=[res.written[0]], graph=None)
    assert render_beats_to_rpy(beats_from_parse_result(pr)) == text
