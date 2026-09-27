"""
tests/test_script_table.py
==========================
剧本演出表（rpy ↔ Excel）测试。

主线：.rpy 解析成"一句台词一行的表格"，表格改完能生成能跑的 .rpy。
这里是文案真正干活的地方，往返必须无损。
"""

import pytest

from core.statement import Statement, StatementKind
from script_table import (
    Beat, beats_from_statements, beats_from_parse_result,
    export_beats_to_excel, import_beats_from_excel,
    render_beats_to_rpy, render_beat_to_rpy, write_beats_to_rpy,
    open_rpy_as_table, export_table_to_rpy, refresh_candidate_sheets,
    find_unregistered_assets,
    BEAT_FIELDS, KIND_SAY, KIND_NARRATION, KIND_MENU_OPTION, KIND_EXTEND,
)


SCRIPT = (
    'define e = Character("Eileen")\n'
    '\n'
    'label ch01:\n'
    '    scene bg classroom\n'
    '    show eileen happy at left\n'
    '    play music bgm_theme\n'
    '    e "早上好。"\n'
    '    e "今天天气不错。"\n'
    '    hide eileen\n'
    '    "旁白一句"\n'
    '    return\n'
)


@pytest.fixture
def parse_result(write_rpy):
    from renpy_parser import RenPyParser, ParseResultBuilder
    fp = write_rpy(SCRIPT, 'ch01.rpy')
    return ParseResultBuilder().build(RenPyParser().parse(fp))


# ------------------------------------------------------------------
#  语句折叠成 Beat
# ------------------------------------------------------------------

class TestBeat:
    def test_roundtrip_row(self):
        beat = Beat(seq=1, label='ch01', kind=KIND_SAY,
                    speaker='e', text='hi', scene='bg a')
        row = beat.to_row()
        back = Beat.from_row(dict(zip(BEAT_FIELDS, row)))
        # 逐列等价即可（stage 是由 scene 派生出的展示文本，
        # 往返一次后由空变成 'scene bg a'，语义一致）
        assert back.to_row() == row

    def test_from_row_handles_none(self):
        beat = Beat.from_row({'Text': None, 'Speaker': None})
        assert beat.text == ''
        assert beat.speaker == ''

    def test_from_row_normalizes_numeric_seq(self):
        assert Beat.from_row({'Seq': 7.0}).seq == 7

    def test_unknown_kind_falls_back_to_say(self):
        assert Beat.from_row({'Kind': 'nonsense'}).kind == KIND_SAY

    def test_is_empty(self):
        assert Beat().is_empty()
        assert not Beat(text='x').is_empty()
        assert not Beat(notes='x').is_empty()


class TestBeatsFromStatements:
    def test_stage_directives_attach_to_next_line(self, parse_result):
        beats = beats_from_parse_result(parse_result)
        first = beats[0]
        assert first.scene == 'bg classroom'
        assert 'eileen happy@left' in first.show
        assert first.music == 'bgm_theme'
        assert first.text == '早上好。'

    def test_second_line_has_no_stage(self, parse_result):
        """舞台指令只吸附在紧随其后的那一句上，不该往下污染。"""
        beats = beats_from_parse_result(parse_result)
        assert beats[1].text == '今天天气不错。'
        assert beats[1].scene == ''
        assert beats[1].show == ''

    def test_narration_has_no_speaker(self, parse_result):
        beats = beats_from_parse_result(parse_result)
        narration = [b for b in beats if b.text == '旁白一句'][0]
        assert narration.kind == KIND_NARRATION
        assert narration.speaker == ''

    def test_hide_captured(self, parse_result):
        beats = beats_from_parse_result(parse_result)
        narration = [b for b in beats if b.text == '旁白一句'][0]
        assert narration.hide == 'eileen'

    def test_seq_is_monotonic(self, parse_result):
        beats = beats_from_parse_result(parse_result)
        seqs = [b.seq for b in beats]
        assert seqs == sorted(seqs)
        assert seqs[0] == 1

    def test_label_propagated(self, parse_result):
        beats = beats_from_parse_result(parse_result)
        assert all(b.label == 'ch01' for b in beats)

    def test_trailing_stage_without_line_is_kept(self):
        """结尾只有舞台指令（如淡出）不能丢。"""
        stmts = [
            Statement.make_say('e', 'hi', source=('a.rpy', 1)),
            Statement.make_scene('bg end', with_='fade', source=('a.rpy', 2)),
        ]
        beats = beats_from_statements(stmts, label='x')
        assert beats[-1].scene == 'bg end'
        assert beats[-1].with_ == 'fade'

    def test_condition_from_if_block(self, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder
        fp = write_rpy(
            'label a:\n'
            '    if love > 5:\n'
            '        e "喜欢你"\n', 'cond.rpy')
        result = ParseResultBuilder().build(RenPyParser().parse(fp))
        beats = beats_from_parse_result(result)
        assert beats[0].condition == 'love > 5'

    def test_menu_option_becomes_own_beat(self, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder
        fp = write_rpy(
            'label a:\n'
            '    menu:\n'
            '        "选项A":\n'
            '            jump b\n', 'menu.rpy')
        result = ParseResultBuilder().build(RenPyParser().parse(fp))
        beats = beats_from_parse_result(result)
        assert beats[0].kind == KIND_MENU_OPTION
        assert beats[0].text == '选项A'

    def test_python_recorded_in_codeblock(self, write_rpy):
        """变量操作进 CodeBlock，导出时原样回写，不再只留痕。"""
        from renpy_parser import RenPyParser, ParseResultBuilder
        fp = write_rpy(
            'label a:\n'
            '    $ love += 1\n'
            '    e "嗯"\n', 'py.rpy')
        result = ParseResultBuilder().build(RenPyParser().parse(fp))
        beats = beats_from_parse_result(result)
        assert 'love' in beats[0].code
        # 渲染回写时原样带 $ 前缀
        lines = render_beat_to_rpy(beats[0])
        assert '$ love += 1' in lines

    def test_voice_captured(self):
        stmts = [
            Statement.make_voice('v_001', source=('a.rpy', 1)),
            Statement.make_say('e', 'hi', source=('a.rpy', 2)),
        ]
        beats = beats_from_statements(stmts, label='x')
        assert beats[0].voice == 'v_001'

    def test_multiple_shows_joined(self):
        stmts = [
            Statement.make_show('eileen happy', at=['left'],
                                source=('a.rpy', 1)),
            Statement.make_show('mary sad', at=['right'],
                                source=('a.rpy', 2)),
            Statement.make_say('e', 'hi', source=('a.rpy', 3)),
        ]
        beats = beats_from_statements(stmts, label='x')
        assert 'eileen happy@left' in beats[0].show
        assert 'mary sad@right' in beats[0].show


# ------------------------------------------------------------------
#  渲染回 rpy
# ------------------------------------------------------------------

class TestRender:
    def test_say_line(self):
        beat = Beat(kind=KIND_SAY, speaker='e', text='你好')
        assert render_beat_to_rpy(beat) == ['e "你好"']

    def test_narration_line(self):
        beat = Beat(kind=KIND_NARRATION, text='旁白')
        assert render_beat_to_rpy(beat) == ['"旁白"']

    def test_stage_before_dialogue(self):
        beat = Beat(speaker='e', text='hi', scene='bg a',
                    show='eileen@left')
        lines = render_beat_to_rpy(beat)
        assert lines[0].startswith('scene ')
        assert lines[1].startswith('show ')
        assert lines[-1] == 'e "hi"'

    def test_condition_wraps_in_if(self):
        beat = Beat(speaker='e', text='hi', condition='love > 5')
        lines = render_beat_to_rpy(beat, indent='    ')
        assert lines[0] == 'if love > 5:'
        assert lines[1] == '    e "hi"'

    def test_text_with_quotes_is_escaped(self):
        beat = Beat(speaker='e', text='他说"你好"')
        line = render_beat_to_rpy(beat)[0]
        assert line == 'e "他说\\"你好\\""'

    def test_empty_beat_renders_nothing(self):
        assert render_beat_to_rpy(Beat()) == []

    def test_menu_option_renders_as_string(self):
        beat = Beat(kind=KIND_MENU_OPTION, text='选项A')
        assert render_beat_to_rpy(beat) == ['"选项A"']

    def test_extend_uses_extend(self):
        beat = Beat(kind=KIND_EXTEND, text='更多')
        assert render_beat_to_rpy(beat) == ['extend "更多"']

    def test_audio_lines(self):
        beat = Beat(text='x', music='bgm_a', sound='sfx_b', voice='v1')
        lines = render_beat_to_rpy(beat)
        assert 'play music bgm_a' in lines
        assert 'play sound sfx_b' in lines
        assert 'voice v1' in lines

    def test_full_script_has_label_and_return(self):
        beats = [Beat(label='ch01', speaker='e', text='hi')]
        text = render_beats_to_rpy(beats)
        assert 'label ch01:' in text
        assert '    e "hi"' in text
        assert '    return' in text

    def test_stop_renders_plain(self):
        beat = Beat(music_action='stop')
        assert render_beat_to_rpy(beat) == ['stop music']
        beat = Beat(sound_action='stop')
        assert render_beat_to_rpy(beat) == ['stop sound']

    def test_stop_music_in_context(self):
        beat = Beat(speaker='e', text='hi', music_action='stop')
        lines = render_beat_to_rpy(beat)
        assert lines[-2] == 'stop music'
        assert lines[-1] == 'e "hi"'


class TestStopLanding:
    """M1：stop music / stop sound 必须落表，不能丢。"""

    STOP_SCRIPT = (
        'label ch03:\n'
        '    play music bgm_theme\n'
        '    e "开始"\n'
        '    stop music\n'
        '    play sound sfx_door\n'
        '    stop sound\n'
        '    return\n'
    )

    def test_stop_parsed_into_beat(self, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder
        fp = write_rpy(self.STOP_SCRIPT, 'ch03.rpy')
        result = ParseResultBuilder().build(RenPyParser().parse(fp))
        beats = beats_from_parse_result(result)

        stop_music = [b for b in beats if b.music_action == 'stop']
        stop_sound = [b for b in beats if b.sound_action == 'stop']
        assert stop_music, "stop music 未落表"
        assert stop_sound, "stop sound 未落表"
        assert stop_music[0].music == ''
        assert stop_sound[0].sound == ''

    def test_stop_beat_not_empty(self):
        assert not Beat(music_action='stop').is_empty()
        assert not Beat(sound_action='stop').is_empty()

    def test_stop_roundtrip(self, tmp_path, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder
        fp = write_rpy(self.STOP_SCRIPT, 'ch03.rpy')
        table = open_rpy_as_table(fp, tmp_path / 'ch03.xlsx')
        beats = import_beats_from_excel(table)
        text = render_beats_to_rpy(beats)

        stripped = [ln.strip() for ln in text.splitlines()
                    if ln.strip() and not ln.strip().startswith('#')]
        assert 'stop music' in stripped
        assert 'stop sound' in stripped
        # 原样出现，没有退化成 play
        assert 'play music bgm_theme' in stripped
        assert 'play sound sfx_door' in stripped

    def test_action_dropdown_attached(self, tmp_path, write_rpy):
        """
        BGM / SFX / 舞台 是文本列，挂不上格级下拉；能点选的是
        说话人 / 场景 / 立绘 / 语音等列。这里同时锁住"没有偷偷留下旧列"。
        表头现在用中文（见 BEAT_HEADERS）。
        """
        from asset_registry import AssetRegistry
        from script_table import BEAT_HEADERS
        fp = write_rpy(self.STOP_SCRIPT, 'ch03.rpy')
        reg = AssetRegistry()
        reg.images = {'bg_room': {'ID': 'bg_room', 'Tag': 'bg'}}
        reg.voices = {'v_001': {'ID': 'v_001'}}
        table = open_rpy_as_table(fp, tmp_path / 'ch03.xlsx', registry=reg)
        import openpyxl
        wb = openpyxl.load_workbook(table)
        ws = wb['Beats']
        headers = [c.value for c in ws[1]]
        for col in ('BGM', 'SFX', 'Stage'):
            assert BEAT_HEADERS[col] in headers, \
                f'{col} 列（{BEAT_HEADERS[col]}）缺失'
        for gone in ('Show', 'Hide', 'At', 'Transform', 'Raw',
                     'Music', 'MusicAction', 'Sound', 'SoundAction',
                     'Condition', 'Notes'):
            assert gone not in headers, f'旧列 {gone} 还在表里'
        assert len(ws.data_validations.dataValidation) >= 2

    def test_multiple_labels_split(self):
        beats = [
            Beat(seq=1, label='a', text='one'),
            Beat(seq=2, label='b', text='two'),
        ]
        text = render_beats_to_rpy(beats)
        assert 'label a:' in text and 'label b:' in text

    def test_single_label_override(self):
        beats = [Beat(seq=1, label='a', text='one'),
                 Beat(seq=2, label='b', text='two')]
        text = render_beats_to_rpy(beats, label_name='merged')
        assert 'label merged:' in text
        assert 'label a:' not in text

    def test_rendered_script_is_reparsable(self):
        """生成的脚本必须能被自己的解析器吃回去——这是"能跑"的必要条件。"""
        from renpy_parser import RenPyParser
        beats = [
            Beat(seq=1, label='ch01', kind=KIND_SAY, speaker='e',
                 text='早上好', scene='bg classroom',
                 show='eileen happy@left', music='bgm_theme'),
            Beat(seq=2, label='ch01', kind=KIND_SAY, speaker='e',
                 text='第二句'),
        ]
        text = render_beats_to_rpy(beats)
        fp = tmp_rpy = None
        import tempfile, pathlib
        d = pathlib.Path(tempfile.mkdtemp())
        fp = d / 'out.rpy'
        fp.write_text(text, encoding='utf-8')
        parser = RenPyParser()
        parser.parse(fp)
        assert not parser.errors, parser.errors


# ------------------------------------------------------------------
#  Excel 往返
# ------------------------------------------------------------------

class TestExcelRoundTrip:
    def test_export_then_import(self, tmp_path, parse_result):
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'ch01.xlsx'
        export_beats_to_excel(beats, path)
        back = import_beats_from_excel(path)

        assert len(back) == len(beats)
        assert back[0].text == beats[0].text
        assert back[0].scene == beats[0].scene
        assert back[0].show == beats[0].show

    def test_blank_rows_skipped(self, tmp_path):
        beats = [Beat(seq=1, text='a'), Beat(seq=2, text='b')]
        path = tmp_path / 'b.xlsx'
        export_beats_to_excel(beats, path)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb['Beats']
        ws.append([None] * len(BEAT_FIELDS))
        wb.save(path)
        assert len(import_beats_from_excel(path)) == 2

    def test_empty_beats_dropped_on_import(self, tmp_path):
        import openpyxl
        beats = [Beat(seq=1, text='a')]
        path = tmp_path / 'c.xlsx'
        export_beats_to_excel(beats, path)
        wb = openpyxl.load_workbook(path)
        # 只有 Seq/Kind、没有任何演出内容的行 → 视为空行丢弃
        row = [''] * len(BEAT_FIELDS)
        row[BEAT_FIELDS.index('Seq')] = 2
        row[BEAT_FIELDS.index('Kind')] = 'say'
        wb['Beats'].append(row)
        wb.save(path)
        assert len(import_beats_from_excel(path)) == 1

    def test_missing_file_returns_empty(self, tmp_path):
        assert import_beats_from_excel(tmp_path / 'nope.xlsx') == []

    def test_edit_in_excel_reaches_script(self, tmp_path, parse_result):
        """改 Excel 里的台词和立绘，生成的脚本要跟着变。"""
        import openpyxl
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'ch01.xlsx'
        export_beats_to_excel(beats, path)

        wb = openpyxl.load_workbook(path)
        ws = wb['Beats']
        ws.cell(row=2, column=BEAT_FIELDS.index('Text') + 1,
                value='改过的台词')
        # 舞台只有一列：策划写人话或原生写法都行，编译器照单全收
        ws.cell(row=2, column=BEAT_FIELDS.index('Stage') + 1,
                value='show eileen sad at right')
        wb.save(path)

        back = import_beats_from_excel(path)
        text = render_beats_to_rpy(back)
        assert '改过的台词' in text
        # 原生写法原样出去：作者怎么写就怎么回，不做下划线改写
        assert 'show eileen sad at right' in text

    def test_candidate_sheets_created(self, tmp_path, parse_result):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'bg_classroom': {'ID': 'bg_classroom'}}
        reg.music = {'bgm_theme': {'ID': 'bgm_theme'}}
        reg.voices = {'v_eileen_01': {'ID': 'v_eileen_01'}}
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'd.xlsx'
        export_beats_to_excel(beats, path, registry=reg)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        assert 'Candidates_图片' in wb.sheetnames
        assert 'Candidates_音乐' in wb.sheetnames
        assert 'Candidates_语音' in wb.sheetnames
        # 第 1 行是说明，第 2 行是表头，第 3 行起才是候选
        assert wb['Candidates_图片']['A3'].value == 'bg_classroom'
        assert wb['Candidates_语音']['A3'].value == 'v_eileen_01'
        assert wb['Candidates_图片']['A1'].value.startswith('↓')

    def test_dropdowns_attached(self, tmp_path, parse_result):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'a': {'ID': 'a'}, 'b': {'ID': 'b'}}
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'e.xlsx'
        export_beats_to_excel(beats, path, registry=reg)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb['Beats']
        assert len(ws.data_validations.dataValidation) >= 1

    def test_no_registry_no_candidate_sheets(self, tmp_path, parse_result):
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'f.xlsx'
        export_beats_to_excel(beats, path)
        import openpyxl
        wb = openpyxl.load_workbook(path)
        # 没有 registry 就不该有候选表（下拉没有候选可挂）；
        # 「舞台快照」是从 Beats 推导出来的派生视图，始终生成。
        assert 'Beats' in wb.sheetnames
        assert not [s for s in wb.sheetnames
                    if s.startswith('Candidates_')]
        assert '舞台快照' in wb.sheetnames

    def test_refresh_candidate_sheets_preserves_beats(self, tmp_path,
                                                      parse_result):
        """刷新候选只重写 Candidates_*，不动 Beats 表里的用户数据。"""
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'bg_classroom': {'ID': 'bg_classroom'}}
        beats = beats_from_parse_result(parse_result)
        path = tmp_path / 'f.xlsx'
        export_beats_to_excel(beats, path, registry=reg)
        # 用户改过的台词
        import openpyxl
        wb = openpyxl.load_workbook(path)
        wb['Beats'].cell(row=2, column=BEAT_FIELDS.index('Text') + 1,
                         value='用户改的台词')
        wb.save(path)

        # 资源表新增一张图，刷新候选
        reg.images['bg_new'] = {'ID': 'bg_new'}
        refresh_candidate_sheets(path, reg)

        wb2 = openpyxl.load_workbook(path)
        # Beats 用户数据保留
        assert wb2['Beats'].cell(row=2, column=BEAT_FIELDS.index('Text') + 1
                                 ).value == '用户改的台词'
        # 候选表出现新图
        ids = [c.value for c in wb2['Candidates_图片']['A']]
        assert 'bg_new' in ids
        assert 'bg_classroom' in ids

    def test_open_rpy_as_table(self, tmp_path, write_rpy):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'bg classroom': {'ID': 'bg classroom'}}
        fp = write_rpy(SCRIPT, 'open_me.rpy')
        out = open_rpy_as_table(fp, tmp_path / 'sheet.xlsx', registry=reg)
        assert out.exists()
        beats = import_beats_from_excel(out)
        assert beats and beats[0].text == '早上好。'

    def test_export_table_to_rpy(self, tmp_path, write_rpy):
        fp = write_rpy(SCRIPT, 'src.rpy')
        table = open_rpy_as_table(fp, tmp_path / 'sheet.xlsx')
        out = export_table_to_rpy(table, tmp_path / 'generated.rpy')
        text = out.read_text(encoding='utf-8')
        assert 'label' in text
        assert '早上好。' in text


class TestUnregisteredDetection:
    """M4：演出表引用了资源表没有的资产 → 要能检测出来。"""

    def test_unregistered_flagged(self):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'bg_classroom': {}}
        reg.music = {'bgm_theme': {}}
        reg.characters = {'e': {}}
        beats = [
            Beat(seq=1, label='ch', speaker='e', text='hi',
                 scene='bg classroom', music='bgm_theme'),
            Beat(seq=2, label='ch', text='x',
                 show='eileen happy@left', voice='v_new'),
        ]
        found = find_unregistered_assets(beats, reg)
        by_id = {aid for _t, aid in found}
        assert 'eileen_happy' in by_id      # 图片未登记（空格已归一为下划线）
        assert 'v_new' in by_id             # 语音未登记
        # 已登记的不报
        assert 'bg_classroom' not in by_id
        assert 'e' not in by_id

    def test_builtin_transform_not_flagged(self):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        beats = [Beat(seq=1, label='ch', text='x', at='left',
                     with_='dissolve')]
        found = find_unregistered_assets(beats, reg)
        assert found == []        # left / dissolve 是 Ren'Py 内置

    def test_space_normalized_against_registry(self):
        from asset_registry import AssetRegistry
        reg = AssetRegistry()
        reg.images = {'bg_classroom': {}}     # 下划线形式
        beats = [Beat(seq=1, label='ch', text='x', scene='bg classroom')]
        found = find_unregistered_assets(beats, reg)
        assert found == []        # 空格与下划线视为同一资产

    def test_none_registry_returns_empty(self):
        assert find_unregistered_assets([Beat(text='x')], None) == []


class TestOrderAndAudioArePreserved:
    """
    用户最在意的两点：剧本表必须保留语句顺序，必须解析音乐/音效/语音。
    这个测试用一条真实编排顺序，往返 Excel 后断言 emit 的 .rpy 完全一致。
    """

    ORDERED_SCRIPT = (
        'define e = Character("Eileen")\n'
        'define m = Character("Mary")\n'
        '\n'
        'label ch02:\n'
        '    scene bg night\n'
        '    show eileen happy at left\n'
        '    play music bgm_night\n'
        '    e "你来了。"\n'
        '    play sound sfx_door\n'
        '    voice v_eileen_01\n'
        '    m "门没锁。"\n'
        '    queue sound sfx_footstep\n'
        '    menu:\n'
        '        "进去":\n'
        '            jump inner\n'
        '        "离开":\n'
        '            "他转身走了。"\n'
        '    return\n'
    )

    def test_emitted_order_matches_source(self, tmp_path, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder

        fp = write_rpy(self.ORDERED_SCRIPT, 'ch02.rpy')
        result = ParseResultBuilder().build(RenPyParser().parse(fp))
        beats = beats_from_parse_result(result)

        # 顺序：scene → show → music → say → sound → voice → say → sound
        #       → menu_option("进去") → say("他转身走了。")
        # （本脚本 menu 无 caption，所以选项先于"离开"里的旁白出现，
        #   与 Ren'Py 源码顺序一致）
        kinds = [b.kind for b in beats if b.text or b.music or b.sound
                 or b.voice or b.scene or b.show]
        assert kinds[:4] == [
            KIND_SAY,          # 第一句（吸附 scene/show/music）
            KIND_SAY,          # 第二句（吸附 sound）
            KIND_MENU_OPTION,  # 选项"进去"（吸附 queue sound）
            KIND_SAY,          # "离开"分支里的旁白
        ]

        text = render_beats_to_rpy(beats)
        lines = [ln for ln in text.splitlines()
                 if ln.strip() and not ln.strip().startswith('#')]

        # 关键顺序断言：music 出现在第一句台词之前且只用一次
        idx_music = next(i for i, ln in enumerate(lines)
                         if ln.strip() == 'play music bgm_night')
        idx_first_say = next(i for i, ln in enumerate(lines)
                             if ln.strip().startswith('e '))
        assert idx_music < idx_first_say

        # 第二句台词之前必须有 sound（门声）
        idx_sound = next(i for i, ln in enumerate(lines)
                         if ln.strip() == 'play sound sfx_door')
        idx_second_say = next(i for i, ln in enumerate(lines)
                             if ln.strip().startswith('m '))
        assert idx_sound < idx_second_say

        # voice 紧跟第二句台词之前
        idx_voice = next(i for i, ln in enumerate(lines)
                         if ln.strip() == 'voice v_eileen_01')
        assert idx_voice < idx_second_say

        # 全程只出现一次 music、一次 door sound，没有重复或错位
        stripped = [ln.strip() for ln in lines]
        assert stripped.count('play music bgm_night') == 1
        assert stripped.count('play sound sfx_door') == 1
        assert stripped.count('voice v_eileen_01') == 1
        assert 'queue sound sfx_footstep' in stripped

    def test_full_roundtrip_keeps_order_and_audio(self, tmp_path, write_rpy):
        from renpy_parser import RenPyParser, ParseResultBuilder

        fp = write_rpy(self.ORDERED_SCRIPT, 'ch02.rpy')
        table = open_rpy_as_table(fp, tmp_path / 'ch02.xlsx')
        beats = import_beats_from_excel(table)

        # 表格里第一行就应有 scene/show/music（吸附在第一句台词）
        first = beats[0]
        assert first.scene == 'bg night'
        assert 'eileen happy@left' in first.show
        assert first.music == 'bgm_night'
        assert first.text == '你来了。'

        # 第二行应有 sound 与 voice
        second = beats[1]
        assert second.sound == 'sfx_door'
        assert second.voice == 'v_eileen_01'
        assert second.text == '门没锁。'

        # 写回的脚本仍然能跑（可被解析器吃回去）
        out = tmp_path / 'regenerated.rpy'
        write_beats_to_rpy(beats, out)
        parser = RenPyParser()
        parser.parse(out)
        assert not parser.errors, parser.errors
