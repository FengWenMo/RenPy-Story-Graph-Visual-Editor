"""
「每参数一列 + 原生下拉」（M4 第二阶段）
=======================================

上一轮（D1）只做到"给「头行」那一格挂**候选注释**"，理由是：
Excel 的数据验证**按列**绑定，而参数名每条语句各不相同 —— 一列装不下
两套枚举；要做"每参数一列"就得让 `BEAT_FIELDS` **动态化**，那是一次
列契约变更，需要单独一轮。

这一轮把那件事做完。四条判据，缺一条都会把用户坑到：

  ① **没有声明项目的表一个字节都不变**（列数都不许变）——
     新功能不许动老项目，这是底线；
  ② **有声明就有列 + 原生下拉**：列名 `mycue·intensity`，候选写进
     Excel 数据验证（`showErrorMessage=False`：手打新值不被拦）；
  ③ **改了那格 = 只改那个参数的值**：折回头行时只替换那一个值 token，
     引号、参数顺序、别的参数、行尾块标记一个字节不动；
  ④ **没改 = 逐字节原样**：`target="ling"` 这种原文不许被吃掉引号 ——
     表里每一格都会把值写回来，没有这条守卫，用户看到的是
     "我只是打开表又导出，引号就没了"。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
from storygraph.story.beat_build import beats_from_parse_result
from storygraph.story.beat_model import BEAT_FIELDS
from storygraph.story.beat_render import render_beats_to_rpy
from storygraph.story.excel_export import export_beats_to_excel
from storygraph.story.excel_import import import_beats_from_excel
from storygraph.story.statement_registry import StatementRegistry

#: 项目声明的自定义块语句（带两个参数，各给一组候选）
RULES = {
    'statements': [
        {'name': 'mycue',
         'label': '闪电提示',
         'columns': [
             {'name': 'intensity', 'label': '强度', 'choices': ['1', '2', '3']},
             {'name': 'target', 'label': '目标', 'choices': ['ling', 'lan']},
         ]},
    ],
}

SCRIPT = (
    'label a:\n'
    '    "第一句"\n'
    '    mycue "闪电" intensity=3 target=ling:\n'
    '        target = ling\n'
    '    "第二句"\n'
    '    return\n'
)


def _project(tmp_path, rules=RULES, script=SCRIPT):
    """造一个**最小真实项目**：`game/` + `storygraph.rules.json` + 一章。"""
    root = tmp_path / 'proj'
    (root / 'game').mkdir(parents=True, exist_ok=True)
    if rules is not None:
        (root / 'storygraph.rules.json').write_text(
            json.dumps(rules, ensure_ascii=False), encoding='utf-8')
    src = root / 'game' / 'ch.rpy'
    src.write_text(script, encoding='utf-8')
    return root, src


def _beats(src):
    pr = IncrementalRenPyAnalyzer().analyze(files=[src], graph=None)[0]
    return beats_from_parse_result(pr)


def _export(tmp_path, beats, statements, name='t.xlsx'):
    x = tmp_path / name
    export_beats_to_excel(beats, x, statements=statements)
    return x


def _header_cells(x):
    wb = openpyxl.load_workbook(x)
    try:
        ws = wb['Beats']
        return [c.value for c in ws[1]], ws
    finally:
        wb.close()


# ---------------------------------------------------------------------------
#  ① 没有声明项目的表：逐列不变
# ---------------------------------------------------------------------------

def test_table_without_declared_statements_is_unchanged(tmp_path):
    """项目没声明任何语句 → 表头**恰好**是基础列（一列都不许多）。"""
    _root, src = _project(tmp_path, rules=None)
    beats = _beats(src)
    x = tmp_path / 'no_spec.xlsx'
    # 不给 statements，且产物在项目外 → 反推不出项目根 → 不许乱加列
    export_beats_to_excel(beats, x)
    headers, _ws = _header_cells(x)
    from storygraph.story.beat_model import BEAT_HEADERS
    assert len(headers) == len(BEAT_FIELDS), headers
    assert headers[0] == BEAT_HEADERS['LineId']


def test_beat_to_row_default_stays_the_same_length(tmp_path):
    """老调用方（`zip(BEAT_FIELDS, beat.to_row())`）一个都不用改。"""
    _root, src = _project(tmp_path)
    for b in _beats(src):
        assert len(b.to_row()) == len(BEAT_FIELDS)
        assert len(b.to_row(spec_fields=['mycue·intensity'])) == \
            len(BEAT_FIELDS) + 1


# ---------------------------------------------------------------------------
#  ② 有声明：列 + 原生下拉
# ---------------------------------------------------------------------------

def test_declared_statement_gets_one_column_per_parameter(tmp_path):
    root, src = _project(tmp_path)
    beats = _beats(src)
    assert any(b.spec_name == 'mycue' for b in beats), (
        '夹具没被识别成声明语句，后面的判据全是空转')
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    assert 'mycue·intensity' in headers, headers
    assert 'mycue·target' in headers, headers
    # 列在**末尾**：基础列的下标不许被挤动（`BEAT_FIELDS.index` 满天飞）
    assert headers[:len(BEAT_FIELDS)] == [
        __import__('storygraph.story.beat_model', fromlist=['x'])
        .BEAT_HEADERS.get(f, f) for f in BEAT_FIELDS]
    # 值从**头行**解出来摊成列
    ic = headers.index('mycue·intensity') + 1
    tc = headers.index('mycue·target') + 1
    row = [b for b in beats if b.spec_name == 'mycue'][0]
    i = beats.index(row) + 2
    assert ws.cell(row=i, column=ic).value == '3'
    assert ws.cell(row=i, column=tc).value == 'ling'


def test_each_parameter_column_has_a_native_dropdown(tmp_path):
    """**原生下拉**：候选写进 Excel 数据验证，且手打新值不被拦。"""
    root, src = _project(tmp_path)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    ic = openpyxl.utils.get_column_letter(
        headers.index('mycue·intensity') + 1)
    hits = [dv for dv in ws.data_validations.dataValidation
            if any(str(rng).startswith(ic) for rng in dv.sqref.ranges)]
    assert hits, [dv.sqref for dv in ws.data_validations.dataValidation]
    dv = hits[0]
    assert dv.type == 'list', dv.type
    assert dv.formula1 == '"1,2,3"', dv.formula1
    assert dv.showErrorMessage is False, '下拉不该变成牢笼（手打要能用）'


def test_choices_too_long_fall_back_to_the_comment(tmp_path):
    """
    候选串超过 Excel 内联上限 → **不挂下拉**（退回顾问式注释），不崩、不坏文件。

    为什么这条要紧：内联候选是塞进公式字符串的，超长时 Excel 会直接
    拒收那条验证规则（严重时判整个文件损坏）——"少一个便利，不要坏一个文件"。
    """
    long_choices = [f'value_{i:03d}' for i in range(40)]
    rules = {'statements': [{'name': 'mycue', 'columns': [
        {'name': 'intensity', 'choices': long_choices}]}]}
    root, src = _project(tmp_path, rules=rules)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    assert 'mycue·intensity' in headers
    col = headers.index('mycue·intensity') + 1
    letter = openpyxl.utils.get_column_letter(col)
    assert not [dv for dv in ws.data_validations.dataValidation
                if any(str(r.sqref).startswith(letter)
                       for r in dv.sqref.ranges)], '超长候选还是挂了下拉'
    # 但"能填什么"必须说得出（注释兜底）
    annotated = [ws.cell(row=r, column=col).comment
                 for r in range(2, ws.max_row + 1)]
    assert any(c is not None and 'value_000' in str(c.text)
               for c in annotated), (
        [str(c.text) if c else None for c in annotated])


def test_comma_inside_choices_is_not_inlined(tmp_path):
    """候选里带逗号 → 内联表达会歧义 → 不挂下拉（不许写出错的下拉）。"""
    rules = {'statements': [{'name': 'mycue', 'columns': [
        {'name': 'target', 'choices': ['ling, lan', 'yu']}]}]}
    root, src = _project(tmp_path, rules=rules)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    letter = openpyxl.utils.get_column_letter(headers.index('mycue·target') + 1)
    assert not [dv for dv in ws.data_validations.dataValidation
                if any(str(r.sqref).startswith(letter)
                       for r in dv.sqref.ranges)]


# ---------------------------------------------------------------------------
#  ③ 改了那格 = 只改那个参数的值
# ---------------------------------------------------------------------------

def _edit_cell(x, header, value, row=0):
    wb = openpyxl.load_workbook(x)
    ws = wb['Beats']
    headers = [c.value for c in ws[1]]
    col = headers.index(header) + 1
    ws.cell(row=2 + row, column=col).value = value
    wb.save(x)


def _row_of_spec(beats, name='mycue'):
    """那条声明语句在第几行（0-based，**不含表头**）—— 别写死 0。"""
    return [i for i, b in enumerate(beats) if b.spec_name == name][0]


def _root_of(tmp_path):
    """夹具项目根（`_project` 造的那个）。"""
    return tmp_path / 'proj'


def _row_of_head(beats, head_text):
    """头行等于 `head_text` 的拍在第几行（0-based，不含表头）。"""
    for i, b in enumerate(beats):
        for attr in ('spec_head', 'loop_header', 'camera_header'):
            if str(getattr(b, attr, '') or '') == head_text:
                return i
    raise AssertionError(f'没找到头行为 {head_text!r} 的拍：'
                         f'{[(b.spec_head, b.loop_header, b.camera_header) for b in beats]}')


def _head_of(beats, attr, prefix=''):
    for b in beats:
        text = str(getattr(b, attr, '') or '')
        if text and (not prefix or prefix in text or text.startswith(prefix)):
            return text
    raise AssertionError(f'没有带 {attr} 的拍（prefix={prefix!r}）')


def _heads(beats, attr):
    """这一批拍里某个头行字段的**全部**取值（保序）—— 判"改了哪一条"要它。"""
    return [str(getattr(b, attr, '') or '')
            for b in beats if str(getattr(b, attr, '') or '')]


def test_editing_one_parameter_column_changes_only_that_value(tmp_path):
    root, src = _project(tmp_path)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    _edit_cell(x, 'mycue·intensity', '7', row=_row_of_spec(beats))

    back = import_beats_from_excel(x)
    heads = [b.spec_head for b in back if b.spec_name == 'mycue']
    assert heads and 'intensity=7' in heads[0], heads
    assert 'target=ling' in heads[0], heads
    out = render_beats_to_rpy(back, label_name='a')
    assert 'mycue "闪电" intensity=7 target=ling:' in out, out


def test_editing_a_parameter_keeps_the_quoted_form_of_the_others(tmp_path):
    """改 A 不许把 B 的引号吃掉（`target="ling"` 是作者的写法）。"""
    script = ('label a:\n'
              '    mycue "闪电" intensity=3 target="ling":\n'
              '        target = ling\n')
    root, src = _project(tmp_path, script=script)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    _edit_cell(x, 'mycue·intensity', '2', row=_row_of_spec(beats))
    back = import_beats_from_excel(x)
    head = [b.spec_head for b in back if b.spec_name == 'mycue'][0]
    assert head == 'mycue "闪电" intensity=2 target="ling":', head


def test_cleared_parameter_column_removes_the_parameter(tmp_path):
    """把那一格清空 = 别带这个参数（不是"写一个空值"）。"""
    root, src = _project(tmp_path)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    _edit_cell(x, 'mycue·intensity', '', row=_row_of_spec(beats))
    back = import_beats_from_excel(x)
    head = [b.spec_head for b in back if b.spec_name == 'mycue'][0]
    assert head == 'mycue "闪电" target=ling:', head


# ---------------------------------------------------------------------------
#  ④ 没改 = 逐字节原样
# ---------------------------------------------------------------------------

def test_untouched_table_round_trips_byte_for_byte(tmp_path):
    """
    打开表 → 什么都不改 → 再导出：**那一行逐字节不变**（含引号与缩进）。

    这条是"每参数一列"能不能加进来的**前提**：表里每一格都会把值写回来，
    一旦折叠时改写了写法，用户看到的就是"工具在改我的文件"。
    """
    script = ('label a:\n'
              '    mycue "闪电"  intensity = 3   target="ling" :\n'
              '        target = ling\n')
    root, src = _project(tmp_path, script=script)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    back = import_beats_from_excel(x)
    head = [b.spec_head for b in back if b.spec_name == 'mycue'][0]
    assert head == 'mycue "闪电"  intensity = 3   target="ling" :', head
    out = render_beats_to_rpy(back, label_name='a')
    assert 'mycue "闪电"  intensity = 3   target="ling" :' in out, out


def test_dirty_parameter_cell_is_sanitized_and_recorded(tmp_path):
    """
    参数格里混进换行 → 压成空格**并且留痕**（W114 的依据）。

    不留痕就是"静默改数据"：用户粘了个多行值，工具悄悄改了还不说。
    """
    root, src = _project(tmp_path)
    beats = _beats(src)
    reg = StatementRegistry.for_project(root)
    x = _export(tmp_path, beats, reg)
    _edit_cell(x, 'mycue·intensity', '7\n8', row=_row_of_spec(beats))
    back = import_beats_from_excel(x)
    b = [b for b in back if b.spec_name == 'mycue'][0]
    assert '\n' not in b.spec_head, repr(b.spec_head)
    assert 'mycue·intensity' in b.sanitized_cells, b.sanitized_cells
    # 压成空格之后**带空格的值必须加引号**（否则 `intensity=7 8` 会被
    # Ren'Py 当成两个位置参数 —— 那是改语义，不是改排版）
    assert 'intensity="7 8"' in b.spec_head, b.spec_head


@pytest.mark.parametrize('rules', [None, RULES])
def test_export_does_not_require_a_registry(tmp_path, rules):
    """不给注册表也得能导出（CLI / 老调用方走的就是这条路）。"""
    _root, src = _project(tmp_path, rules=rules)
    beats = _beats(src)
    x = tmp_path / 'plain.xlsx'
    export_beats_to_excel(beats, x)
    headers, _ws = _header_cells(x)
    assert len(headers) >= len(BEAT_FIELDS)


# ---------------------------------------------------------------------------
#  ⑤ **已识别语句**也要"每参数一列"（`camera` 的层名）
# ---------------------------------------------------------------------------
#
# 自定义语句的形状由项目声明（`rules.json` 的 `columns`）；`camera` 这类
# **标准语句**没人替它声明 —— 于是"每参数一列"只覆盖了一半。
# 补法只有两种：**猜**（拿"第几个词"去对参数名）或**我们自己声明一次**
# （`spec_args.BUILTIN_ARG_SCHEMAS`，按官方语法写死）。本仓库对猜一贯不收。
#
# ⚠ `camera` 的层名是**位置**参数：编辑它必须**替换第 0 个 token**
# （`camera background:` → `camera screens:`），写成 `layer=screens` 是**语法错**。
# 位置号因此进列名（`camera·layer@0`）—— 读表那一侧拿不到 schema，
# 只认列名才判得对。

CAMERA_SCRIPT = (
    'label a:\n'
    '    "前"\n'
    '    camera:\n'
    '        perspective True\n'
    '    camera background:\n'
    '        perspective True\n'
    '        xpos 0\n'
    '    "后"\n'
)


def _camera_beats(tmp_path):
    _root, src = _project(tmp_path, rules=None, script=CAMERA_SCRIPT)
    return _beats(src)


def test_recognized_statement_gets_a_positional_column(tmp_path):
    beats = _camera_beats(tmp_path)
    reg = StatementRegistry.for_project(_root_of(tmp_path))
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    assert 'camera·layer@0' in headers, headers
    col = headers.index('camera·layer@0') + 1
    rows = [ws.cell(row=r, column=col).value
            for r in range(2, ws.max_row + 1)]
    assert '' in rows or None in rows, rows          # `camera:` 没有层名
    assert 'background' in rows, rows                # `camera background:` 有


def test_positional_column_edit_replaces_the_nth_token(tmp_path):
    """改层名 = **替换第 0 个位置 token**，不许写成 `layer=…`（那是语法错）。"""
    beats = _camera_beats(tmp_path)
    reg = StatementRegistry.for_project(_root_of(tmp_path))
    x = _export(tmp_path, beats, reg)
    headers, _ws = _header_cells(x)
    assert 'camera·layer@0' in headers
    row = _row_of_head(beats, 'camera background:')
    _edit_cell(x, 'camera·layer@0', 'screens', row=row)

    back = import_beats_from_excel(x)
    heads = _heads(back, 'camera_header')
    assert heads == ['camera:', 'camera screens:'], heads
    assert 'layer=' not in ''.join(heads), heads
    out = render_beats_to_rpy(back, label_name='a')
    assert 'camera screens:' in out, out
    assert '        xpos 0' in out, out               # 块体一个字不动


def test_positional_column_cleared_removes_the_token(tmp_path):
    beats = _camera_beats(tmp_path)
    reg = StatementRegistry.for_project(_root_of(tmp_path))
    x = _export(tmp_path, beats, reg)
    row = _row_of_head(beats, 'camera background:')
    _edit_cell(x, 'camera·layer@0', '', row=row)
    back = import_beats_from_excel(x)
    heads = _heads(back, 'camera_header')
    assert heads == ['camera:', 'camera:'], heads


def test_untouched_positional_column_keeps_the_line_byte_for_byte(tmp_path):
    """没改 → 逐字节原样（含 `camera background:` 里那些空格）。"""
    script = ('label a:\n'
              '    camera   background  :\n'
              '        perspective True\n')
    _root, src = _project(tmp_path, rules=None, script=script)
    beats = _beats(src)
    reg = StatementRegistry.for_project(_root_of(tmp_path))
    x = _export(tmp_path, beats, reg)
    back = import_beats_from_excel(x)
    heads = [b.camera_header for b in back if b.camera_header]
    assert heads == ['camera   background  :'], heads


def test_positional_column_has_no_dropdown_when_choices_are_empty(tmp_path):
    """
    层名**没有候选**（运行期概念，项目里没人声明过有哪些层）→ 不挂下拉。

    编一组候选就是"诱导人写错"：本仓库对猜一贯不收。
    """
    beats = _camera_beats(tmp_path)
    reg = StatementRegistry.for_project(_root_of(tmp_path))
    x = _export(tmp_path, beats, reg)
    headers, ws = _header_cells(x)
    letter = openpyxl.utils.get_column_letter(headers.index('camera·layer@0') + 1)
    assert not [dv for dv in ws.data_validations.dataValidation
                if any(str(r).startswith(letter) for r in dv.sqref.ranges)]


# ---------------------------------------------------------------------------
#  ⑥ 叶语句（`with` / `pause` / `window`）：参数列折进**「代码」列那一行**
# ---------------------------------------------------------------------------
#
# 这三条走的是「代码」列（`Beat.raw` + `beat.ordered`），**不是** `opaque`。
# 为什么不把它们搬到 `opaque`（那样做参数列最简单）：搬过去就是**换渲染通道**
# —— 它们现在与 `$` 代码按 `ordered` **交错**输出（`show screen` / `pause` /
# `hide screen` / `$ …` 这种演出），换通道会动顺序，而顺序一错就是
# "相邻舞台指令过表翻序"（记忆缺口 4，静默改演出）。
#
# 所以参数列折的是**那一行**，而且必须**同步 `ordered`**
# （`_ordered_code_lines` 是"池行 == 记录行"才认，一处不同就整份退回历史顺序）。

LEAF_SCRIPT = (
    'label a:\n'
    '    scene bg room\n'
    '    with dissolve\n'
    '    pause 1.0\n'
    '    window hide\n'
    '    e "台词"\n'
)


def _leaf_beats(tmp_path):
    _root, src = _project(tmp_path, rules=None, script=LEAF_SCRIPT)
    return _beats(src)


def test_leaf_statements_get_one_column_per_parameter(tmp_path):
    beats = _leaf_beats(tmp_path)
    x = _export(tmp_path, beats, StatementRegistry.for_project(_root_of(tmp_path)))
    headers, ws = _header_cells(x)
    assert 'with·transition@0' in headers, headers
    assert 'pause·seconds@0' in headers, headers
    assert 'window·action@0' in headers, headers
    vals = {}
    for header in ('with·transition@0', 'pause·seconds@0', 'window·action@0'):
        col = headers.index(header) + 1
        vals[header] = [ws.cell(row=r, column=col).value
                        for r in range(2, ws.max_row + 1)]
    assert 'dissolve' in vals['with·transition@0'], vals
    assert '1.0' in vals['pause·seconds@0'], vals
    assert 'hide' in vals['window·action@0'], vals


def test_window_action_column_has_an_inline_dropdown(tmp_path):
    """`window` 的动作是**固定三个取值** → 内联下拉（不依赖资源表）。"""
    beats = _leaf_beats(tmp_path)
    x = _export(tmp_path, beats, StatementRegistry.for_project(_root_of(tmp_path)))
    headers, ws = _header_cells(x)
    letter = openpyxl.utils.get_column_letter(headers.index('window·action@0') + 1)
    hits = [dv for dv in ws.data_validations.dataValidation
            if any(str(r).startswith(letter) for r in dv.sqref.ranges)]
    assert hits, [str(dv.sqref) for dv in ws.data_validations.dataValidation]
    assert hits[0].formula1 == '"show,hide,auto"', hits[0].formula1


def test_editing_a_leaf_parameter_keeps_the_line_order(tmp_path):
    """
    **改一个叶语句参数，不许动任何一行的顺序**（这是本轮的核心判据）。

    夹具刻意做成"raw 与 code 交错"（`show screen` / `pause` / `hide screen` /
    `$ …`）—— 真机手机消息那一段就是这个形状。改 `pause` 的值之后，
    产物里这些行的**相对顺序必须一字不差**，而且逐行 diff 只允许一行不同。
    """
    script = ('label a:\n'
              '    show screen msg\n'
              '    pause 0.5\n'
              '    $ msg_add("一")\n'
              '    hide screen msg\n'
              '    $ msg_add("二")\n')
    _root, src = _project(tmp_path, rules=None, script=script)
    beats = _beats(src)
    x = _export(tmp_path, beats, StatementRegistry.for_project(_root_of(tmp_path)))

    def _product(bs):
        return render_beats_to_rpy(bs, label_name='a')

    baseline = _product(import_beats_from_excel(x))
    order = [ln.strip() for ln in baseline.split('\n') if ln.strip()]

    headers, _ws = _header_cells(x)
    row = [i for i, b in enumerate(beats) if 'pause' in (b.raw or '')][0]
    _edit_cell(x, 'pause·seconds@0', '1.5', row=row)
    changed = _product(import_beats_from_excel(x))

    assert 'pause 1.5' in changed, changed
    new_order = [ln.strip() for ln in changed.split('\n') if ln.strip()]
    assert new_order == [t if t != 'pause 0.5' else 'pause 1.5' for t in order], (
        f'顺序被改了：\n--- 改前 ---\n{baseline}\n--- 改后 ---\n{changed}')
    diff = [i for i, (a, b) in enumerate(zip(baseline.split('\n'),
                                             changed.split('\n'))) if a != b]
    assert len(diff) == 1, [baseline.split('\n')[i] for i in diff]


def test_untouched_leaf_columns_round_trip_byte_for_byte(tmp_path):
    beats = _leaf_beats(tmp_path)
    x = _export(tmp_path, beats, StatementRegistry.for_project(_root_of(tmp_path)))
    back = import_beats_from_excel(x)
    out = render_beats_to_rpy(back, label_name='a')
    src_text = Path(_root_of(tmp_path) / 'game' / 'ch.rpy').read_text(
        encoding='utf-8')
    for ln in src_text.split('\n'):
        if ln.strip() and not ln.strip().startswith('#'):
            assert ln in out, f'没改却对不上：{ln!r}\n{out}'


def test_two_identical_leaf_statements_in_one_beat_are_not_guessed(tmp_path):
    """
    一拍里**两条同类**叶语句（两个 `pause`）→ 那一格**留空**，改它也不生效。

    一格装不下两条：挑一条改就是**改了另一条指令**。宁可这一格什么都不做
    （用户想改可以直接改「代码」列）。
    """
    script = ('label a:\n'
              '    pause 0.5\n'
              '    pause 1.0\n'
              '    e "台词"\n')
    _root, src = _project(tmp_path, rules=None, script=script)
    beats = _beats(src)
    assert len(beats) == 1, [(b.raw, b.text) for b in beats]
    x = _export(tmp_path, beats, StatementRegistry.for_project(_root_of(tmp_path)))
    headers, ws = _header_cells(x)
    col = headers.index('pause·seconds@0') + 1
    assert not (ws.cell(row=2, column=col).value or ''), (
        ws.cell(row=2, column=col).value)

    _edit_cell(x, 'pause·seconds@0', '9.9', row=0)
    back = import_beats_from_excel(x)
    assert 'pause 0.5' in (back[0].raw or ''), back[0].raw
    assert 'pause 1.0' in (back[0].raw or ''), back[0].raw
    assert '9.9' not in (back[0].raw or ''), back[0].raw
