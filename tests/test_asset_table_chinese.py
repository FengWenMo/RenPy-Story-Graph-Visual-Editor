"""
资源表**中英双认**：表头改成中文，但读写与老表都不许出问题
==========================================================

用户 2026-10-07："资源配置表…建议改为中文"。改表头是**契约级**动作：
仓库里 `parse_asset_tables` / 改名 / 删行 / 搜索 / 合并**全按 `row['ID']`、
`row['Path']` 取值**。所以做法是**只在读写边界翻译**：

    · 写：表头写中文（`ASSET_FIELD_LABELS`），取值仍用规范键；
    · 读：中文表头映射回规范键（`_canon_field` / `_sheet_headers`）；
    · 结果：**老表（英文表头）继续能读能改**，用户手上的表不用改名。

这个文件钉四条：新表是中文 / 读回来仍是规范键 / 老表照样能读能改 /
改列名之后**编辑类操作（改路径、按 ID 删行）仍然找得到列**
（表头一变，`'Path' not in headers` 会静默不生效 —— 那是最阴的后果）。
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl                                            # noqa: E402

from storygraph.assets.asset_registry import (             # noqa: E402
    ASSET_FIELD_LABELS, _canon_field, _sheet_headers,
    delete_rows_by_id, parse_asset_tables, read_excel_rows,
    update_table_paths, write_dict_to_excel,
)


def _table(tmp_path, headers, rows, name='images.xlsx') -> Path:
    d = tmp_path / 'asset_tables'
    d.mkdir(parents=True, exist_ok=True)
    p = d / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(p)
    return p


def test_written_headers_are_chinese(tmp_path):
    """新表：表头是中文，但**数据语义一个字节没变**（仍是同一批键值）。"""
    p = tmp_path / 'img.xlsx'
    write_dict_to_excel(p, {'ling_happy': {'ID': 'ling_happy',
                                           'Path': 'images/ling_happy.png',
                                           'Tag': 'ling',
                                           'Attributes': 'happy',
                                           'Category': 'character',
                                           'Notes': ''}},
                        ['ID', 'Path', 'Tag', 'Attributes',
                         'Category', 'Notes'])
    wb = openpyxl.load_workbook(p)
    try:
        raw = [c.value for c in wb.active[1]]
    finally:
        wb.close()
    assert raw == ['编号', '真实文件', '角色标签', '表情', '类别', '备注'], raw
    # **`编号`(ID) 整列隐藏**（用户 2026-10-07 定案：它就是那个"本该藏起来、
    # 没什么用的内部编号"；藏起来 = 与"换位往前挪"同样效果、零风险）。
    wb2 = openpyxl.load_workbook(p)
    try:
        assert wb2.active.column_dimensions['A'].hidden is True or \
            wb2.active.column_dimensions['A'].hidden is True
    finally:
        wb2.close()
    # 隐藏只是观感：**读得到**，改名/删行/搜索全不受影响
    rows = read_excel_rows(p)
    assert rows[0]['ID'] == 'ling_happy'
    assert rows[0]['Path'] == 'images/ling_happy.png'
    assert ASSET_FIELD_LABELS['Path'] == '真实文件'


def test_read_maps_back_to_canonical_keys(tmp_path):
    """读：中文表头 → 规范键（下游 `row['ID']` 那套代码一行都不用改）。"""
    p = _table(tmp_path,
               ['编号', '真实文件', '角色标签', '表情', '类别', '备注'],
               [['ling_happy', 'images/ling_happy.png', 'ling', 'happy',
                 'character', '']])
    rows = read_excel_rows(p)
    assert rows[0]['ID'] == 'ling_happy'
    assert rows[0]['Path'] == 'images/ling_happy.png'
    assert rows[0]['Tag'] == 'ling'


def test_old_english_table_still_reads_and_updates(tmp_path):
    """老表（英文表头）继续能读、能改 —— 用户手上的表不用改名。"""
    p = _table(tmp_path,
               ['ID', 'Path', 'Tag', 'Attributes', 'Category', 'Notes'],
               [['ling_happy', 'images/old.png', 'ling', 'happy',
                 'character', '人工备注']])
    assert read_excel_rows(p)[0]['ID'] == 'ling_happy'

    # 改路径：这条路按表头找 `Path` 列 —— 英文表头必须照样找得到。
    # ⚠ `updates` 是 **ID → 新 Path**（不是 path → path）——第一版传错，返回 0。
    assert update_table_paths(tmp_path / 'asset_tables', 'images',
                              {'ling_happy': 'images/new.png'}) == 1
    assert read_excel_rows(p)[0]['Path'] == 'images/new.png'
    # 人工备注不许被顺手清掉
    assert read_excel_rows(p)[0]['Notes'] == '人工备注'


def test_editing_still_finds_columns_after_headers_are_chinese(tmp_path):
    """
    **最阴的后果**：表头变中文之后，`'Path' not in headers` 会让改名/删行
    **静默不生效**（不报错、什么都没做）。这条判据盯着它。
    """
    p = _table(tmp_path, ['编号', '真实文件', '音效参数'], [])
    wb = openpyxl.load_workbook(p)
    try:
        ws = wb.active
        assert _sheet_headers(ws)[:2] == ['ID', 'Path']
        assert _canon_field('真实文件') == 'Path'
        assert _canon_field('用户自己加的列') == '用户自己加的列'   # 不认识的别改
    finally:
        wb.close()

    _table(tmp_path, ['编号', '真实文件'], [['a1', 'sfx/a.wav']],
           name='sounds.xlsx')
    d = tmp_path / 'asset_tables'
    assert update_table_paths(d, 'sounds', {'a1': 'sfx/b.wav'}) == 1
    assert delete_rows_by_id(d, 'sounds', ['a1']) == 1


def test_swapped_header_order_is_safe_to_append(tmp_path):
    """
    **列顺序换位必须是安全的**（用户 2026-10-07 要"与内容编号换位置往前挪"）。

    第十五轮试过直接改表头顺序 → 4 条判据红（占位行的 ID 被读成路径）：
    根因是"**按行追加的写入点**固定按某个顺序写**值**"，头一换序就错位。
    正确前提是**追加按文件自己的表头逐列取值**（`append_rows_to_table`
    已经是这么写的）。这条判据把"换位安全"钉住：即使表头顺序是
    `真实文件 | 编号`，追加进的值也必须落在**名字**对应的列下。
    """
    d = tmp_path / 'asset_tables'
    d.mkdir(parents=True, exist_ok=True)
    p = d / 'sounds.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['真实文件', '编号', '备注'])        # ← 故意换过序
    wb.save(p)

    from storygraph.assets.asset_registry import append_rows_to_table
    append_rows_to_table(d, 'sounds', [{'ID': 'a1', 'Path': 'sfx/a.wav',
                                        'Notes': '备注'}])
    rows = read_excel_rows(p)
    assert rows[0]['ID'] == 'a1', rows
    assert rows[0]['Path'] == 'sfx/a.wav', rows
    assert rows[0]['Notes'] == '备注', rows


def test_parse_asset_tables_reads_chinese_headers(tmp_path):
    """解析器那条路（注册表）也要认中文表头。"""
    _table(tmp_path, ['编号', '真实文件', '角色标签'],
           [['ling_happy', 'images/ling_happy.png', 'ling']])
    reg = parse_asset_tables(tmp_path / 'asset_tables')
    assert 'ling_happy' in reg.images, list(reg.images)
