"""
判据 · **E2 候选接入编辑器**（2026-10-10）
==========================================

E2 的一句话目标：**"下拉里放什么"由引擎决定，人一个都不用填。**

所以这里要钉住三件事：

1. **引擎优先**：采过 `engine_facts.json` → 分组与值都从引擎来
   （组 = `layeredimage` 的 group，值 = 引擎枚举的属性）；
2. **定义表降级为覆盖层**：项目声明的 `attr_groups` **只用来改中文名**
   （`outfit` → 「服装」），不再指定"哪些词属于哪组"；
3. **没采过就一字不变**：`facts` 为空 → 走原来的老路（这是通用性底线，
   也是"没采到 ≠ 引擎什么都没给"那条规矩的落地）。

另外带上 **N3**（同组冲突工具自查）与 **N2/N5**（覆盖内置语句要报）。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.app import beats                                      # noqa: E402
from storygraph.app.session_graph import GraphSessionMixin            # noqa: E402
from storygraph.core import engine_facts as EF                        # noqa: E402


def _facts(attrs, candidates):
    """
    一份**可用**的引擎事实（`collection` 两个闸门都为真）。

    ⚠ 为什么必须带 `collection`：`unrecognized_attrs` 只在事实**可用**时才下结论 ——
    这是"采不到就不宣称"那条规矩的落点（第一次写这批用例时正因漏了它而红，
    是**判据自己错了**，不是实现错了）。
    """
    return EF.build_facts({
        'attributes': [(EF.SRC_PROBE, attrs)],
        'candidates': [(EF.SRC_PROBE, candidates)],
    }, collection={'parse_clean': True, 'probe_complete': True})


LAN_ATTRS = {
    'lan': [{'group': 'expressions', 'attribute': 'happy'},
            {'group': 'expressions', 'attribute': 'sad'},
            {'group': 'outfit', 'attribute': 'clothes2'},
            {'group': 'outfit', 'attribute': 'base_hand_up'}],
}
#: 引擎给的顺序就是"适合展示"的顺序 —— 注意 `outfit` 的词排在 `sad` 前面
LAN_CANDS = {'lan': ['happy', 'clothes2', 'base_hand_up', 'sad', 'zero_shot']}


class AttrGroup:
    """项目声明的词组（与 `authoring.spec_groups.AttrGroup` 同形）。"""

    def __init__(self, name, label, options=()):
        self.name, self.label, self.options = name, label, list(options)

    def title(self):
        return self.label or self.name or '（未命名）'


_Group = AttrGroup          # 用例里的旧名字，保持可读


# ---------------------------------------------------------------------------
#  纯层：分组候选
# ---------------------------------------------------------------------------

class TestGroupedChoices:
    def test_groups_come_from_engine(self):
        got = EF.grouped_choices(_facts(LAN_ATTRS, LAN_CANDS), 'lan')
        names = [g['group'] for g in got]
        assert names[:2] == ['expressions', 'outfit'], got

    def test_values_keep_engine_order(self):
        got = EF.grouped_choices(_facts(LAN_ATTRS, LAN_CANDS), 'lan')
        outfit = next(g for g in got if g['group'] == 'outfit')
        assert outfit['values'] == ['clothes2', 'base_hand_up']

    def test_untagged_candidates_are_not_lost(self):
        """引擎有、但没归到组的候选 → 兜底一栏，**绝不丢词**。"""
        got = EF.grouped_choices(_facts(LAN_ATTRS, LAN_CANDS), 'lan')
        rest = [g for g in got if g['group'] == '']
        assert rest and rest[0]['values'] == ['zero_shot']

    def test_rename_only_changes_the_label(self):
        got = EF.grouped_choices(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                 {'outfit': '服装'})
        outfit = next(g for g in got if g['group'] == 'outfit')
        assert outfit['label'] == '服装'
        assert outfit['values'] == ['clothes2', 'base_hand_up']   # 结构没被改

    def test_unknown_tag_is_empty(self):
        assert EF.grouped_choices(_facts(LAN_ATTRS, LAN_CANDS), 'nobody') == []


class TestRenameOverlay:
    def test_reads_objects_and_dicts(self):
        got = EF.rename_overlay([_Group('outfit', '服装'),
                                 {'name': 'expressions', 'label': '表情'},
                                 {'name': 'x', 'label': ''}])
        assert got == {'outfit': '服装', 'expressions': '表情'}

    def test_options_are_ignored(self):
        """归并语义不再走这条路 —— 结构已经是引擎的事实。"""
        got = EF.rename_overlay([_Group('outfit', '服装', ['swim_*'])])
        assert got == {'outfit': '服装'}


# ---------------------------------------------------------------------------
#  纯层：N3 同组冲突
# ---------------------------------------------------------------------------

class TestGroupConflicts:
    def test_two_in_one_group_conflicts(self):
        got = EF.group_conflicts(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                 ['happy', 'sad'])
        assert got == [{'group': 'expressions', 'attrs': ['happy', 'sad']}]

    def test_different_groups_are_fine(self):
        assert EF.group_conflicts(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                  ['happy', 'clothes2']) == []

    def test_single_attr_is_fine(self):
        assert EF.group_conflicts(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                  ['happy']) == []

    def test_unknown_tag_is_empty(self):
        assert EF.group_conflicts(_facts(LAN_ATTRS, LAN_CANDS), 'nobody',
                                  ['a', 'b']) == []


# ---------------------------------------------------------------------------
#  app 层：引擎优先
# ---------------------------------------------------------------------------

class TestExpressionGroupsFor:
    def test_engine_first_does_not_need_symbols(self):
        """引擎事实给了，就不该再要求符号表 —— 编剧一个都不用填。"""
        got = beats.expression_groups_for('lan', None, [], _facts(LAN_ATTRS,
                                                                 LAN_CANDS))
        assert [g['group'] for g in got][:2] == ['expressions', 'outfit']

    def test_engine_first_ignores_declared_options(self):
        """项目声明里的 `options`（旧机制）不再决定结构，只改说法。"""
        got = beats.expression_groups_for(
            'lan', None, [_Group('outfit', '服装', ['swim_*'])],
            _facts(LAN_ATTRS, LAN_CANDS))
        outfit = next(g for g in got if g['group'] == 'outfit')
        assert outfit['label'] == '服装'
        assert outfit['values'] == ['clothes2', 'base_hand_up']

    def test_without_facts_the_old_path_is_untouched(self):
        """**通用性底线**：没采过 → 原来的行为一字不变（空声明仍返回 `[]`）。"""
        assert beats.expression_groups_for('lan', None, [], None) == []

    def test_falls_back_when_engine_has_no_groups_for_tag(self):
        facts = _facts({'other': [{'group': 'g', 'attribute': 'x'}]},
                       {'other': ['x']})
        # 引擎没给 lan 的分组 → 走项目声明（这里是空）→ 空
        assert beats.expression_groups_for('lan', None, [], facts) == []


class _Sym:
    """最小符号表：只提供 `attrs_of_tag`。"""

    def __init__(self, table):
        self._t = table

    def attrs_of_tag(self, tag):
        return list(self._t.get(tag) or [])


class TestOldPathStillWorks:
    def test_declared_groups_still_group(self):
        sym = _Sym({'lan': ['happy', 'sad', 'clothes2']})
        got = beats.expression_groups_for('lan', sym, [_Group('outfit', '服装',
                                                             ['clothes2'])])
        labels = [g['label'] for g in got]
        assert '服装' in labels


class TestBeatGroupConflicts:
    def _beat(self, show):
        b = beats.Beat()
        b.show = show
        return b

    def test_detects_same_group_selection(self):
        """实测 lint 对这类冲突**一声不吭**（`show lan happy sad` → exit 0），
        所以必须由工具自己算。"""
        got = beats.beat_group_conflicts(self._beat('lan happy sad'),
                                         _facts(LAN_ATTRS, LAN_CANDS))
        assert got and got[0]['tag'] == 'lan'
        assert set(got[0]['attrs']) == {'happy', 'sad'}

    def test_no_conflict_across_groups(self):
        assert beats.beat_group_conflicts(self._beat('lan happy clothes2'),
                                          _facts(LAN_ATTRS, LAN_CANDS)) == []

    def test_without_facts_nothing_is_claimed(self):
        assert beats.beat_group_conflicts(self._beat('lan happy sad'), {}) == []


# ---------------------------------------------------------------------------
#  会话：引擎事实的读法（含"采不到 ≠ 没有"）
# ---------------------------------------------------------------------------

class _Stub(GraphSessionMixin):
    """
    只借 `GraphSessionMixin` 的这几个方法 —— 它们只用到 `project_root`。
    继承（而不是鸭子类型）是因为 `engine_facts_warnings` 内部会调 `self.engine_facts()`。
    """

    def __init__(self, root):
        self.project_root = root


def _write_facts(root, **collection):
    facts = EF.build_facts({'images': [(EF.SRC_PROBE, ['lan'])]},
                           collection={'parse_clean': True,
                                       'probe_complete': True,
                                       **collection})
    EF.facts_path(root).write_text(EF.dumps(facts), encoding='utf-8')
    return facts


class TestCardNote:
    """卡片底部那条提示 —— **文案由语义层拼**（前端只显示），所以这里能纯测。"""

    def _beat(self, show, speaker='lyj', text=''):
        b = beats.Beat()
        b.show = show
        b.speaker = speaker
        b.text = text
        return b

    def test_same_group_conflict_is_in_the_note(self):
        note = beats.card_note(self._beat('lan happy sad'), None,
                               _facts(LAN_ATTRS, LAN_CANDS))
        assert 'expressions' in note
        assert 'happy' in note and 'sad' in note

    def test_no_conflict_no_group_mention(self):
        note = beats.card_note(self._beat('lan happy clothes2'), None,
                               _facts(LAN_ATTRS, LAN_CANDS))
        assert 'expressions' not in note and 'outfit' not in note

    def test_without_facts_nothing_is_claimed(self):
        note = beats.card_note(self._beat('lan happy sad'), None, None)
        assert '组里同时选了' not in note

    def test_parts_are_joined_with_a_separator(self):
        """多句时用 `；` 连成一行（`set_note` 是一行 QLabel）。"""
        note = beats.card_note(self._beat('lan happy sad'), None,
                               _facts(LAN_ATTRS, LAN_CANDS))
        assert '\n' not in note


class TestUnrecognizedAttrs:
    """**只标注、不过滤**（用户 2026-10-11 定的口径）。"""

    def test_flags_only_the_unknown_one(self):
        got = EF.unrecognized_attrs(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                    ['happy', 'lan_angry'])
        assert got == ['lan_angry']

    def test_known_attrs_are_not_flagged(self):
        assert EF.unrecognized_attrs(_facts(LAN_ATTRS, LAN_CANDS), 'lan',
                                     ['happy', 'clothes2']) == []

    def test_no_facts_means_no_claim(self):
        assert EF.unrecognized_attrs(None, 'lan', ['whatever']) == []
        assert EF.unrecognized_attrs({}, 'lan', ['whatever']) == []

    def test_uncollectable_facts_mean_no_claim(self):
        """坑 4：采集不可用的事实不能拿来下结论。"""
        bad = EF.build_facts({'attributes': [(EF.SRC_PROBE, LAN_ATTRS)]},
                             collection={'parse_clean': True,
                                         'probe_complete': False})
        assert EF.unrecognized_attrs(bad, 'lan', ['lan_angry']) == []

    def test_engine_silent_about_tag_means_no_claim(self):
        """引擎对这个 tag 什么都没说 → **不下结论**（否则非 layeredimage 的 tag 全被误标）。"""
        facts = _facts({'other': [{'group': 'g', 'attribute': 'x'}]},
                       {'other': ['x']})
        assert EF.unrecognized_attrs(facts, 'nobody', ['whatever']) == []


class TestBeatUnknownAttrs:
    def _beat(self, show):
        b = beats.Beat()
        b.show = show
        return b

    def test_reports_only_the_unknown(self):
        got = beats.beat_unknown_attrs(self._beat('lan happy lan_angry'),
                                       _facts(LAN_ATTRS, LAN_CANDS))
        assert got == [{'tag': 'lan', 'attr': 'lan_angry'}]

    def test_at_transform_is_not_an_attribute(self):
        """`close_happy@i11` 里的 `i11` 是位置变换，不该被当成属性。"""
        got = beats.beat_unknown_attrs(self._beat('lan happy@i11'),
                                       _facts(LAN_ATTRS, LAN_CANDS))
        assert got == []

    def test_no_facts_no_claim(self):
        assert beats.beat_unknown_attrs(self._beat('lan lan_angry'), None) == []


class TestExcelAnnotation:
    """剧本表**立绘格挂批注**：只提示，不改数据、不改下拉、不动资源表。"""

    def _export(self, tmp_path, facts, show='lan happy lan_angry'):
        from openpyxl import load_workbook
        from storygraph.story.beat_model import Beat
        from storygraph.story.excel_export import BEAT_FIELDS, export_beats_to_excel

        b = Beat()
        b.show = show
        out = tmp_path / 't.xlsx'
        export_beats_to_excel([b], out, facts=facts)
        ws = load_workbook(out).active
        col = BEAT_FIELDS.index('Show') + 1
        return ws.cell(row=2, column=col)

    def test_comment_is_attached_when_something_is_unknown(self, tmp_path):
        cell = self._export(tmp_path, _facts(LAN_ATTRS, LAN_CANDS))
        assert cell.comment is not None
        text = cell.comment.text
        assert 'lan_angry' in text
        assert '不拦你' in text                # 明说"只是提示"

    def test_cell_value_is_not_modified(self, tmp_path):
        """**批注不改数据** —— 候选与导入仍以实际存在的资源为准。"""
        cell = self._export(tmp_path, _facts(LAN_ATTRS, LAN_CANDS))
        assert cell.value == 'lan happy lan_angry'

    def test_no_facts_no_comment(self, tmp_path):
        """没采过引擎事实 → 一个批注都不挂（行为与从前一字不变）。"""
        cell = self._export(tmp_path, None)
        assert cell.comment is None

    def test_show_keywords_are_not_attributes(self, tmp_path):
        """`show lan at center` 里的 `at` 是**关键字**，不是属性（实测踩到过）。"""
        cell = self._export(tmp_path, _facts(LAN_ATTRS, LAN_CANDS),
                            show='lan happy at center')
        assert cell.comment is None

    def test_empty_cell_gets_no_comment(self, tmp_path):
        """⚠ 批注必须描述**它所在的那一格** —— 格子是空的就别说话。"""
        cell = self._export(tmp_path, _facts(LAN_ATTRS, LAN_CANDS), show='')
        assert cell.comment is None

    def test_asset_tables_are_not_touched_at_all(self, tmp_path):
        """
        ⚠⚠ "提示放哪"那张表的**第一行**：**资源表一字不动**。

        资源表 (images.xlsx / Candidates_*) 是**人工维护的资源登记**，
        机器不往里写语义判断、也不加会漂移的列；候选表的下拉必须全量
        （"一切以实际有的资源为准"）。所以导出剧本表时，
        `game/asset_tables/` 下**一个字节都不许变**。

        做法：先放一份"资源表"，记下**逐字节指纹**，导出后再比。
        """
        import hashlib
        from storygraph.assets.asset_registry import _TABLE_FIELDS, write_dict_to_excel
        from storygraph.story.beat_model import Beat
        from storygraph.story.excel_export import export_beats_to_excel

        adir = tmp_path / 'game' / 'asset_tables'
        adir.mkdir(parents=True)
        write_dict_to_excel(adir / 'images.xlsx', {
            'lan_angry': {'ID': 'lan_angry', 'Path': 'game/images/lan_angry.png',
                          'Tag': 'lan', 'Attributes': 'lan_angry',
                          'Category': 'character'},
        }, _TABLE_FIELDS['image'])
        write_dict_to_excel(adir / 'Candidates_角色.xlsx',
                            {'lan': {'ID': 'lan', '候选': 'happy,sad'}},
                            ['ID', '候选'])

        def snap():
            return {p.name: hashlib.sha1(p.read_bytes()).hexdigest()
                    for p in sorted(adir.glob('*.xlsx'))}

        before = snap()
        b = Beat()
        b.show = 'lan happy lan_angry'
        export_beats_to_excel([b], tmp_path / 't.xlsx',
                              facts=_facts(LAN_ATTRS, LAN_CANDS))
        after = snap()
        assert before == after, '资源表 / 候选表被改动了 —— 那是红线'
        assert len(before) == 2, before


class TestTableVsEngine:
    """资源表 ↔ 引擎 的对账：**只报告，不删**（路线图 §8.9）。"""

    def test_table_only_is_reported(self):
        got = EF.table_vs_engine(_facts(LAN_ATTRS, LAN_CANDS),
                                 {'lan': ['happy', 'mask', '_ss_cut']})
        assert got['table_only'] == {'lan': ['mask', '_ss_cut']}

    def test_engine_only_is_reported(self):
        got = EF.table_vs_engine(_facts(LAN_ATTRS, LAN_CANDS),
                                 {'lan': ['happy']})
        assert 'clothes2' in got['engine_only']['lan']

    def test_no_facts_no_claim(self):
        assert EF.table_vs_engine(None, {'lan': ['x']}) == {}
        assert EF.table_vs_engine({}, {'lan': ['x']}) == {}

    def test_uncollectable_facts_no_claim(self):
        bad = EF.build_facts({'attributes': [(EF.SRC_PROBE, LAN_ATTRS)]},
                             collection={'parse_clean': True,
                                         'probe_complete': False})
        assert EF.table_vs_engine(bad, {'lan': ['x']}) == {}

    def test_tag_the_engine_says_nothing_about_is_skipped(self):
        got = EF.table_vs_engine(_facts(LAN_ATTRS, LAN_CANDS),
                                 {'nobody': ['a', 'b']})
        assert got == {'table_only': {}, 'engine_only': {},
                       'spelling_only': {}}

    def test_spelling_differences_are_not_counted_as_differences(self):
        """真机逼出来的那一桶：`bg1.1` vs `bg1_1` **不算差异**（但也不丢）。

        ⚠ 桶里**两边都留**（表里的 `bg1.1` 与引擎里的 `bg1_1`）——
        这样人一眼能看到**配对关系**；只留一边就看不出跟谁重了。
        """
        got = EF.table_vs_engine(
            _facts({'bg': [{'group': 'g', 'attribute': 'bg1_1'}]},
                   {'bg': ['bg1_1']}),
            {'bg': ['bg1.1']})
        assert got['table_only'] == {} and got['engine_only'] == {}
        assert sorted(got['spelling_only']['bg']) == ['bg1.1', 'bg1_1']

    def test_spelling_bucket_never_hides_a_real_difference(self):
        """写法归一**只用来消除假警报**，不许把真差异吞掉（契约 T1 的同款纪律）。"""
        got = EF.table_vs_engine(
            _facts({'bg': [{'group': 'g', 'attribute': 'bg1_1'}]},
                   {'bg': ['bg1_1']}),
            {'bg': ['bg1.1', 'mask']})
        assert got['table_only'] == {'bg': ['mask']}
        assert sorted(got['spelling_only']['bg']) == ['bg1.1', 'bg1_1']


class TestDiffFacts:
    """漂移报告（E3）：纯函数、不需要引擎 —— CI 里直接跑。"""

    def _f(self, **kw):
        payload = {k: [(EF.SRC_PROBE, v)] for k, v in kw.items()}
        return EF.build_facts(payload, collection={'parse_clean': True,
                                                   'probe_complete': True})

    def test_identical_is_same(self):
        a = self._f(labels=['start', 'end'])
        assert EF.diff_facts(a, a)['same'] is True

    def test_added_label(self):
        d = EF.diff_facts(self._f(labels=['start']),
                          self._f(labels=['start', 'end']))
        assert d['same'] is False
        assert d['added'] == {'labels': ['end']}

    def test_removed_transform(self):
        d = EF.diff_facts(self._f(transforms=['fade', 'dissolve']),
                          self._f(transforms=['fade']))
        assert d['removed'] == {'transforms': ['dissolve']}

    def test_nested_attribute_drift_is_per_tag(self):
        """映射型事实按 tag 展开 —— 报告能精确到"哪个角色多了哪个属性"。"""
        d = EF.diff_facts(
            self._f(attributes={'lan': [{'group': 'g', 'attribute': 'happy'}]}),
            self._f(attributes={'lan': [{'group': 'g', 'attribute': 'happy'},
                                        {'group': 'g', 'attribute': 'sad'}]}))
        assert 'attributes.lan' in d['added']

    def test_engine_version_change_counts_as_drift(self):
        a = EF.build_facts({}, engine={'version': '8.4.1'},
                           collection={'parse_clean': True,
                                       'probe_complete': True})
        b = EF.build_facts({}, engine={'version': '8.5.3'},
                           collection={'parse_clean': True,
                                       'probe_complete': True})
        d = EF.diff_facts(a, b)
        assert d['same'] is False
        assert d['engine']['version'] == ('8.4.1', '8.5.3')


class TestEngineAlignSection:
    """体检报告里那一段（有才对；没有就整段不出现）。"""

    def _report(self, align):
        return {'ok': True, 'problems': 0, 'counts': {'images': 3},
                'missing': {}, 'unregistered': {}, 'outside_game': {},
                'duplicate_ids': [], 'duplicate_paths': [], 'no_path': [],
                'placeholders': [], 'statements': [], 'cds': [],
                'engine_align': align}

    def test_section_appears_with_data(self):
        from storygraph.app.asset_report import format_asset_health
        text = format_asset_health(self._report(
            {'table_only': {'lan': ['mask', '_ss_cut']}, 'engine_only': {}}))
        assert '⑨ 资源表 ↔ 引擎 对账' in text
        assert 'mask' in text
        assert '不影响导入' in text          # 明说"不拦你"

    def test_section_absent_without_data(self):
        from storygraph.app.asset_report import format_asset_health
        assert '⑨' not in format_asset_health(self._report({}))

    def test_section_absent_when_nothing_to_say(self):
        from storygraph.app.asset_report import format_asset_health
        assert '⑨' not in format_asset_health(
            self._report({'table_only': {}, 'engine_only': {},
                          'spelling_only': {}}))

    def test_spelling_bucket_is_a_count_not_a_list(self):
        """写法差异只出一行计数 —— 不摊开成清单（否则又变成噪声）。"""
        from storygraph.app.asset_report import format_asset_health
        text = format_asset_health(self._report(
            {'table_only': {}, 'engine_only': {},
             'spelling_only': {'bg': ['bg1.1', 'bg1.2']}}))
        assert '只是写法不同' in text and '2 项' in text
        # 摊开成清单的样子是 `bg: bg1.1、bg1.2` —— 这种行不该出现
        assert 'bg: bg1.1' not in text


class TestCliDriftGate:
    """E3 门禁：`engine-facts --check-only` 的**退出码约定**（CI 全靠它）。"""

    def _facts_file(self, path, **kw):
        payload = {k: [(EF.SRC_PROBE, v)] for k, v in kw.items()}
        facts = EF.build_facts(payload, engine={'version': '8.5.3'},
                               collection={'parse_clean': True,
                                           'probe_complete': True})
        path.write_text(EF.dumps(facts), encoding='utf-8')
        return facts

    def _run(self, argv):
        from storygraph.cli import main
        try:
            return main(argv)
        except SystemExit as e:                     # pragma: no cover
            return int(e.code or 0)

    def test_no_drift_exits_zero(self, tmp_path, capsys):
        cur = tmp_path / 'cur.json'
        base = tmp_path / 'base.json'
        self._facts_file(cur, labels=['start'])
        self._facts_file(base, labels=['start'])
        code = self._run(['engine-facts', '--check-only',
                          '--project', str(tmp_path),
                          '--out', str(cur), '--baseline', str(base)])
        assert code == 0
        assert '没有变化' in capsys.readouterr().out

    def test_drift_exits_one_and_prints_the_change(self, tmp_path, capsys):
        cur = tmp_path / 'cur.json'
        base = tmp_path / 'base.json'
        self._facts_file(base, labels=['start'])
        self._facts_file(cur, labels=['start', 'ending'])
        code = self._run(['engine-facts', '--check-only',
                          '--project', str(tmp_path),
                          '--out', str(cur), '--baseline', str(base)])
        assert code == 1
        out = capsys.readouterr().out
        assert 'ending' in out and '漂移' in out

    def test_missing_baseline_is_blocking(self, tmp_path, capsys):
        cur = tmp_path / 'cur.json'
        self._facts_file(cur, labels=['start'])
        code = self._run(['engine-facts', '--check-only',
                          '--project', str(tmp_path), '--out', str(cur)])
        assert code == 2
        assert 'baseline' in capsys.readouterr().err

    def test_broken_file_is_blocking(self, tmp_path, capsys):
        cur = tmp_path / 'cur.json'
        base = tmp_path / 'base.json'
        cur.write_text('{ 不是 json', encoding='utf-8')
        self._facts_file(base, labels=['start'])
        code = self._run(['engine-facts', '--check-only',
                          '--project', str(tmp_path),
                          '--out', str(cur), '--baseline', str(base)])
        assert code == 2


class TestSessionEngineFacts:
    def test_missing_file_is_empty(self, tmp_path):
        assert GraphSessionMixin.engine_facts(_Stub(tmp_path)) == {}

    def test_uncollectable_file_is_empty(self, tmp_path):
        """**坑 4**：采集不可用的那份不能拿来当依据。"""
        _write_facts(tmp_path, probe_complete=False)
        assert GraphSessionMixin.engine_facts(_Stub(tmp_path)) == {}

    def test_valid_file_is_read_and_cached(self, tmp_path):
        _write_facts(tmp_path)
        stub = _Stub(tmp_path)
        first = GraphSessionMixin.engine_facts(stub)
        assert first and EF.value_of(first, 'images') == ['lan']
        assert GraphSessionMixin.engine_facts(stub) is first     # 命中缓存

    def test_cache_invalidates_when_file_changes(self, tmp_path):
        _write_facts(tmp_path)
        stub = _Stub(tmp_path)
        GraphSessionMixin.engine_facts(stub)
        path = EF.facts_path(tmp_path)
        facts = EF.build_facts({'images': [(EF.SRC_PROBE, ['lan', 'ye'])],
                                'labels': [(EF.SRC_PROBE, ['start'])]},
                               collection={'parse_clean': True,
                                           'probe_complete': True})
        # 确保 mtime 变化（同一纳秒内写入会命中旧缓存）
        path.write_text(EF.dumps(facts), encoding='utf-8')
        stat = path.stat()
        import os
        os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 10 ** 9))
        again = GraphSessionMixin.engine_facts(stub)
        assert EF.value_of(again, 'labels') == ['start']

    def test_broken_json_is_empty_not_crash(self, tmp_path):
        EF.facts_path(tmp_path).write_text('{ not json', encoding='utf-8')
        assert GraphSessionMixin.engine_facts(_Stub(tmp_path)) == {}


class TestSessionWarnings:
    def test_overridden_builtin_is_reported(self, tmp_path):
        facts = EF.build_facts({}, collection={'parse_clean': True,
                                              'probe_complete': True},
                               derived={'overridden_builtins': [
                                   {'name': 'pause', 'file': 'aa.rpy',
                                    'line': 3, 'reason': '内置'}]})
        EF.facts_path(tmp_path).write_text(EF.dumps(facts), encoding='utf-8')
        warnings = GraphSessionMixin.engine_facts_warnings(_Stub(tmp_path))
        assert warnings and 'pause' in warnings[0]

    def test_no_facts_no_warnings(self, tmp_path):
        assert GraphSessionMixin.engine_facts_warnings(_Stub(tmp_path)) == []


# ---------------------------------------------------------------------------
#  接进真实 JSON：真机事实的形状能被 E2 消费
# ---------------------------------------------------------------------------

REAL = Path(__import__('os').environ.get(
    'STORYGRAPH_REAL_PROJECT', r'D:\千连万理\RenPyTest\sr_general'))


class TestAgainstRealFacts:
    """拿真机采出来的 `engine_facts.json` 当输入 —— 没有就 skip。"""

    def _load(self):
        path = EF.facts_path(REAL)
        if not path.is_file():
            pytest.skip(f'没有真机事实文件：{path}（先跑 cli engine-facts）')
        return json.loads(path.read_text(encoding='utf-8'))

    def test_real_tag_groups_are_consumable(self):
        facts = self._load()
        assert EF.attributes_of(facts, 'lan'), '真机 lan 应当有属性'
        got = EF.grouped_choices(facts, 'lan', {'outfit': '服装'})
        labels = {g['label'] for g in got}
        assert '服装' in labels, labels
        outfit = next(g for g in got if g['group'] == 'outfit')
        assert 'clothes2' in outfit['values']

    def test_real_beat_produces_display_values(self):
        facts = self._load()
        got = beats.expression_groups_for('lan', None, [], facts)
        assert got, '真机 lan 应当能直接给出分组候选'
        total = sum(len(g['values']) for g in got)
        assert total >= 5, total


# ---------------------------------------------------------------------------
#  一维一格（2026-10-12）：`show ye base_handdown gratified` 的两个词
# ---------------------------------------------------------------------------

class _SlotSym:
    """
    最小符号表：够 `attr_slot_rows` 用（`attrs_of_tag` + 项目根）。

    ⚠ `attrs_of_tag` 给的是**资源表的全部属性**（含 `base_handdown_mask`
    这类内部件）—— 真机就是这个形状，所以判据也照这个形状造。
    """

    def __init__(self, table, project_root=None, words=None):
        self._t = table
        self.registry = type('R', (), {'project_root': project_root})()
        self.project_root = project_root
        self.expressions = dict(words or {})
        self.project_words = {'expression': set(self.expressions),
                              'position': set()}

    def attrs_of_tag(self, tag):
        return list(self._t.get(tag) or [])

    def expression_word_for(self, tag, attr):
        return ''


#: 真机 `ye` 的形状：两个组（`outfit` 2 个值 / `expressions` 3 个值），
#: 外加两个不属于任何组的内部件（掩码 / 切图）。
YE_ATTRS = {
    'ye': [{'group': 'expressions', 'attribute': 'ahah'},
           {'group': 'expressions', 'attribute': 'gratified'},
           {'group': 'expressions', 'attribute': 'happy'},
           {'group': 'outfit', 'attribute': 'base_handdown'},
           {'group': 'outfit', 'attribute': 'base_swimsuit'}],
}
#: 引擎给的候选顺序（`gratified` 排在 `happy` 前，姿势词排在表情后面 —— 与真机一致）。
#: ⚠ `base_handdown_mask` 这类内部件**不在候选里**（真机就是这样：`candidates`
#: 是引擎枚举的"可显示属性"，掩码/切图进不去）。
YE_CANDS = {'ye': ['ahah', 'gratified', 'happy',
                   'base_handdown', 'base_swimsuit']}
YE_SYM = {'ye': ['ahah', 'gratified', 'happy', 'base_handdown',
                 'base_swimsuit', 'base_handdown_mask', '_hd_cut']}


def _ye_sym(project_root=None):
    return _SlotSym(YE_SYM, project_root=project_root)


def _ye_facts():
    return _facts(YE_ATTRS, YE_CANDS)


class TestAttrSlotRows:
    def test_two_dimensions_become_two_slots(self):
        """用户原始诉求：`base_handdown` + 表情 = **两格**，不是挤在一个格子里。

        ⚠ 顺序**跟着引擎给的分组顺序**走（真机 `ye` 是 `expressions` 在前）——
        工具不重排，重排等于把引擎"适合展示的顺序"丢掉。
        """
        rows = beats.attr_slot_rows('ye base_handdown gratified', _ye_sym(),
                                    facts=_ye_facts())
        assert [r['group'] for r in rows] == ['expressions', 'outfit'], rows
        assert [r['value'] for r in rows] == ['gratified', 'base_handdown']

    def test_each_slot_offers_its_own_group_only(self):
        """服装那一格的候选里**不该混进表情**（反之亦然）。"""
        rows = beats.attr_slot_rows('ye base_handdown happy', _ye_sym(),
                                    facts=_ye_facts())
        by_group = {r['group']: r for r in rows}
        assert by_group['outfit']['options'] == ['base_handdown', 'base_swimsuit']
        assert 'happy' in by_group['expressions']['options']
        assert 'base_handdown' not in by_group['expressions']['options']

    def test_one_word_is_the_old_single_chip(self):
        """**通用性底线**：一个词 → 走老路（`show ling happy` 行为一字不变）。"""
        assert beats.attr_slot_rows('ye happy', _ye_sym(),
                                    facts=_ye_facts()) == []

    def test_unknown_tag_falls_back(self):
        assert beats.attr_slot_rows('bg1.1', _ye_sym(),
                                    facts=_ye_facts()) == []

    def test_without_facts_and_layers_it_claims_nothing(self):
        """没采过引擎事实、也读不到项目里的 `layeredimage` → **不猜**。"""
        assert beats.attr_slot_rows('ye base_handdown gratified', _ye_sym()) == []

    def test_engine_facts_beat_the_project_declaration(self):
        """声明只当改名层：`outfit` → 「服装」，结构仍由引擎给。"""
        rows = beats.attr_slot_rows(
            'ye base_handdown gratified', _ye_sym(), groups=[_Group('outfit', '服装')],
            facts=_ye_facts())
        outfit = next(r for r in rows if r['group'] == 'outfit')
        assert outfit['label'] == '服装'

    def test_mask_words_are_preserved_as_extras(self):
        """**不认识的词一个字节都不丢**（`base_handdown_mask` 这类内部件）。"""
        rows = beats.attr_slot_rows('ye base_handdown base_handdown_mask gratified',
                                    _ye_sym(), facts=_ye_facts())
        assert rows, '两个维度都在 → 应当给格子'
        assert rows[0]['extras'] == ['base_handdown_mask'], rows[0]


class TestApplyAttrSlot:
    def test_switching_expression_keeps_the_outfit(self):
        """**这一条就是用户那个 bug**：换表情不许把 `base_handdown` 吃掉。"""
        got = beats.apply_attr_slot('ye base_handdown gratified', 'expressions',
                                    'ahah', symbols=_ye_sym(),
                                    facts=_ye_facts())
        assert 'base_handdown' in got and 'gratified' not in got, got
        assert got.split()[0] == 'ye'

    def test_switching_outfit_keeps_the_expression(self):
        got = beats.apply_attr_slot('ye base_handdown gratified', 'outfit',
                                    'base_swimsuit', symbols=_ye_sym(),
                                    facts=_ye_facts())
        assert 'gratified' in got and 'base_swimsuit' in got, got
        assert 'base_handdown' not in got.split(), got

    def test_clearing_the_outfit_drops_only_that_word(self):
        got = beats.apply_attr_slot('ye base_handdown gratified', 'outfit', '',
                                    symbols=_ye_sym(), facts=_ye_facts())
        assert got == 'ye gratified', got

    def test_unknown_group_changes_nothing(self):
        """认不出的维度 → **原样返回**（宁可不改，也不猜着改）。"""
        spec = 'ye base_handdown gratified'
        assert beats.apply_attr_slot(spec, 'nope', 'x', symbols=_ye_sym(),
                                     facts=_ye_facts()) == spec

    def test_single_chip_spec_is_left_alone(self):
        """只有一格时这条路**不许**动它（老路由 `_set_field` 负责）。"""
        assert beats.apply_attr_slot('ye happy', 'expressions', 'ahah',
                                     symbols=_ye_sym(),
                                     facts=_ye_facts()) == 'ye happy'


class TestSilentDefaultAttrs:
    """
    "写了却被引擎丢掉"的那一半 —— 用户问的正是"为什么不报"。

    ⚠ 这条**要 `layer_model` 才查**（不像别的检查只靠引擎事实）：因为"谁顶掉谁"
    是叠层语义。给不出模型 → **一个字都不说**（与 `unrecognized_attrs` 同规矩）。
    ⚠ 还要求**先有 N3 同组冲突**（`only_groups`）：卡片每翻一页都调它，而没冲突
    的拍占绝大多数 —— 那一条路**一次磁盘都不该摸**。
    """

    def _model(self, tmp_path):
        from storygraph.story.layered_model import load_layer_model
        game = tmp_path / 'game'
        game.mkdir(parents=True, exist_ok=True)
        (game / 'defs.rpy').write_text(
            'layeredimage lan:\n'
            '    group expressions:\n'
            '        attribute ahah:\n'
            '            "a.png"\n'
            '        attribute happy default:\n'
            '            "h.png"\n'
            '        attribute sad:\n'
            '            "s.png"\n'
            '    group outfit:\n'
            '        attribute clothes2:\n'
            '            "c.png"\n', encoding='utf-8')
        return load_layer_model(tmp_path)

    def _beat(self, show):
        b = beats.Beat()
        b.show = show
        return b

    def test_second_word_in_one_group_is_reported(self, tmp_path):
        """`show lan sad ahah`：同组两个 → 引擎只用**声明靠前**的那个。"""
        model = self._model(tmp_path)
        got = beats.silent_default_attrs(self._beat('lan sad ahah'),
                                         layer_model=model,
                                         only_groups={'expressions'})
        assert got and got[0]['attr'] in ('sad', 'ahah'), got
        assert got[0]['tag'] == 'lan'

    def test_across_groups_nothing_is_reported(self, tmp_path):
        model = self._model(tmp_path)
        assert beats.silent_default_attrs(self._beat('lan sad clothes2'),
                                          layer_model=model,
                                          only_groups={'expressions'}) == []

    def test_without_a_model_nothing_is_claimed(self):
        assert beats.silent_default_attrs(self._beat('lan sad ahah'),
                                          only_groups={'expressions'}) == []

    def test_without_a_conflict_group_nothing_is_looked_up(self):
        """**省法**：没有同组冲突 → 连模型都不去建（返回 `[]`）。"""
        assert beats.silent_default_attrs(self._beat('lan sad ahah'),
                                          only_groups=set()) == []

    def test_card_note_stays_quiet_without_a_project_root(self):
        note = beats.card_note(self._beat('lan sad ahah'), None,
                               _facts(LAN_ATTRS, LAN_CANDS))
        assert '没生效' not in note

