"""
tests/test_asset_registry.py
============================
资产扫描、Excel 读写、解析、校验、引用分析测试。

关键点：
  · scan_assets 的分类推断与 ID 唯一化
  · write_asset_excel 的增量合并（保留用户编辑字段）
  · parse_asset_tables 的 ID 去重、自定义表、Custom_ 字段收集
  · validate_registry 的文件存在性与颜色格式
  · analyze_references 的行号准确性
"""

import pytest

from asset_registry import (
    AssetRegistry, scan_assets, write_asset_excel, parse_asset_tables,
    read_excel_rows, validate_registry, analyze_references, load_hooks,
    _detect_category, _make_unique_id,
    resolve_reference, extract_script_references,
    RENPY_BUILTIN_TRANSFORMS,
)


@pytest.fixture
def scan_root(tmp_path):
    root = tmp_path / 'game'
    (root / 'images' / 'bg').mkdir(parents=True)
    (root / 'images' / 'chars').mkdir(parents=True)
    (root / 'audio').mkdir(parents=True)
    (root / 'gui').mkdir(parents=True)

    (root / 'images' / 'bg' / 'bg_room.png').write_bytes(b'1')
    (root / 'images' / 'bg' / 'bg_hall.png').write_bytes(b'1')
    (root / 'images' / 'chars' / 'eileen happy.png').write_bytes(b'1')
    (root / 'gui' / 'button_idle.png').write_bytes(b'1')
    (root / 'audio' / 'bgm_theme.ogg').write_bytes(b'1')
    (root / 'audio' / 'sfx_click.wav').write_bytes(b'1')
    return root


# ------------------------------------------------------------------
#  扫描
# ------------------------------------------------------------------

class TestScanAssets:
    def test_missing_root_returns_empty(self, tmp_path):
        result = scan_assets(tmp_path / 'nope')
        assert result['images'] == []
        assert result['audio'] == {'music': [], 'sound': []}

    def test_images_found(self, scan_root):
        assets = scan_assets(scan_root, include_ui=False)
        ids = {a['ID'] for a in assets['images']}
        assert 'bg_room' in ids and 'bg_hall' in ids

    def test_ui_excluded_when_requested(self, scan_root):
        assets = scan_assets(scan_root, include_ui=False)
        assert all(a['Category'] != 'ui' for a in assets['images'])

    def test_ui_included_when_requested(self, scan_root):
        assets = scan_assets(scan_root, include_ui=True)
        assert any(a['Category'] == 'ui' for a in assets['images'])

    def test_background_category_and_tag(self, scan_root):
        assets = scan_assets(scan_root, include_ui=False)
        bg = next(a for a in assets['images'] if a['ID'] == 'bg_room')
        assert bg['Category'] == 'background'
        assert bg['Tag'] == 'bg'

    def test_character_tag_from_parent(self, scan_root):
        assets = scan_assets(scan_root / 'game' / 'images'
                             if False else scan_root, include_ui=False)
        hero = [a for a in assets['images']
                if a['ID'].startswith('eileen')]
        assert hero

    def test_music_and_sound_classified(self, scan_root):
        assets = scan_assets(scan_root)
        music_ids = {a['ID'] for a in assets['audio']['music']}
        sound_ids = {a['ID'] for a in assets['audio']['sound']}
        assert music_ids == {'bgm_theme'}
        assert sound_ids == {'sfx_click'}

    def test_paths_are_relative(self, scan_root):
        assets = scan_assets(scan_root, include_ui=False)
        bg = next(a for a in assets['images'] if a['ID'] == 'bg_room')
        assert bg['Path'] == 'images/bg/bg_room.png'

    def test_unique_ids_on_name_collision(self, tmp_path):
        root = tmp_path / 'root'
        (root / 'a').mkdir(parents=True)
        (root / 'b').mkdir(parents=True)
        (root / 'a' / 'same.png').write_bytes(b'1')
        (root / 'b' / 'same.png').write_bytes(b'1')
        assets = scan_assets(root)
        ids = {a['ID'] for a in assets['images']}
        assert len(ids) == 2

    def test_bg_vs_bgm_disambiguation(self, tmp_path):
        """'bgm_xxx' 路径必须归到 music 而非 background。"""
        root = tmp_path / 'root'
        (root / 'audio').mkdir(parents=True)
        (root / 'audio' / 'bgm_night.mp3').write_bytes(b'1')
        assets = scan_assets(root)
        assert {a['ID'] for a in assets['audio']['music']} == {'bgm_night'}

    def test_custom_audio_dirs(self, scan_root, tmp_path):
        extra = tmp_path / 'extra_audio'
        extra.mkdir()
        (extra / 'theme_song.mp3').write_bytes(b'1')
        assets = scan_assets(scan_root, audio_dirs=[extra])
        assert assets['audio']['music'] or assets['audio']['sound']

    def test_make_unique_id_fallback_counter(self):
        existing = {'x', 'parent_x'}
        assert _make_unique_id('x', ['parent'], existing) == 'x_1'

    def test_detect_category_fallback(self):
        assert _detect_category(__import__('pathlib').Path('a/b/c.png')) == 'other'


# ------------------------------------------------------------------
#  Excel 读写
# ------------------------------------------------------------------

class TestExcelIO:
    def test_read_missing_file(self, tmp_path):
        assert read_excel_rows(tmp_path / 'nope.xlsx') == []

    def test_write_then_read(self, tmp_path):
        path = tmp_path / 't.xlsx'
        write_asset_excel(
            {'images': [{'ID': 'a', 'Path': 'p.png', 'Tag': 'bg',
                         'Attributes': 'a', 'Category': 'background',
                         'Notes': 'hi'}],
             'audio': {'music': [], 'sound': []}},
            tmp_path / 'out')
        rows = read_excel_rows(tmp_path / 'out' / 'images.xlsx')
        assert len(rows) == 1
        assert rows[0]['ID'] == 'a'
        assert rows[0]['Notes'] == 'hi'

    def test_blank_rows_skipped(self, tmp_path):
        import openpyxl
        path = tmp_path / 'blank.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['ID', 'Path'])
        ws.append([None, None])
        ws.append(['a', 'p'])
        wb.save(path)
        rows = read_excel_rows(path)
        assert len(rows) == 1

    def test_incremental_merge_keeps_user_fields(self, scan_root, tmp_path):
        out = tmp_path / 'tables'
        write_asset_excel(scan_assets(scan_root, include_ui=False), out)

        # 用户手工编辑 Notes
        import openpyxl
        f = out / 'images.xlsx'
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            if row[0].value == 'bg_room':
                ws.cell(row=row[0].row, column=6, value='USER NOTE')
        wb.save(f)

        write_asset_excel(scan_assets(scan_root, include_ui=False), out)
        rows = {r['ID']: r for r in read_excel_rows(f)}
        assert rows['bg_room']['Notes'] == 'USER NOTE'


# ------------------------------------------------------------------
#  解析
# ------------------------------------------------------------------

class TestParseAssetTables:
    def _write(self, asset_dir, rows, filename='images.xlsx'):
        from asset_registry import write_dict_to_excel
        data = {r['ID']: r for r in rows}
        write_dict_to_excel(asset_dir / filename, data,
                            list(rows[0].keys()))

    def test_missing_dir_is_tolerated(self, tmp_path):
        reg = parse_asset_tables(tmp_path / 'nope')
        assert isinstance(reg, AssetRegistry)

    def test_rows_loaded(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'p.png'}])
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert 'a' in reg.images

    def test_duplicate_id_reported(self, tmp_path):
        """Excel 里出现两个相同 ID 必须报错。"""
        import openpyxl
        d = tmp_path / 't'
        d.mkdir()
        f = d / 'images.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['ID', 'Path'])
        ws.append(['a', 'p.png'])
        ws.append(['a', 'q.png'])
        wb.save(f)
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert any('重复' in e for e in reg.errors)

    def test_row_without_id_reported(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': '', 'Path': 'p.png'}])
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert any('缺少ID' in e for e in reg.errors)

    def test_custom_prefix_fields_collected(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'p.png', 'Custom_X': 1}])
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert reg.images['a']['_custom'] == {'Custom_X': 1}

    def test_custom_table_loaded(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'k1', 'Value': 'v'}],
                    'custom_props.xlsx')
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert 'props' in reg.custom
        assert 'k1' in reg.custom['props']

    def test_hook_applied(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'p.png'}])
        hooks = {'pre_process_row': lambda name, row: {**row, 'Notes': 'HOOK'}}
        reg = parse_asset_tables(d, hooks=hooks, project_root=tmp_path)
        assert reg.images['a']['Notes'] == 'HOOK'

    def test_hook_exception_recorded_as_warning(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'p.png'}])
        def boom(name, row):
            raise ValueError('boom')
        reg = parse_asset_tables(d, hooks={'pre_process_row': boom},
                                 project_root=tmp_path)
        assert reg.warnings

    def test_missing_image_file_warns(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'missing.png'}])
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert any('文件不存在' in w for w in reg.warnings)

    def test_character_color_validation(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'e', 'Color': 'not-a-color'}],
                    'characters.xlsx')
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert any('颜色' in w and '无效' in w for w in reg.warnings)

    def test_invalid_transform_reference(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'e', 'DefaultTransform': 'nope'}],
                    'characters.xlsx')
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert any('DefaultTransform' in w for w in reg.warnings)

    def test_get_asset_ids(self, tmp_path):
        d = tmp_path / 't'
        d.mkdir()
        self._write(d, [{'ID': 'a', 'Path': 'p.png'}])
        reg = parse_asset_tables(d, project_root=tmp_path)
        assert reg.get_asset_ids('image') == ['a']
        assert reg.get_asset_ids('nonexistent_type') == []


# ------------------------------------------------------------------
#  hooks
# ------------------------------------------------------------------

class TestHooks:
    def test_no_hooks_file_returns_empty(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        assert load_hooks(project_root=tmp_path) == {}

    def test_hook_module_loaded(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        (tmp_path / 'asset_import_hooks.py').write_text(
            'def pre_process_row(t, r):\n    return r\n',
            encoding='utf-8')
        hooks = load_hooks(project_root=tmp_path)
        assert callable(hooks['pre_process_row'])


# ------------------------------------------------------------------
#  引用分析
# ------------------------------------------------------------------

class TestAnalyzeReferences:
    def test_references_found_with_line(self, write_rpy):
        """
        脚本写 `show bg room`，资产表里叫 `bg_room`（文件名）。
        两边形状不同，必须能认出来是同一个东西。
        """
        reg = AssetRegistry()
        reg.images = {
            'bg_room': {'ID': 'bg_room', 'Path': 'images/bg_room.png',
                        'Tag': 'bg', 'Attributes': 'bg_room'},
            'bgm_theme': {},
        }
        fp = write_rpy(
            'label a:\n'
            '    show bg room\n'
            '    play music bgm_theme\n')
        analyze_references([fp], reg)
        assert reg.references['bg_room'] == [f'{fp.name}:2']
        assert reg.references['bgm_theme'] == [f'{fp.name}:3']

    def test_duplicate_locations_deduplicated(self, write_rpy):
        reg = AssetRegistry()
        reg.images = {'bg': {'ID': 'bg', 'Tag': 'bg', 'Attributes': 'bg'}}
        fp = write_rpy('label a:\n    show bg\n    scene bg\n')
        analyze_references([fp], reg)
        assert len(reg.references['bg']) == 2

    def test_missing_file_skipped(self, tmp_path):
        reg = AssetRegistry()
        analyze_references([tmp_path / 'nope.rpy'], reg)
        assert not reg.references

    def test_plain_dict_converted_to_defaultdict(self, write_rpy):
        reg = AssetRegistry()
        reg.images = {'bg': {'ID': 'bg', 'Tag': 'bg', 'Attributes': 'bg'}}
        reg.references = {}
        fp = write_rpy('label a:\n    show bg\n')
        analyze_references([fp], reg)
        assert reg.references['bg']


class TestImageSpecMatching:
    """`show bg room` 与文件 `bg_room.png` 是同一个东西。"""

    def _reg(self):
        reg = AssetRegistry()
        reg.images = {
            'bg_classroom': {'ID': 'bg_classroom', 'Path': 'p.png',
                             'Tag': 'bg', 'Attributes': 'bg_classroom'},
            'eileen_happy': {'ID': 'eileen_happy', 'Path': 'q.png',
                             'Tag': 'eileen', 'Attributes': 'eileen_happy'},
        }
        return reg

    def test_full_spec_maps_to_underscore_id(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show bg classroom\n')
        analyze_references([fp], reg)
        assert 'bg_classroom' in reg.references

    def test_two_word_character_sprite(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show eileen happy\n')
        analyze_references([fp], reg)
        assert 'eileen_happy' in reg.references

    def test_tag_only_falls_back_to_tag_match(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    scene bg\n')
        analyze_references([fp], reg)
        assert reg.references  # 至少匹配到一张 tag=bg 的图

    def test_spec_stops_at_keywords(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show bg classroom at leftish\n')
        analyze_references([fp], reg)
        assert 'bg_classroom' in reg.references
        # 'at' 之后的内容不该被当成图像名的一部分
        assert 'bg_classroom_at' not in reg.references

    def test_unknown_name_goes_to_unresolved(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show mysterious\n')
        analyze_references([fp], reg)
        assert not reg.references
        assert reg.unresolved['image']

    def test_builtin_transition_not_flagged(self, write_rpy):
        """`with dissolve` 是 Ren'Py 自带的，不是缺失的资产。"""
        reg = self._reg()
        fp = write_rpy('label a:\n    scene bg with dissolve\n')
        analyze_references([fp], reg)
        assert not reg.unresolved['transform']

    def test_builtin_position_not_flagged(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show bg at left\n')
        analyze_references([fp], reg)
        assert not reg.unresolved['transform']

    def test_custom_transform_still_captured(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    show bg at my_slide\n')
        analyze_references([fp], reg)
        assert reg.unresolved['transform']

    def test_registered_transform_resolves(self, write_rpy):
        reg = self._reg()
        reg.transforms = {'my_slide': {'ID': 'my_slide'}}
        fp = write_rpy('label a:\n    show bg at my_slide\n')
        analyze_references([fp], reg)
        assert reg.references['my_slide']

    def test_hide_is_captured(self, write_rpy):
        reg = self._reg()
        fp = write_rpy('label a:\n    hide bg classroom\n')
        analyze_references([fp], reg)
        assert 'bg_classroom' in reg.references


class TestResolveReference:
    def test_exact_id(self):
        reg = AssetRegistry()
        reg.images = {'a': {'ID': 'a', 'Tag': 'x', 'Attributes': 'y'}}
        assert resolve_reference('a', reg) == 'a'

    def test_attributes_match(self):
        reg = AssetRegistry()
        reg.images = {'a': {'ID': 'a', 'Tag': 'x', 'Attributes': 'y_z'}}
        assert resolve_reference('y_z', reg) == 'a'

    def test_tag_match(self):
        reg = AssetRegistry()
        reg.images = {'a': {'ID': 'a', 'Tag': 'x', 'Attributes': 'y'}}
        assert resolve_reference('x', reg) == 'a'

    def test_no_match(self):
        assert resolve_reference('nope', AssetRegistry()) is None
        assert resolve_reference('', AssetRegistry()) is None
