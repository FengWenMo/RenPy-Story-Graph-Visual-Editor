"""
边界②：CLI 也能走**增量**，第三类内容不再只能靠 UI 保住
=========================================================

交付评审（2026-10-04）列的第二条边界：

> **第三类只在"增量"路径保住**。CLI 的 `table xlsx2rpy` 只有 xlsx、
> 没有源 .rpy → patcher 无源 → **永远走全量** → 只能用于**新建**；
> **改已有章节必须用 UI 卡牌编辑器**（有 `script_path` → patcher → 保住）。

这一条现在不成立了：`--source <这一章原来的 .rpy>` 让命令行和编辑器
**等价** —— 原文逐字节保留，只换改过的那几行。

但"能走增量"只解决了一半。另一半是**补丁没配上时怎么办**：

    `--source X --out 新路径` 时目标是新的 → P0 那条"手写护栏"不适用
    → 静默退回全量 → 源里的 `define` / `init python` 在产物里没了，
    而退出码是 0。

真实项目 `game/pomodorro/pomodorro.rpy` 实测就是这样：产物里
`define t = Character("Roro", image="dororo")`、`image dororo neutral = …`
等 **37 处顶层语句整批消失**。所以这个文件守两件事：

  ① 能走增量时**真的走增量**（原文保住、只改改过的那行）；
  ② 走不了增量、又要丢顶层语句时**拒绝写盘并说出来**（`OPAQUE_LOST`），
     而不是静默写一份"能跑但少了定义"的脚本。

②的逃生舱是 `--force`（降级为 warn，仍然报出来）—— 与 `HANDWRITTEN_TARGET`
同一性质：说的不是"产物自己坏了"，而是"这个目标配不上这份产物"。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
from storygraph.story.beat_build import beats_from_parse_result
from storygraph.story.excel_export import export_beats_to_excel

#: 一章"真实形状"的脚本：第三类（不进表）与拍混在一起。
CHAPTER = (
    '# 第一章 开场\n'
    'default score = 0\n'
    'define e = Character("Eileen")\n'
    '\n'
    'init python:\n'
    '    def helper(x):\n'
    '        return x * 2\n'
    '\n'
    'label ch01:\n'
    '    e "开场台词"\n'
    '    e "第二句"\n'
    '    return\n'
)

OPAQUE_TOKENS = ('default score = 0', 'define e = Character("Eileen")',
                 'init python:', 'def helper(x):')


def _write_chapter(tmp_path, text: str = CHAPTER,
                   name: str = 'ch01.rpy') -> Path:
    p = tmp_path / name
    p.write_text(text, encoding='utf-8')
    return p


def _table_of(src: Path, tmp_path) -> Path:
    """`.rpy` → `.xlsx`（走真实的那条解析链）。"""
    pr, _ = IncrementalRenPyAnalyzer().analyze(files=[src], graph=None)
    beats = beats_from_parse_result(pr)
    xl = tmp_path / (src.stem + '.xlsx')
    export_beats_to_excel(beats, xl)
    return xl


def _edit_dialogue(xl: Path, old: str, new: str) -> None:
    """把表里某一格台词改掉（模拟"文案在 Excel 里改了一句"）。"""
    from openpyxl import load_workbook
    wb = load_workbook(xl)
    ws = wb['Beats']
    headers = [c.value for c in ws[1]]
    col = headers.index('台词') + 1
    hit = 0
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=col).value or '') == old:
            ws.cell(row=row, column=col).value = new
            hit += 1
    assert hit, f'表里没有 {old!r} 这一格'
    wb.save(xl)


# ---------------------------------------------------------------------------
#  1. 能走增量 → 真的走增量（原文保住）
# ---------------------------------------------------------------------------

class TestIncrementalPath:
    def test_merge_incrementally_uses_the_source(self, tmp_path):
        from storygraph.pipeline.compile import merge_incrementally
        src = _write_chapter(tmp_path)
        pr, _ = IncrementalRenPyAnalyzer().analyze(files=[src], graph=None)
        beats = beats_from_parse_result(pr)

        out = merge_incrementally(beats, 'label ch01:\n    pass\n', src)
        assert out.used_patch, f'给了源却没走增量：{out.reason}'
        assert 'init python:' in out.text
        assert 'default score = 0' in out.text

    def test_no_source_means_full_generation(self, tmp_path):
        """没给源就是原来的行为（新建文件的正当路径），原因要说出来。"""
        from storygraph.pipeline.compile import merge_incrementally
        src = _write_chapter(tmp_path)
        pr, _ = IncrementalRenPyAnalyzer().analyze(files=[src], graph=None)
        beats = beats_from_parse_result(pr)
        out = merge_incrementally(beats, 'label ch01:\n    pass\n', None)
        assert not out.used_patch
        assert out.reason, '没走增量却没给原因 —— 这就是"静默退回全量"'

    def test_cli_source_keeps_opaque_and_changes_one_line(self, tmp_path,
                                                          capsys):
        """
        端到端：`table xlsx2rpy --source` → 只改一句台词，第三类原样。

        这条就是"边界②"的正向证明：**命令行也能安全地改已有章节**。
        """
        import difflib

        from storygraph.cli import main

        src = _write_chapter(tmp_path)
        before = src.read_text(encoding='utf-8')
        xl = _table_of(src, tmp_path)
        _edit_dialogue(xl, '开场台词', '改过的台词')
        out = tmp_path / 'out.rpy'

        code = main(['table', 'xlsx2rpy', str(xl), '--source', str(src),
                     '--out', str(out)])
        log = capsys.readouterr()
        assert code == 0, (log.out + log.err)
        assert out.exists(), '没有产物'

        text = out.read_text(encoding='utf-8')
        for token in OPAQUE_TOKENS:
            assert token in text, f'第三类丢了：{token}\n{text}'
        assert '改过的台词' in text

        diff = [l for l in difflib.unified_diff(before.splitlines(),
                                                text.splitlines(), n=0,
                                                lineterm='')
                if l[:1] in '+-' and l[:3] not in ('---', '+++')]
        # 生成的标记（2 行）+ 改掉的那一句（1 减 1 加）= 4
        assert len(diff) <= 4, f'改了 1 句却动了这么多行：{diff}'

    def test_cli_reports_that_it_used_the_patch(self, tmp_path, capsys):
        """走没走增量必须说出来 —— 用户要能确认第三类被保住了。"""
        from storygraph.cli import main
        src = _write_chapter(tmp_path)
        xl = _table_of(src, tmp_path)
        _edit_dialogue(xl, '开场台词', 'x')
        code = main(['table', 'xlsx2rpy', str(xl), '--source', str(src),
                     '--out', str(tmp_path / 'o.rpy')])
        out = capsys.readouterr().out
        assert code == 0
        assert '增量补丁' in out, out


# ---------------------------------------------------------------------------
#  2. 走不了增量又要丢顶层语句 → 拒写（不许静默降级）
# ---------------------------------------------------------------------------

class TestOpaqueLostIsRefused:
    def _mismatched_table(self, tmp_path):
        """
        造一张**和源对不上**的表：表来自另一个剧本。

        补丁必然失败（区间对不上），于是退回全量 —— 而全量会丢顶层语句。
        """
        other = tmp_path / 'other.rpy'
        other.write_text('label other:\n    e "另一章"\n', encoding='utf-8')
        return _table_of(other, tmp_path)

    def test_full_fallback_refuses_to_drop_top_level_statements(
            self, tmp_path, capsys):
        from storygraph.cli import main

        src = _write_chapter(tmp_path)
        before = src.read_text(encoding='utf-8')
        xl = self._mismatched_table(tmp_path)
        out = tmp_path / 'new_out.rpy'          # 目标不存在 → 手写护栏不管用

        code = main(['table', 'xlsx2rpy', str(xl), '--source', str(src),
                     '--out', str(out)])
        err = capsys.readouterr().err
        assert code == 1, '丢了 37 处定义却报成功'
        assert not out.exists(), '拦下了却还是写出了文件'
        assert 'OPAQUE_LOST' in err, err
        assert src.read_text(encoding='utf-8') == before, '源文件被改写'

    def test_force_still_reports_the_loss(self, tmp_path, capsys):
        """逃生舱：`--force` 能过，但**必须**仍然把丢的东西报出来。"""
        from storygraph.cli import main

        src = _write_chapter(tmp_path)
        xl = self._mismatched_table(tmp_path)
        out = tmp_path / 'forced.rpy'

        code = main(['table', 'xlsx2rpy', str(xl), '--source', str(src),
                     '--out', str(out), '--force'])
        out_text = capsys.readouterr().out
        assert code == 0
        assert out.exists()
        assert 'OPAQUE_LOST' in out_text or '顶层语句' in out_text, (
            f'强制覆盖时也必须报出来（沉默不是选项）：\n{out_text}')

    def test_without_source_nothing_changes(self, tmp_path, capsys):
        """没给 `--source` 时判据无从谈起 —— 不能凭空拦人。"""
        from storygraph.cli import main
        src = _write_chapter(tmp_path)
        xl = self._mismatched_table(tmp_path)
        out = tmp_path / 'plain.rpy'

        code = main(['table', 'xlsx2rpy', str(xl), '--out', str(out)])
        assert code == 0, capsys.readouterr().err
        assert out.exists()


# ---------------------------------------------------------------------------
#  3. 门禁的两条护栏互不干扰
# ---------------------------------------------------------------------------

class TestGuardInteractions:
    def test_merged_product_may_overwrite_handwritten_target(self, tmp_path):
        """增量产物**可以**写回手写源文件（它包含原文全部内容）。"""
        from storygraph.story.generation_guard import write_generation
        target = _write_chapter(tmp_path)
        text = '# 由剧本表自动生成\n\n' + target.read_text(encoding='utf-8')

        # 交的是**底稿路径**（2026-10-05 实测修：以前这里传的是布尔 `True`，
        # 于是"我走了增量"这句话本身就能越过护栏 —— 见下一条用例）。
        art = write_generation(text, target, merged_from_source=target)
        assert art.written, '增量产物被手写护栏拦下了（边界②等于没做）'

    def test_merged_claim_does_not_cover_someone_elses_file(self, tmp_path):
        """
        **底稿不是目标时，手写护栏照旧拦**（哪怕调用方说"这次是增量"）。

        为什么单开一条（2026-10-05 实测修）：`merged_from_source` 从前是个
        布尔，调用方声明一句"我走了增量"就能越过护栏 —— 而"编辑甲章、
        把产物写到乙章的手写文件上"正是靠这条护栏挡住的：那份产物里有
        甲章的全部内容，**却没有乙章的** `default` / `init python`。
        """
        from storygraph.story.generation_guard import write_generation
        base = _write_chapter(tmp_path, name='ch01.rpy')
        other = _write_chapter(tmp_path, name='ch02.rpy')
        text = '# 由剧本表自动生成\n\n' + base.read_text(encoding='utf-8')

        art = write_generation(text, other, merged_from_source=base)
        assert not art.written, '乙章的手写文件被甲章的增量产物覆盖了'
        assert [i.code for i in art.blocking] == ['HANDWRITTEN_TARGET']

    def test_merged_product_is_still_syntax_checked(self, tmp_path):
        """让路的**只有**手写那一条：产物自己坏了照样拦。"""
        from storygraph.story.generation_guard import write_generation
        target = _write_chapter(tmp_path)
        bad = '# 由剧本表自动生成\n\nlabel a:\n    $ l oh "3"\n'
        art = write_generation(bad, target, merged_from_source=target)
        assert not art.written
        assert [i.code for i in art.blocking] == ['PYTHON_SYNTAX_ERROR']

    def test_opaque_loss_detects_top_level_statements(self, tmp_path):
        from storygraph.story.generation_guard import opaque_loss
        src = CHAPTER
        lost = opaque_loss(src, 'label ch01:\n    e "x"\n')
        heads = [t for _n, t in lost]
        assert any('default score' in h for h in heads), heads
        assert any('init python' in h for h in heads), heads
        # 产物里都有 → 一处都不该报
        assert not opaque_loss(src, src)

    def test_opaque_loss_ignores_comments_and_names(self, tmp_path):
        """`# 注释` 与别的 label 不属于"顶层语句"，不许误报。"""
        from storygraph.story.generation_guard import opaque_loss
        src = '# 文件头注释\nlabel a:\n    e "x"\n\nlabel b:\n    e "y"\n'
        assert not opaque_loss(src, 'label a:\n    e "x"\n')
