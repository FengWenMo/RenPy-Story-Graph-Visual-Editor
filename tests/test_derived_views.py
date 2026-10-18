"""
tests/test_derived_views.py
===========================
D1/D2（用户 2026-10-19）：递推型派生视图 —— **承认它，但做对**。

用户原话：

> 「舞台快照/时间轴是"本拍 = 上一拍剩下的 + 本拍变化"，这是递推……
> 它不该是公式，这不是没做到，是做对了。所以要补的不是公式化，是让快照不骗人。」

> D1 | 表头写「生成于 …」，并给一个显式「刷新视图」入口
> D2 | 时间轴加一列「本拍变化」（相对上一拍），让人看出递推点在哪

判据：

① 两张派生视图的表头都写着**生成时间**（不然它看起来和真值一样权威）；
② 有一个**显式**重算入口（`refresh_stage_views`），重算后时间会变、Beats 一格不动；
③ 重算**不许造出「舞台快照1」「演出时间轴1」**（`create_sheet` 同名会追加序号 ——
   "刷新"正是最容易踩的地方）；
④ 时间轴的「本拍变化」= 快照里同一行的那份文字（同一份推导，不许两套说法）。
"""

from __future__ import annotations

import openpyxl
import pytest

from storygraph.story.beat_model import Beat, assign_identity


def _beats():
    beats = [
        Beat(seq=1, label='ch01', speaker='', text='开场',
             scene='bg room'),
        Beat(seq=2, label='ch01', speaker='e', text='你好',
             show='ling happy@left'),
        Beat(seq=3, label='ch01', speaker='e', text='只说话'),
        Beat(seq=4, label='ch01', speaker='', text='换音乐',
             bgm='bgm/theme.ogg'),
    ]
    assign_identity(beats)
    return beats


@pytest.fixture
def table(tmp_path):
    from storygraph.story.excel_export import export_beats_to_excel
    out = tmp_path / 'ch01.xlsx'
    export_beats_to_excel(_beats(), out)
    return out


class TestTheStamp:
    def test_both_views_say_when_they_were_generated(self, table):
        wb = openpyxl.load_workbook(table)
        for name in ('舞台快照', '演出时间轴'):
            row1 = [str(c.value or '') for c in wb[name][1]]
            assert any('生成于' in v for v in row1), (name, row1)

    def test_the_stamp_does_not_shift_the_columns(self, table):
        """生成时间写在表头右侧的空列里，第 6 列仍然是「本拍变化」。"""
        wb = openpyxl.load_workbook(table)
        assert wb['舞台快照'].cell(row=1, column=6).value == '本拍变化'
        assert wb['演出时间轴'].cell(row=1, column=6).value == '本拍变化'

    def test_the_stamp_says_how_to_recompute(self, table):
        wb = openpyxl.load_workbook(table)
        row1 = ' '.join(str(c.value or '') for c in wb['舞台快照'][1])
        assert '刷新派生视图' in row1, row1


class TestExplicitRefresh:
    def test_refresh_recomputes_without_duplicating_sheets(self, table):
        from storygraph.story.excel_export import refresh_stage_views

        refresh_stage_views(table)
        wb = openpyxl.load_workbook(table)
        assert '舞台快照1' not in wb.sheetnames, wb.sheetnames
        assert '演出时间轴1' not in wb.sheetnames, wb.sheetnames
        assert wb.sheetnames.count('舞台快照') == 1

    def test_refresh_does_not_touch_beats(self, table):
        from storygraph.story.excel_export import refresh_stage_views
        from storygraph.story.excel_import import import_beats_from_excel

        before = [b.text for b in import_beats_from_excel(table)]
        refresh_stage_views(table)
        after = [b.text for b in import_beats_from_excel(table)]
        assert before == after

    def test_the_stamp_moves_forward(self, table):
        """旧时间戳改成明显过去的时间，重算后必须变 —— 否则"刷新"没生效。"""
        import storygraph.story.excel_export as ex

        wb = openpyxl.load_workbook(table)
        ws = wb['舞台快照']
        hit = False
        for c in ws[1]:
            if '生成于' in str(c.value or ''):
                c.value = ex.VIEW_STAMP_PREFIX + '2000-01-01 00:00'
                hit = True
        assert hit
        wb.save(table)
        ex.refresh_stage_views(table)
        wb = openpyxl.load_workbook(table)
        row1 = ' '.join(str(c.value or '') for c in wb['舞台快照'][1])
        assert '2000-01-01' not in row1, row1

    def test_refreshing_candidates_does_not_duplicate_the_timeline(self, table):
        from storygraph.assets.asset_registry import AssetRegistry
        from storygraph.story.excel_export import refresh_candidate_sheets

        refresh_candidate_sheets(table, AssetRegistry())
        wb = openpyxl.load_workbook(table)
        assert '演出时间轴1' not in wb.sheetnames, wb.sheetnames


class TestTheTimelineShowsWhatChanged:
    def test_its_changes_column_matches_the_snapshot(self, table):
        """同一份推导：时间轴与快照里同一拍的「本拍变化」必须一字不差。"""
        wb = openpyxl.load_workbook(table)
        full = {}
        ws = wb['舞台快照']
        for r in range(2, ws.max_row + 1):
            lid = str(ws.cell(row=r, column=1).value or '')
            if lid:
                full[lid] = str(ws.cell(row=r, column=6).value or '')
        tl = wb['演出时间轴']
        assert tl.max_row >= 2
        for r in range(2, tl.max_row + 1):
            lid = str(tl.cell(row=r, column=1).value or '')
            if not lid:
                continue
            assert str(tl.cell(row=r, column=6).value or '') == full.get(lid, '')

    def test_every_timeline_row_has_a_change(self, table):
        wb = openpyxl.load_workbook(table)
        tl = wb['演出时间轴']
        for r in range(2, tl.max_row + 1):
            assert str(tl.cell(row=r, column=6).value or '').strip(), r
