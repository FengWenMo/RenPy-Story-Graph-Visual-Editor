"""
D1：把声明的 `choices` 接进 **Excel** 与**卡片**
================================================

M4 让 `columns` / `choices` 变成"可读的字段"，但**读得到不等于用得上** ——
项目写了候选却没人看，那它只是一段注释。这里把它接到人真的会碰到的地方：

  ① **Excel**：导出时给「自定义语句·头行」那一格挂**候选注释**
     （鼠标一悬停就看到能填什么）；
  ② **校验**：`W120`（缺声明的必需参数）/ `W121`（值不在候选里）——
     声明过的结构不校验，等于没声明；
  ③ **卡片**：`spec_arg_choices_for()` 给 UI 一组 `{参数名: [候选]}`。

⚠ 三条判断口径（都是"不猜"）：
  · 没声明 → **什么都不做**（不报、不注释、不给候选）；
  · 位置参数**不当成**有名字的参数（拿"第几个词"去对声明里的名字是猜）；
  · `W121` 是**警告不是错误**：候选是"常见取值"，不是白名单。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from storygraph.app.beats import spec_arg_choices_for
from storygraph.parsing.renpy_parser import IncrementalRenPyAnalyzer
from storygraph.pipeline.validate import spec_arg_values, validate_all
from storygraph.story.beat_build import beats_from_parse_result
from storygraph.story.beat_model import BEAT_FIELDS
from storygraph.story.excel_export import export_beats_to_excel
from storygraph.story.statement_registry import (StatementRegistry,
                                                 clear_cache,
                                                 spec_arg_choices)

SCRIPT = (
    'label a:\n'
    '    "第一句"\n'
    '    mycue "闪电" intensity=3 target=ling:\n'
    '        hold 0.5\n'
    '    "第二句"\n'
    '    return\n'
)

MANIFEST = {'statements': [{
    'name': 'mycue', 'block': True, 'label': '闪电动画',
    'columns': [
        {'name': 'intensity', 'choices': ['1', '2', '3'], 'required': True},
        {'name': 'target', 'choices': ['ling', 'lan']},
        {'name': 'notes'},                      # 无候选、非必需
    ]}]}


@pytest.fixture(autouse=True)
def _fresh_cache():
    clear_cache()
    yield
    clear_cache()


def _project(tmp_path, manifest=MANIFEST, script=SCRIPT):
    root = tmp_path / 'proj'
    game = root / 'game'
    game.mkdir(parents=True, exist_ok=True)
    f = game / 'ch.rpy'
    f.write_text(script, encoding='utf-8')
    if manifest is not None:
        (root / 'storygraph.rules.json').write_text(
            json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    return root, f


def _beats(path):
    pr = IncrementalRenPyAnalyzer().analyze(files=[path], graph=None)[0]
    return beats_from_parse_result(pr)


# ---------------------------------------------------------------------------
#  ① Excel：候选注释
# ---------------------------------------------------------------------------

def test_head_cell_gets_a_choices_comment(tmp_path):
    root, path = _project(tmp_path)
    beats = _beats(path)
    x = root / 'game' / 'asset_tables' / 'out.xlsx'
    x.parent.mkdir(parents=True, exist_ok=True)
    export_beats_to_excel(beats, x)

    wb = openpyxl.load_workbook(x)
    ws = wb['Beats']
    col = BEAT_FIELDS.index('SpecHead') + 1
    row = None
    for i in range(2, ws.max_row + 1):
        if ws.cell(row=i, column=col).value:
            row = i
            break
    assert row, '没找到写了头行的那一行'
    cell = ws.cell(row=row, column=col)
    assert cell.comment is not None, '头行那一格没有候选注释'
    text = cell.comment.text
    assert 'mycue' in text
    assert '1 / 2 / 3' in text, text
    assert 'ling / lan' in text, text
    assert '[必需]' in text, text
    assert '闪电动画' in text, '`label` 也要带上（人话说明）'


def test_no_comment_when_nothing_declared(tmp_path):
    """没声明 → 一格注释都不加（不猜候选，也不刷屏）。"""
    root, path = _project(tmp_path, manifest=None)
    beats = _beats(path)
    x = root / 'game' / 'asset_tables' / 'out.xlsx'
    x.parent.mkdir(parents=True, exist_ok=True)
    export_beats_to_excel(beats, x)
    ws = openpyxl.load_workbook(x)['Beats']
    col = BEAT_FIELDS.index('SpecHead') + 1
    for i in range(2, ws.max_row + 1):
        assert ws.cell(row=i, column=col).comment is None


# ---------------------------------------------------------------------------
#  ② 校验：W120 / W121
# ---------------------------------------------------------------------------

def _codes(issues):
    return [i.code for i in issues]


def test_missing_required_param_is_reported(tmp_path):
    root, path = _project(tmp_path, manifest={'statements': [{
        'name': 'mycue', 'block': True,
        'columns': [{'name': 'intensity', 'required': True}]}]})
    beats = _beats(path)
    # 把头行里的 intensity 抹掉，只留 target
    for b in beats:
        if b.spec_name:
            b.spec_head = 'mycue "闪电" target=ling:'
    rep = validate_all(beats, project_root=root)
    assert 'W120' in _codes(rep.issues), [i.format() for i in rep.issues]


def test_value_outside_choices_is_reported(tmp_path):
    root, path = _project(tmp_path)
    beats = _beats(path)
    for b in beats:
        if b.spec_name:
            b.spec_head = 'mycue "闪电" intensity=9 target=ling:'
    rep = validate_all(beats, project_root=root)
    hits = [i for i in rep.issues if i.code == 'W121']
    assert hits, [i.format() for i in rep.issues]
    assert 'intensity' in hits[0].message


def test_declared_and_correct_produces_no_spec_issue(tmp_path):
    """声明齐了、值也在候选里 → 一条都不报（别把正常当异常）。"""
    root, path = _project(tmp_path)
    beats = _beats(path)
    rep = validate_all(beats, project_root=root)
    assert 'W120' not in _codes(rep.issues), _codes(rep.issues)
    assert 'W121' not in _codes(rep.issues), _codes(rep.issues)


def test_undeclared_statement_is_not_validated(tmp_path):
    """没声明就没得校验 —— **不许乱报**。"""
    root, path = _project(tmp_path, manifest=None)
    beats = _beats(path)
    rep = validate_all(beats, project_root=root)
    assert 'W120' not in _codes(rep.issues)
    assert 'W121' not in _codes(rep.issues)


def test_positional_args_are_not_treated_as_named():
    """
    位置参数**不当成**有名字的参数。

    `mycue 3 ling` 里那两个词跟声明里的 `intensity` / `target` 没有
    可推导的对应关系；硬按位置对齐是**猜**，猜错会指着一处不存在的
    错误让人去改。
    """
    vals = spec_arg_values('mycue "闪电" 3 ling')
    assert vals == {}, vals
    vals = spec_arg_values('mycue "闪电" intensity=3 target="ling"')
    assert vals == {'intensity': '3', 'target': 'ling'}, vals


def test_choices_value_is_a_warning_not_an_error(tmp_path):
    """候选是"常见取值"不是白名单 → 只能是**警告**，不能拦。"""
    root, path = _project(tmp_path)
    beats = _beats(path)
    for b in beats:
        if b.spec_name:
            b.spec_head = 'mycue "闪电" intensity=99 target=ling:'
    rep = validate_all(beats, project_root=root)
    hits = [i for i in rep.issues if i.code == 'W121']
    assert hits and hits[0].severity != 'error', hits[0]


# ---------------------------------------------------------------------------
#  ③ 卡片：候选提供者
# ---------------------------------------------------------------------------

def test_card_choices_provider(tmp_path):
    root, path = _project(tmp_path)
    reg = StatementRegistry.for_project(root)
    beats = _beats(path)
    beat = next(b for b in beats if b.spec_name)
    got = spec_arg_choices_for(beat, reg)
    assert got == {'intensity': ['1', '2', '3'],
                   'target': ['ling', 'lan']}, got
    # `notes` 没有候选 → 不进映射（UI 那一栏就是普通输入框）
    assert 'notes' not in got
    _ = path


def test_card_choices_provider_is_empty_when_undeclared(tmp_path):
    root, path = _project(tmp_path, manifest=None)
    reg = StatementRegistry.for_project(root)
    beats = _beats(path)
    beat = next(b for b in beats if b.opaque)
    assert spec_arg_choices_for(beat, reg) == {}
    assert spec_arg_choices(reg, 'mycue') == []
