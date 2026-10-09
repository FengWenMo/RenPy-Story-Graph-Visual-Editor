"""
角色**别名列**（`.txt` 导入用）· 判据
=====================================

用户 2026-10-10 定案：别名放在 `characters.xlsx` 的**「别名」列**，
而且要能**多个别名**（"妈妈 / Mom / MM 都是同一个人"）。

这个文件的判据分两半：

  **A. 匹配语义**（`txt_import.CharacterMap` 的四层）
  **B. 表结构契约**（新列必须登记到位，且**不许**踩两个坑）
     · ⚠ 别名列**不能**进 `_CHAR_IDENTITY_KEYS` —— 那个元组是**重命名识别**用的
       （"两边非空且不同 → 当成换了角色"），人一编辑别名就会被误判成换角色；
     · ⚠ 别名列**不能被「从脚本导入角色表」冲掉**（角色同步是 append-only，
       已有行不重写 —— 这条要是坏了，用户填的别名会静默消失）。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from storygraph.assets.asset_engine import AssetEngine          # noqa: E402
from storygraph.assets.asset_registry import (                  # noqa: E402
    ASSET_FIELD_LABELS, _TABLE_FIELDS, ensure_asset_tables,
    parse_asset_tables, write_dict_to_excel)
from storygraph.authoring.txt_import import (                   # noqa: E402
    CharacterMap, convert_text, split_aliases)

# ==================== A. 匹配语义 ====================

ROWS = {
    'mm': {'ID': 'mm', 'Name': '妈妈', 'Aliases': 'MOM|母亲|老妈'},
    'ld': {'ID': 'ld', 'Name': '校领导', 'Aliases': 'LINGDAO, 领导'},
    'jc': {'ID': 'jc', 'Name': '「蒋忱」', 'Aliases': ''},
}


def _cm():
    return CharacterMap(ROWS)


@pytest.mark.parametrize('sep', ['|', ',', '，', '、', ';', '；', ' '])
def test_aliases_accept_every_common_separator(sep):
    """一个人可以有**多个**别名，分隔符不该逼用户记一种。"""
    assert split_aliases(f'a{sep}b{sep}c') == ['a', 'b', 'c']


def test_split_aliases_handles_empty():
    assert split_aliases(None) == []
    assert split_aliases('') == []
    assert split_aliases('  | ,  ') == []


def test_four_layers_all_resolve():
    cm = _cm()
    assert cm.lookup('mm') == ('mm', 'variable')       # L0
    assert cm.lookup('MM') == ('mm', 'variable')       # L1 归一
    assert cm.lookup('妈妈') == ('mm', 'display')      # L2 显示名
    assert cm.lookup('MOM') == ('mm', 'alias')         # L3 别名
    assert cm.lookup('LINGDAO') == ('ld', 'alias')
    assert cm.lookup('领导') == ('ld', 'alias')
    assert cm.lookup('「蒋忱」') == ('jc', 'display')
    assert cm.lookup('XIAOYI') == (None, '')           # L4 没命中


def test_multiple_aliases_for_one_person():
    """`妈妈`/`MOM`/`母亲`/`老妈` 都是同一个人 —— 全部落到 `mm`。"""
    cm = _cm()
    for token in ('妈妈', 'MOM', '母亲', '老妈', 'mm', 'Mm'):
        assert cm.lookup(token)[0] == 'mm', token


def test_alias_makes_the_conversion_stop_complaining():
    cm = _cm()
    r1 = convert_text('MOM"吃饭了"', characters=cm)
    assert r1.unmapped == {}
    assert '    mm "吃饭了"' in r1.rpy_text
    assert r1.ok


def test_unmapped_warning_points_at_the_alias_column():
    r = convert_text('XIAOYI"姐姐"', characters=_cm())
    assert r.unmapped == {'XIAOYI': 1}
    assert any('别名' in w for w in r.warnings)


def test_conflicting_aliases_are_reported_not_silently_picked():
    """同一个写法指到两个角色 → **报出来**。"""
    cm = CharacterMap({'mm': {'ID': 'mm', 'Name': '妈妈', 'Aliases': 'MOM'},
                       'mom2': {'ID': 'mom2', 'Name': '母亲',
                                'Aliases': 'MOM'}})
    assert cm.conflicts, '撞车必须留下记录'
    r = convert_text('MOM"一"', characters=cm)
    assert any('冲突' in w for w in r.warnings)
    assert r.conflicts


def test_conflict_between_alias_and_someone_elses_display_name():
    cm = CharacterMap({'a': {'ID': 'a', 'Name': '甲', 'Aliases': '乙'},
                       'b': {'ID': 'b', 'Name': '乙', 'Aliases': ''}})
    assert cm.conflicts
    # 但映射按"**强层赢**"定：显示名（正名）压过别人的别名
    assert cm.lookup('乙') == ('b', 'display')


def test_same_character_from_several_layers_is_not_a_conflict():
    """`narrator` 的变量名与显示名相同 —— 这不是冲突。"""
    cm = CharacterMap({'narrator': {'ID': 'narrator', 'Name': 'narrator',
                                    'Aliases': ''}})
    assert cm.conflicts == []


def test_candidates_are_offered_for_unmapped_tokens():
    """未映射时给**候选**让人挑（我们不猜，但要帮人省事）。"""
    cm = CharacterMap({'mom': {'ID': 'mom', 'Name': '妈妈', 'Aliases': ''},
                       'mom2': {'ID': 'mom2', 'Name': '母亲', 'Aliases': ''}})
    assert 'mom' in cm.candidates('MOMX')


# ==================== B. 表结构契约 ====================

def test_alias_column_is_registered_in_the_table_schema():
    assert 'Aliases' in _TABLE_FIELDS['character'], \
        '新表要带「别名」列（否则用户没地方填）'
    assert ASSET_FIELD_LABELS.get('Aliases') == '别名', \
        '中文表头边界翻译要登记（否则表头对不上）'


def test_alias_column_is_NOT_in_the_identity_keys():
    """
    ⚠ 别名是**人写字段**，不能参与"重命名识别" —— 否则用户一填别名，
    引擎就会把它当成"换了角色"（`_looks_like_rename` 的判据是"两边非空且不同"）。
    """
    assert 'Aliases' not in AssetEngine._CHAR_IDENTITY_KEYS


def test_a_new_character_table_carries_the_alias_column(tmp_path):
    reg = parse_asset_tables(tmp_path)          # 空目录 → 建表
    assert reg is not None
    p = tmp_path / 'characters.xlsx'
    write_dict_to_excel(p, {'mm': {'ID': 'mm', 'Name': '妈妈',
                                   'Aliases': 'MOM'}},
                        _TABLE_FIELDS['character'])
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    headers = [c.value for c in ws[1]]
    assert '别名' in headers, headers        # 写出去是**中文表头**
    assert headers.index('别名') < headers.index('备注')


def test_alias_survives_a_read_write_roundtrip(tmp_path):
    p = tmp_path / 'characters.xlsx'
    write_dict_to_excel(p, {'mm': {'ID': 'mm', 'Name': '妈妈',
                                   'Aliases': 'MOM|母亲'}},
                        _TABLE_FIELDS['character'])
    rows = parse_asset_tables(tmp_path).characters
    assert rows['mm']['Aliases'] == 'MOM|母亲'
    write_dict_to_excel(p, rows, _TABLE_FIELDS['character'])
    again = parse_asset_tables(tmp_path).characters
    assert again['mm']['Aliases'] == 'MOM|母亲'


def test_an_unknown_column_is_kept_as_is(tmp_path):
    """用户手加的列原样保留（工具既有规矩）—— 别名列也走这条路。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(['ID', 'Name', 'Color', 'ImageTag', 'Notes', 'Aliases', '我加的列'])
    ws.append(['mm', '妈妈', None, None, '', 'MOM', 'x'])
    wb.save(tmp_path / 'characters.xlsx')
    rows = parse_asset_tables(tmp_path).characters
    assert rows['mm']['Aliases'] == 'MOM'
    assert rows['mm']['我加的列'] == 'x'


@pytest.mark.skipif(not (Path(r'D:\千连万理\RenpyTest\sr_general')).is_dir(),
                    reason='真机项目不在本机')
def test_the_real_character_table_still_reads(tmp_path):
    """
    真机角色表是**英文表头的老表**（没有别名列）—— 必须照旧能读
    （新列是"以后填"，不是"必须存在"）。
    """
    from storygraph.pipeline.build import load_asset_registry
    reg = load_asset_registry(Path(r'D:\千连万理\RenpyTest\sr_general'))
    rows = reg.characters
    assert len(rows) > 10
    cm = CharacterMap(rows)
    assert cm.lookup('JC')[0] == 'jc'
    assert cm.lookup('妈妈')[0] == 'mm'         # 真机里 mm 的显示名就是「妈妈」
