"""
tests/test_table_export_inplace_guard.py
========================================
**界面那条「从剧本表导出」的路，也必须过「原地覆盖 + 非增量 = 拒写」这道闸。**

2026-10-21 实测（P1，P0 的同一个病换了个入口）：

  1. 上一轮导出过这一章 → 文件首行有生成标记；
  2. 程序员在同一个 `.rpy` 里加了 `pause 0.5` / 一句台词 / `$ flag = True`；
  3. 策划在表里改一句 → 界面「从剧本表导出」。

  `export_table` → `apply_export` 报 `ok=True / written=True`，
  程序员那三行**全没了**，源文件被重写、首行还多了一条重复的生成标记。

为什么只有界面这条路漏（CLI 那条早就拦得住）：

  · `export_table` 的**预检**（`evaluate_generation`）拿得到 `source_text`，
    但它只查"顶层语句会不会丢"（`OPAQUE_LOST`）——label 里的 `pause` / `$`
    不在它的判据里；
  · 真正落盘的 `apply_export` 只往下传了 `merged_from`，**没传
    `source_text` / `source_path`** —— 而门禁那条"原地 + 非增量 → 拒写"
    要先有 `source_text` 才进得去；
  · 于是只剩手工护栏在判，而它按"有没有生成标记"判 —— 带标记的文件它
    一律放行（那正是"再导出一次"的正常路径）。

判据（每条都用文件内容，不看回执）：
  · 表过期 + 原地 → **不写**，源文件逐字节不变，程序员的 3 行还在；
  · 表没改 → 原地照旧能写（不能把正常路堵死）；
  · 结果包里带 `source_path`（落盘那一步全靠它）。
"""

from pathlib import Path

MARK = '# 由剧本表自动生成，请勿在此文件手工改动结构'

#: 上一轮导出过的样子：首行带生成标记。
SENTINEL = f'''{MARK}

label ch01:
    e "第一句"
    e "第二句"
    return
'''

#: 程序员后来加的三行（表里没有）。
PROG_EXTRA = ('    pause 0.5\n'
              '    e "程序员新加的一句。"\n'
              '    $ flag = True\n')


def _project(tmp_path):
    root = tmp_path / 'proj'
    game = root / 'game'
    game.mkdir(parents=True)
    script = game / 'ch01.rpy'
    script.write_text(SENTINEL, encoding='utf-8')
    return root, script


def _set_cell(table: Path, old: str, new: str) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(str(table))
    ws = wb['Beats']
    heads = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1]) if c.value}
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, heads['台词']).value or '').strip() == old:
            ws.cell(r, heads['台词']).value = new
            wb.save(str(table))
            return 1
    wb.save(str(table))
    return 0


def _stale(tmp_path):
    """造出"表过期"：先导表，程序员再改源，策划再改表。"""
    from storygraph.app import tables
    root, script = _project(tmp_path)
    res = tables.generate_table(script, project_root=root)
    assert res['ok'], res['error']
    table = res['table_path']

    cur = script.read_text(encoding='utf-8')
    script.write_text(cur.replace('    e "第一句"\n',
                                  '    e "第一句"\n' + PROG_EXTRA),
                      encoding='utf-8')
    before = script.read_bytes()
    assert _set_cell(table, '第二句', '改过的一句') == 1
    return root, script, table, before


class TestStaleSourceInPlace:
    def test_apply_export_refuses_and_leaves_the_file_alone(self, tmp_path):
        from storygraph.app import tables
        root, script, table, before = _stale(tmp_path)

        res = tables.export_table(table, project_root=root, source_rpy=script)
        # 目标就是这一章自己（`explicit_src` 那条规则）
        assert Path(res['out_path']).resolve() == script.resolve()
        done = tables.apply_export(res, project_root=root)

        assert done['ok'] is False and done['written'] is False, (
            '表过期了却写了盘 —— 这就是"从剧本表导出"吃掉程序员代码的那条路')
        assert script.read_bytes() == before, '源文件被改了，应当一个字节不动'

    def test_programmer_lines_survive(self, tmp_path):
        from storygraph.app import tables
        root, script, table, _ = _stale(tmp_path)
        res = tables.export_table(table, project_root=root, source_rpy=script)
        tables.apply_export(res, project_root=root)
        now = script.read_text(encoding='utf-8')
        assert 'pause 0.5' in now
        assert '程序员新加的一句。' in now
        assert '$ flag = True' in now
        assert '改过的一句' not in now, '拒写了却还是把表里的改动写进去了'

    def test_refusal_says_why(self, tmp_path):
        from storygraph.app import tables
        root, script, table, _ = _stale(tmp_path)
        res = tables.export_table(table, project_root=root, source_rpy=script)
        done = tables.apply_export(res, project_root=root)
        why = str(done.get('error') or '')
        # 不能只说"这是手写文件"—— 真因是"表和源码对不上"
        assert '对不上' in why or '增量' in why, f'拒绝理由没说清原因：{why}'

    def test_result_package_carries_source_path(self, tmp_path):
        """`source_path` 必须随结果包走到落盘那一步（否则闸门进不去）。"""
        from storygraph.app import tables
        root, script, table, _ = _stale(tmp_path)
        res = tables.export_table(table, project_root=root, source_rpy=script)
        assert res.get('source_path')
        assert Path(res['source_path']).resolve() == script.resolve()


class TestNormalEditStillWrites:
    def test_incremental_in_place_still_lands(self, tmp_path):
        """表与源对得上时，界面那条路照旧能写（只改那一行）。"""
        from storygraph.app import tables
        root, script = _project(tmp_path)
        res = tables.generate_table(script, project_root=root)
        table = res['table_path']
        assert _set_cell(table, '第二句', '改过的一句') == 1

        out = tables.export_table(table, project_root=root, source_rpy=script)
        assert out['blocking'] == [], out['blocking']
        done = tables.apply_export(out, project_root=root)
        assert done['ok'] is True, done.get('error')
        now = script.read_text(encoding='utf-8')
        assert '改过的一句' in now
        assert '第二句' not in now
        # 原来的生成标记还在（不该被再插一条）
        assert now.count(MARK) == 1, f'生成标记重复了：{now!r}'


class TestRemovedLinesAreNamed:
    """
    回执必须点名"删掉了源文件里哪几行"。

    2026-10-21：`changed_lines` 把三种动作混成一个数 —— 换一句台词、插一行、
    **删掉源里一行 `show`**。用户清空某一拍的舞台列时，源文件里那行会跟着走；
    只报"改动 2 行"，他的第一反应是"我只改了一句，怎么少了一行"。
    """

    def test_patch_result_carries_the_deleted_lines(self, tmp_path):
        from storygraph.app import tables
        root = tmp_path / 'proj'
        (root / 'game').mkdir(parents=True)
        script = root / 'game' / 'ch01.rpy'
        script.write_text('label ch01:\n'
                          '    scene bg room\n'
                          '    show ling happy\n'
                          '    ling "第一句"\n'
                          '    return\n', encoding='utf-8')
        table = tables.generate_table(script, project_root=root)['table_path']

        import openpyxl
        wb = openpyxl.load_workbook(str(table))
        ws = wb['Beats']
        heads = {str(c.value).strip(): i + 1 for i, c in enumerate(ws[1])
                 if c.value}
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, heads['立绘']).value or '').strip():
                ws.cell(r, heads['立绘']).value = None   # 策划清空立绘
        wb.save(str(table))

        res = tables.export_table(table, project_root=root, source_rpy=script)
        assert res['used_patch'] is True, res['note']
        assert res['removed_preview'], '删了整行却没记下来'
        assert any('show ling happy' in ln for ln in res['removed_preview'])
        # 回执的人话里也要有
        assert '删掉' in res['note'], res['note']

    def test_cli_report_points_at_the_deleted_lines(self):
        """`cli. _report_incremental` 的输出里得看得见那几行。"""
        import io
        from storygraph.pipeline.compile import CompileResult
        from storygraph.cli import _report_incremental

        result = CompileResult(rpy_text='label a:/n    return\n',
                               changed_lines=2,
                               used_patch=True,
                               removed_preview=('show ling happy',))
        buf = io.StringIO()
        _report_incremental(result, stream=buf)
        out = buf.getvalue()
        assert '删掉' in out
        assert 'show ling happy' in out
