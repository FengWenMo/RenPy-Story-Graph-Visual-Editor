"""
asset_registry.py
==================
资产扫描、Excel 读写、解析与校验。

对外提供的接口（被 asset_engine.py 使用）：
  · AssetRegistry           资产注册表数据结构
  · scan_assets             从目录扫描资产（图片/音频）
  · write_asset_excel       将扫描结果写入 Excel
  · parse_asset_tables      从 Excel 目录构建 AssetRegistry
  · load_hooks              加载自定义钩子模块
  · analyze_references      扫描 .rpy 脚本填充 registry.references
  · validate_registry       校验文件存在性、依赖引用、颜色格式
"""

import re
from pathlib import Path
from collections import defaultdict
import openpyxl
import importlib.util


# ===================== 常量 =====================
SUPPORTED_IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif'}
SUPPORTED_AUDIO_EXT = {'.ogg', '.mp3', '.wav', '.m4a', '.flac', '.aac', '.opus'}

CATEGORY_KEYWORDS = {
    'background': ['bg', 'background', 'backdrop', 'scene'],
    'cg': ['cg', 'event', 'cutscene', 'illust'],
    'character': ['character', 'char', 'sprite', 'avatar', 'portrait',
                  'figure', 'people'],
    'ui': ['gui', 'ui', 'dutui', 'interface', 'system', 'icon', 'button',
           'menu'],
    'music': ['bgm', 'music', 'theme', 'loop', 'ost'],
    'sound': ['sfx', 'sound', 'effect', 'se', 'voice', 'click', 'hit'],
}


# ===================== 数据结构 =====================

class AssetRegistry:
    """
    资产注册表。

    · characters / images / transforms / music / sounds: {id: row_dict}
    · custom: {table_name: {id: row_dict}}
    · references: {asset_id: ['file.rpy:line', ...]}
    · errors / warnings: 校验过程中收集的诊断信息
    """

    def __init__(self):
        self.characters = {}
        self.images = {}
        self.transforms = {}
        self.music = {}
        self.sounds = {}
        self.custom = {}
        self.references = defaultdict(list)
        self.errors = []
        self.warnings = []

    def get_asset_ids(self, asset_type):
        mapping = {
            'character': self.characters,
            'image': self.images,
            'transform': self.transforms,
            'music': self.music,
            'sound': self.sounds,
        }
        return list(mapping.get(asset_type, {}).keys())


# ===================== Excel 读写 =====================

def read_excel_rows(file_path) -> list:
    """读取 Excel 首个 sheet，返回 [{header: value, ...}, ...]。"""
    file_path = Path(file_path)
    if not file_path.exists():
        return []
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_dict = {}
        for i, header in enumerate(headers):
            if header is not None:
                row_dict[header] = row[i] if i < len(row) else None
        rows.append(row_dict)
    wb.close()
    return rows


def write_dict_to_excel(file_path, data_dict: dict, field_order: list):
    """将 {id: row_dict} 写入 Excel。"""
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(field_order)
    for _asset_id, row_data in data_dict.items():
        excel_row = [row_data.get(field, '') for field in field_order]
        ws.append(excel_row)
    wb.save(file_path)


# ===================== 智能扫描 =====================

def _detect_category(rel_path: Path) -> str:
    """根据路径中的关键词推断资产类别。"""
    path_str = rel_path.as_posix().lower()
    for cat, kws in CATEGORY_KEYWORDS.items():
        if any(kw in path_str for kw in kws):
            # bg 与 bgm 冲突时，优先 bgm（音乐）
            if cat == 'background' and 'bgm' in path_str:
                continue
            return cat
    parent = rel_path.parent.name.lower()
    for cat, kws in CATEGORY_KEYWORDS.items():
        if parent in kws:
            return cat
    return 'other'


def _make_unique_id(stem: str, parent_dirs: list, existing_ids: set) -> str:
    """生成唯一 ID：优先用 stem，冲突时逐步添加父目录名。"""
    if stem not in existing_ids:
        existing_ids.add(stem)
        return stem
    for i in range(len(parent_dirs) - 1, -1, -1):
        new_id = f"{parent_dirs[i]}_{stem}"
        if new_id not in existing_ids:
            existing_ids.add(new_id)
            return new_id
    counter = 1
    while True:
        new_id = f"{stem}_{counter}"
        if new_id not in existing_ids:
            existing_ids.add(new_id)
            return new_id
        counter += 1


def scan_assets(scan_root, include_ui=True, audio_dirs=None) -> dict:
    """
    扫描目录下的资产，返回 {'images': [...], 'audio': {'music': [...], 'sound': [...]}}。

    scan_root:   扫描根目录（Path 或 str）
    include_ui:  是否包含 UI 相关图片
    audio_dirs:  自定义音频目录列表；None 则自动探测
    """
    scan_root = Path(scan_root)
    if not scan_root.exists():
        return {'images': [], 'audio': {'music': [], 'sound': []}}

    image_assets = []
    audio_music = []
    audio_sound = []

    # 音频目录自动探测
    if audio_dirs is None:
        candidates = [scan_root / 'audio', scan_root / 'music',
                      scan_root / 'sounds', scan_root / 'bgm']
        existing = [d for d in candidates if d.exists()]
        audio_dirs = existing if existing else [scan_root]

    existing_image_ids = set()

    # ---------- 图片扫描 ----------
    for file in scan_root.rglob('*'):
        if not file.is_file() or file.suffix.lower() not in SUPPORTED_IMAGE_EXT:
            continue
        rel_path = file.relative_to(scan_root)
        rel_str = rel_path.as_posix().lower()

        if not include_ui:
            if any(kw in rel_str
                   for kw in ('gui', 'ui', 'dutui', 'interface', 'system')):
                continue

        category = _detect_category(rel_path)
        parent_dirs = [p.name for p in rel_path.parents if p != scan_root]

        tag = ''
        attributes = file.stem
        if category == 'character':
            char_idx = None
            for i, part in enumerate(rel_path.parts):
                if any(kw in part.lower()
                       for kw in CATEGORY_KEYWORDS['character']):
                    char_idx = i
                    break
            if char_idx is not None and char_idx + 1 < len(rel_path.parts):
                possible_name = rel_path.parts[char_idx + 1]
                if possible_name == file.name:
                    tag = file.stem.split('_')[0]
                else:
                    tag = possible_name
            else:
                tag = file.stem.split('_')[0]
        elif category == 'background':
            tag = 'bg'
        elif category == 'cg':
            tag = 'cg'
        elif category == 'ui':
            tag = 'ui'

        unique_id = _make_unique_id(file.stem, parent_dirs, existing_image_ids)
        image_assets.append({
            'ID': unique_id,
            'Path': rel_path.as_posix(),
            'Tag': tag,
            'Attributes': attributes,
            'Category': category,
            'Notes': '',
        })

    # ---------- 音频扫描 ----------
    for adir in audio_dirs:
        adir = Path(adir)
        if not adir.exists():
            continue
        for file in adir.rglob('*'):
            if not file.is_file() or file.suffix.lower() not in SUPPORTED_AUDIO_EXT:
                continue
            try:
                rel_path = file.relative_to(scan_root)
            except ValueError:
                rel_path = file
            rel_str = rel_path.as_posix().lower()
            stem = file.stem
            if any(kw in rel_str for kw in CATEGORY_KEYWORDS['music']):
                audio_music.append({
                    'ID': f"bgm_{stem}" if not stem.startswith('bgm_')
                          else stem,
                    'Path': rel_path.as_posix(),
                    'Loop': True,
                    'Volume': 1.0,
                    'Notes': '',
                })
            else:
                audio_sound.append({
                    'ID': f"sfx_{stem}" if not stem.startswith('sfx_')
                          else stem,
                    'Path': rel_path.as_posix(),
                    'Notes': '',
                })

    return {
        'images': image_assets,
        'audio': {'music': audio_music, 'sound': audio_sound},
    }


# ===================== Excel 生成 =====================

def _merge_existing(existing: dict, asset: dict, auto_fields: list) -> dict:
    """若资产已在 Excel 中存在，保留用户编辑字段，仅覆盖自动字段。"""
    aid = asset['ID']
    if aid in existing:
        merged = existing[aid].copy()
        for key in auto_fields:
            merged[key] = asset.get(key, '')
        return merged
    return asset


def write_asset_excel(assets_dict: dict, output_dir, split_by_category=False):
    """将 scan_assets 的结果写入 Excel（支持增量更新）。"""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # ---------- 图片 ----------
    image_assets = assets_dict['images']
    if split_by_category:
        groups = defaultdict(list)
        for a in image_assets:
            groups[a['Category']].append(a)
        for cat, assets in groups.items():
            file_path = output_dir / f'images_{cat}.xlsx'
            existing = {r['ID']: r for r in read_excel_rows(file_path) if r.get('ID')}
            new_data = {}
            for a in assets:
                new_data[a['ID']] = _merge_existing(
                    existing, a, ['Path', 'Tag', 'Attributes', 'Category'])
            write_dict_to_excel(
                file_path, new_data,
                ['ID', 'Path', 'Tag', 'Attributes', 'Category', 'Notes'])
    else:
        file_path = output_dir / 'images.xlsx'
        existing = {r['ID']: r for r in read_excel_rows(file_path) if r.get('ID')}
        new_data = {}
        for a in image_assets:
            new_data[a['ID']] = _merge_existing(
                existing, a, ['Path', 'Tag', 'Attributes', 'Category'])
        write_dict_to_excel(
            file_path, new_data,
            ['ID', 'Path', 'Tag', 'Attributes', 'Category', 'Notes'])

    # ---------- 音乐 ----------
    music_file = output_dir / 'music.xlsx'
    existing = {r['ID']: r for r in read_excel_rows(music_file) if r.get('ID')}
    new_data = {}
    for a in assets_dict['audio']['music']:
        new_data[a['ID']] = _merge_existing(existing, a, ['Path', 'Loop', 'Volume'])
    write_dict_to_excel(
        music_file, new_data,
        ['ID', 'Path', 'Loop', 'Volume', 'Notes'])

    # ---------- 音效 ----------
    sound_file = output_dir / 'sounds.xlsx'
    existing = {r['ID']: r for r in read_excel_rows(sound_file) if r.get('ID')}
    new_data = {}
    for a in assets_dict['audio']['sound']:
        new_data[a['ID']] = _merge_existing(existing, a, ['Path'])
    write_dict_to_excel(
        sound_file, new_data,
        ['ID', 'Path', 'Notes'])


# ===================== 钩子加载 =====================

def load_hooks(config_path=None, project_root=None):
    """
    加载自定义钩子模块。

    config_path=None 时按顺序搜索：
      1. project_root / 'asset_import_hooks.py'（推荐）
      2. CWD / 'asset_import_hooks.py'（兼容旧行为）
    均不存在则返回空字典。

    支持两个钩子：
      - pre_process_row(table_name, row) -> row
      - post_process_registry(registry) -> None
    """
    candidates = []
    if config_path is not None:
        candidates.append(Path(config_path))
    else:
        if project_root is not None:
            candidates.append(Path(project_root) / 'asset_import_hooks.py')
        candidates.append(Path('asset_import_hooks.py'))

    target = None
    for c in candidates:
        if c.exists():
            target = c
            break

    if target is None:
        return {}

    try:
        spec = importlib.util.spec_from_file_location(
            'asset_import_hooks', str(target)
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return {
            'pre_process_row': getattr(module, 'pre_process_row', None),
            'post_process_registry': getattr(module, 'post_process_registry', None),
        }
    except Exception as e:
        print(f"[asset_registry] 加载钩子失败: {e}")
        return {}


# ===================== Excel 解析 =====================

# 标准表定义：(registry 属性名, 文件名)
_STANDARD_TABLES = {
    'characters': 'characters.xlsx',
    'images': 'images.xlsx',
    'transforms': 'transforms.xlsx',
    'music': 'music.xlsx',
    'sounds': 'sounds.xlsx',
}


def parse_asset_tables(asset_dir, hooks=None, project_root=None) -> AssetRegistry:
    """
    解析资产表目录，构建 AssetRegistry。

    asset_dir:    存放 xlsx 的目录
    hooks:        load_hooks 的返回值
    project_root: 项目根目录；None 则默认 asset_dir.parent
    """
    asset_dir = Path(asset_dir)
    registry = AssetRegistry()
    if hooks is None:
        hooks = {}
    if project_root is None:
        project_root = asset_dir.parent

    pre_hook = hooks.get('pre_process_row')

    # ---------- 标准表 ----------
    for key, filename in _STANDARD_TABLES.items():
        file = asset_dir / filename
        if not file.exists():
            continue
        rows = read_excel_rows(file)
        for row in rows:
            if pre_hook is not None:
                try:
                    row = pre_hook(key, row)
                except Exception as e:
                    registry.warnings.append(
                        f"{filename}: pre_process_row 钩子异常 {e}")
            asset_id = row.get('ID')
            if not asset_id:
                registry.errors.append(f"{filename}: 缺少ID字段")
                continue
            target_dict = getattr(registry, key)
            if asset_id in target_dict:
                registry.errors.append(f"{filename}: ID '{asset_id}' 重复")
            target_dict[asset_id] = row
            custom_fields = {k: v for k, v in row.items()
                             if k.startswith('Custom_')}
            if custom_fields:
                row['_custom'] = custom_fields

    # ---------- 自定义表 ----------
    for file in asset_dir.glob('custom_*.xlsx'):
        table_name = file.stem.replace('custom_', '')
        rows = read_excel_rows(file)
        custom_dict = {}
        for row in rows:
            if 'ID' not in row:
                registry.warnings.append(f"{file.name}: 无ID列，跳过该行")
                continue
            custom_dict[row['ID']] = row
        registry.custom[table_name] = custom_dict

    # ---------- 后处理钩子 ----------
    post_hook = hooks.get('post_process_registry')
    if post_hook is not None:
        try:
            post_hook(registry)
        except Exception as e:
            registry.warnings.append(f"post_process_registry 钩子异常 {e}")

    validate_registry(registry, project_root)
    return registry


# ===================== 校验 =====================

_COLOR_PATTERN = re.compile(r'^#(?:[0-9a-fA-F]{3}){1,2}$')


def validate_registry(registry: AssetRegistry, project_root):
    """检查文件存在性、依赖引用、颜色格式等。"""
    project_root = Path(project_root)

    for img_id, data in registry.images.items():
        p = data.get('Path')
        if p and not (project_root / p).exists():
            registry.warnings.append(f"图片 '{img_id}' 文件不存在: {p}")

    for type_key in ('music', 'sounds'):
        for aid, data in getattr(registry, type_key).items():
            p = data.get('Path')
            if p and not (project_root / p).exists():
                registry.warnings.append(f"{type_key} '{aid}' 文件不存在: {p}")

    image_tags = {d.get('Tag') for d in registry.images.values() if d.get('Tag')}
    for char_id, data in registry.characters.items():
        if data.get('ImageTag') and data['ImageTag'] not in image_tags:
            registry.warnings.append(
                f"角色 '{char_id}' 的 ImageTag '{data['ImageTag']}' "
                f"未找到对应图片标签")
        if data.get('DefaultTransform') and \
                data['DefaultTransform'] not in registry.transforms:
            registry.warnings.append(
                f"角色 '{char_id}' 的 DefaultTransform "
                f"'{data['DefaultTransform']}' 未定义")

    for char_id, data in registry.characters.items():
        color = data.get('Color')
        if color and not _COLOR_PATTERN.match(str(color)):
            registry.warnings.append(
                f"角色 '{char_id}' 的颜色 '{color}' 格式无效")


# ===================== 脚本引用分析 =====================

_RENPY_ASSET_PATTERNS = [
    # show / scene / hide image_id
    (re.compile(r'^\s*(?:show|scene|hide)\s+([a-zA-Z_]\w*)',
                re.MULTILINE), 'image'),
    # play / queue / stop music "bgm_xxx"
    (re.compile(r'^\s*(?:play|queue|stop)\s+music\s+["\']?([a-zA-Z_]\w*)',
                re.MULTILINE), 'music'),
    # play / queue / stop sound "sfx_xxx"
    (re.compile(r'^\s*(?:play|queue|stop)\s+sound\s+["\']?([a-zA-Z_]\w*)',
                re.MULTILINE), 'sound'),
    # show ... at transform_id
    (re.compile(r'\bat\s+([a-zA-Z_]\w*)', re.MULTILINE), 'transform'),
]


def analyze_references(script_files, registry: AssetRegistry):
    """
    扫描 .rpy 脚本中的资产引用，填充 registry.references。

    registry.references: {asset_id: ['file.rpy:line', ...]}
    """
    if registry.references is None:
        registry.references = defaultdict(list)
    elif not isinstance(registry.references, defaultdict):
        converted = defaultdict(list)
        for k, v in registry.references.items():
            converted[k] = list(v)
        registry.references = converted

    for fp in script_files:
        fp = Path(fp)
        if not fp.exists():
            continue
        try:
            text = fp.read_text(encoding='utf-8', errors='ignore')
        except Exception:
            continue

        for pattern, _kind in _RENPY_ASSET_PATTERNS:
            for match in pattern.finditer(text):
                asset_id = match.group(1)
                line_no = text[:match.start()].count('\n') + 1
                loc = f"{fp.name}:{line_no}"
                if loc not in registry.references[asset_id]:
                    registry.references[asset_id].append(loc)


# ===================== 命令行自测 =====================

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("用法: python asset_registry.py <扫描目录> [输出目录]")
        sys.exit(0)

    scan_root = Path(sys.argv[1])
    output_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else Path.cwd() / 'asset_tables'

    print(f"扫描目录: {scan_root}")
    assets = scan_assets(scan_root, include_ui=False)
    write_asset_excel(assets, output_dir, split_by_category=False)
    print(f"已生成资产表到 {output_dir}")

    reg = parse_asset_tables(output_dir)
    print(f"图片: {len(reg.images)}, 音乐: {len(reg.music)}, 音效: {len(reg.sounds)}")
    if reg.errors:
        print("错误:", reg.errors)
    if reg.warnings:
        print("警告:", reg.warnings[:10],
              "..." if len(reg.warnings) > 10 else "")